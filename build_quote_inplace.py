# -*- coding: utf-8 -*-
import openpyxl, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
FONT="맑은 고딕"; ACC='_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
DARK="404040";GREEN_HL="82DC28";OLIVE="62A81B";DGREEN="082A08";LGRN="EAF3DA";WHITE="FFFFFF";INK="262626"
thin=Side(style="thin",color="BFBFBF");bd=Border(left=thin,right=thin,top=thin,bottom=thin)
SRC="/root/.claude/uploads/bc09908a-8d16-5be3-806b-53899e2cede8/ff188eec-BAT_________260618_____.xlsx"
wb=openpyxl.load_workbook(SRC)
ds=wb["견적상세"]; ms=wb["월별 액션플랜"]; qp=wb["견적제안"]
def Cc(r,c,v=None,b=False,col=INK,fill=None,al="center",fmt=None,sz=10,wrap=False,border=True):
    cc=ds.cell(r,c)
    if v is not None: cc.value=v
    cc.font=Font(name=FONT,size=sz,bold=b,color=col)
    if fill: cc.fill=PatternFill("solid",fgColor=fill)
    cc.alignment=Alignment(horizontal=al,vertical="center",wrap_text=wrap)
    if border: cc.border=bd
    if fmt: cc.number_format=fmt
    return cc

# A) 값 수정
for row,val in [(7,20),(16,40),(31,24),(40,20),(49,28),(64,34)]:
    assert ds.cell(row,2).value=="마이크로 (UGC)",row; ds.cell(row,7).value=val
assert ds.cell(65,2).value=="나노"; ds.cell(65,7).value=44
for row,val in [(27,20),(36,30),(45,40),(60,50),(74,60)]:
    assert ds.cell(row,2).value=="어필리에이트",row; ds.cell(row,7).value=val
for i,v in enumerate([20,40,24,20,28,34]): ms.cell(5,3+i).value=v
ms.cell(6,8).value=44
ms.cell(19,1).value="어필리에이트 (월 활성)"
ms.cell(19,9).value="=SUM(C19:H19)"; ms.cell(19,10).value="=SUM(C19:H19)*B19"

# 견적제안 숫자 동기화
def rep(cell,subs):
    v=qp[cell].value
    for a,b in subs: v=v.replace(a,b)
    qp[cell].value=v
rep("E13",[("총 500건","총 471건")])
rep("E15",[("178,000,000","163,000,000"),("마이크로 197","마이크로 166"),("나노 198","나노 200")])
rep("E18",[("38,200,000","53,200,000"),("어필리에이트 60","어필리에이트 월활성 210인·월")])

# B) 전체합계 상단 삽입
N=20
ds.insert_rows(5,N)
pat=re.compile(r'(\$?[A-Z]{1,2}\$?)(\d+)')
def bump(m): return m.group(1)+str(int(m.group(2))+N)
for row in ds.iter_rows(min_row=5+N, max_row=ds.max_row):
    for c in row:
        if isinstance(c.value,str) and c.value.startswith("="):
            c.value=pat.sub(bump, c.value)

META={"마이크로 (UGC)":("IG·TT","직장인 루틴형 2535","#올더베러 #오늘도베러",500000,False),
"나노":("IG","웰니스 입문형 2030","#올더베러 #오늘도베러 #웰니스입문",250000,False),
"매크로":("IG·TT·YT","헬시플레저 광역 4050+","#올더베러 #올영세일",800000,False),
"KOL 무가시딩":("IG·YT","30대 웰니스 KOL","#올더베러 #오늘도베러",100000,False),
"X (트위터) 시딩":("X","리뷰·밈 2030","#올더베러 #올영템",500000,False),
"네이버 블로그":("Blog","정보탐색 검색자","#올더베러 #웰니스구미추천",300000,False),
"EGC 콘텐츠":("IG·TT","정보 신뢰 구매고려층","#올더베러 #직원피셜",1000000,False),
"Half EGC 콘텐츠":("IG·TT","건기식 관심 직장인","#올더베러 #웰니스루틴",1250000,False),
"콘텐츠 제작 (KV·영상·이미지)":("제작","—","—",1550000,False),
"파워페이지":("SNS","구매 고의도 검색자","#올영추천템 #웰니스추천",800000,False),
"커뮤니티 체험단":("뷰티앱·커뮤니티","카테고리 관여 高 사용자","#올리브영 #웰니스어워드",5000000,True),
"챌린저스":("챌린저스 앱","올영 액티브 유저","#올영1위 #올더베러챌린지",5000000,True),
"어필리에이트":("올영 쇼핑 큐레이터","세일즈 경험 크리에이터","#올영세일 #장바구니",100000,False)}
AGG=[("마이크로 (UGC)",166),("나노",200),("매크로",15),("KOL 무가시딩",60),("X (트위터) 시딩",15),
("네이버 블로그",15),("EGC 콘텐츠",12),("Half EGC 콘텐츠",6),("콘텐츠 제작 (KV·영상·이미지)",6),
("파워페이지",9),("커뮤니티 체험단",3),("챌린저스",2),("어필리에이트",210)]
hdr=["항목","채널","캠페인 및 활용 목적 / 앵글","타겟 페르소나","필수 해시태그","목표수량","단가","마크업(20%)","총비용"]
hf={7:GREEN_HL,8:OLIVE,9:OLIVE,10:DGREEN}; ht={7:INK,8:WHITE,9:WHITE,10:WHITE}
ds.merge_cells(start_row=5,start_column=2,end_row=5,end_column=10)
Cc(5,2,"■ 전체 합계 (6개월 취합)",b=True,col=DGREEN,fill=LGRN,al="left")
for i,h in enumerate(hdr,2): Cc(6,i,h,b=True,col=ht.get(i,WHITE),fill=hf.get(i,DARK),al="center",wrap=True)
r=7
for nm,qty in AGG:
    ch,per,tag,price,mk=META[nm]
    label=nm+(" (월 활성 합)" if nm=="어필리에이트" else "")
    Cc(r,2,label,b=True,al="left",wrap=True); Cc(r,3,ch); Cc(r,4,"6개월 합산",al="left",sz=9,col="808080")
    Cc(r,5,per,al="left",wrap=True,sz=9); Cc(r,6,tag,al="left",wrap=True,sz=9,col=DGREEN)
    Cc(r,7,qty,b=True); Cc(r,8,price,fmt=ACC)
    Cc(r,9,(f"=H{r}*G{r}*0.2" if mk else 0),fmt=ACC,col=(DGREEN if mk else INK))
    Cc(r,10,(f"=H{r}*G{r}*1.2" if mk else f"=H{r}*G{r}"),fmt=ACC,b=True)
    r+=1
ds.merge_cells(start_row=20,start_column=2,end_row=20,end_column=6)
Cc(20,2,"총 합계 (원고료+마크업, VAT 별도)",b=True,col=WHITE,fill=DARK,al="right")
Cc(20,7,"=SUM(G7:G19)",b=True,col=WHITE,fill=DARK); Cc(20,8,"",fill=DARK)
Cc(20,9,"=SUM(I7:I19)",fmt=ACC,b=True,col=WHITE,fill=DARK); Cc(20,10,"=SUM(J7:J19)",fmt=ACC,b=True,col=WHITE,fill=DGREEN)
Cc(21,9,"부가세(10%)",b=True,al="right"); Cc(21,10,"=J20*0.1",fmt=ACC,b=True)
Cc(22,9,"합계(VAT 포함)",b=True,al="right",col=WHITE,fill=DARK); Cc(22,10,"=J20*1.1",fmt=ACC,b=True,col=WHITE,fill=DARK)
ds.merge_cells(start_row=24,start_column=2,end_row=24,end_column=10)
Cc(24,2,"■ 월별 상세 (7월 ~ 12월)",b=True,col=DGREEN,fill=LGRN,al="left")

wb.save("proposal/올더베러_견적서_BAT.xlsx")
print("saved")
