# -*- coding: utf-8 -*-
"""디자인 언어 검증용 샘플 4장."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from template_lib import *

prs = Presentation()
prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]  # blank

# ---- S1: 섹션라벨 + 3 넘버링 서클 ----
s = prs.slides.add_slide(BL)
section_label(s, "전략", "STRATEGY", "우연을 의도로 — 각인 설계의 3단계")
textbox(s, 0.9, 1.55, 11.53, 0.5,
        [("올더베러가 1위가 아닌 이유는 안 팔려서가 아니라 ‘기억에 남지 않아서’ 입니다.",16,INK,False)],
        align=PP_ALIGN.CENTER)
xs=[3.0, 6.665, 10.33]
num_circle(s, xs[0], 2.6, "01", "노출 설계", "마이크로·나노 인플루언서 시딩으로\n다수·반복 노출의 첫 바퀴를 돌린다")
num_circle(s, xs[1], 2.6, "02", "신뢰 각인", "후기·랭킹·1위 신호가 누적되며\n‘많이 선택받는 제품’으로 각인된다")
num_circle(s, xs[2], 2.6, "03", "전환 가속", "확보한 신뢰 증거를 올영세일\n구매 전환으로 연결한다")
hline(s, 3.46, 3.06, 2.74, LGRAY, 1.0)
hline(s, 7.12, 3.06, 2.74, LGRAY, 1.0)
page_num(s, 1)

# ---- S2: 2x2 아이콘 카드 ----
s = prs.slides.add_slide(BL)
section_label(s, "솔루션", "SOLUTION", "캠페인을 굴리는 4개의 엔진")
cards=[("인플루언서 시딩","마이크로·나노 중심의 진성 후기 볼륨으로 첫 바퀴를 돌립니다."),
       ("바이럴 확산","확신·챌린저스·커뮤니티로 신뢰 신호를 증폭합니다."),
       ("검색 자산화","GEO·AEO로 검색 결과를 브랜드 자산으로 축적합니다."),
       ("전환·랭킹","위닝 소재·퍼포먼스 광고로 올영세일 전환을 극대화합니다.")]
pos=[(1.5,1.7),(6.9,1.7),(1.5,4.25),(6.9,4.25)]
for (t,d),(x,y) in zip(cards,pos):
    icon_card(s, x,y,4.9,2.2, t, d)
page_num(s, 2)

# ---- S3: 바차트 ----
s = prs.slides.add_slide(BL)
section_label(s, "진단", "DIAGNOSIS", "나노 인플루언서의 압도적 효율")
textbox(s, 1.5, 1.7, 4.5, 3.5,
        [("왜 마이크로·나노 중심인가",18,INK,True),("",8,INK,False),
         ("나노 인플루언서의 평균 인게이지먼트는 6.23%로 메가 대비 월등하고,",13,GRAY,False),("",6,INK,False),
         ("전환율은 2~3배, 소비자 73%가 ‘진짜 후기’를 선호합니다.",13,GRAY,False)])
vbar_chart(s, 6.6, 2.0, 5.4, 3.4,
           [("나노",0.62),("마이크로",0.41),("미드",0.24),("메가",0.13)])
page_num(s, 3)

# ---- S4: 간지(대각선 분할) ----
s = prs.slides.add_slide(BL)
rect(s, 0,0,13.33,7.5, fill=INK)            # bg(=photo placeholder)
rtri(s, -1, 0, 9.5, 7.5, rot=0, fill=GREEN) # 좌하단 그린 대각선
textbox(s, 0.9, 2.6, 7, 0.4, [("PART 02",16,YELLOW,True,3)])
textbox(s, 0.9, 3.05, 8, 1.4, [("전략 — 우연을 의도로",40,WHITE,True)])
oval(s, 11.0, 0.8, 1.3, fill=WHITE)
textbox(s, 11.0, 0.8, 1.3, 1.3, [("02",26,GREEN,True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
page_num(s, 4)

prs.save("/tmp/sample.pptx")
print("saved /tmp/sample.pptx")
