# -*- coding: utf-8 -*-
"""케어플러스 더마 PDRN 수딩패치 8월 MKT — 나노·마이크로 시딩 믹스 견적서.
올더베러 최종 픽스 견적 단가(VAT별도, 마크업·운영비 일괄 15%)를 고정 기준으로 사용.
  · 나노(IG)      : 정가 275,000 → 할인가 225,000 + 마크업 33,750 = 건당 258,750
  · 마이크로(UGC) : 정가 450,000 → 할인가 350,000 + 마크업 52,500 = 건당 402,500
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NANO_LIST, NANO_DISC, NANO_MK, NANO_C = 275000, 225000, 33750, 258750
MICR_LIST, MICR_DISC, MICR_MK, MICR_C = 450000, 350000, 52500, 402500

DGREEN="1F6B4C"; LGRN="E4F0EA"; DARK="2B2B2B"; GREY="7A7A7A"; WHITE="FFFFFF"
HEADER="34495E"; HFILL="D9E1E8"; ACC="C0392B"; YEL="FFF4CC"
thin=Side(style="thin",color="BFBFBF")
BORD=Border(left=thin,right=thin,top=thin,bottom=thin)

def C(ws,r,c,v,b=False,col=None,fill=None,al="center",sz=10,wrap=False,border=True,it=False):
    cell=ws.cell(row=r,column=c,value=v)
    cell.font=Font(name="맑은 고딕",bold=b,size=sz,color=(col or "222222"),italic=it)
    cell.alignment=Alignment(horizontal=al,vertical="center",wrap_text=wrap)
    if fill: cell.fill=PatternFill("solid",fgColor=fill)
    if border: cell.border=BORD
    return cell
def mrg(ws,r1,c1,r2,c2): ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)

wb=openpyxl.Workbook()

# ============ SHEET 1: 견적 요약(3안) ============
ws=wb.active; ws.title="견적_3안"
widths=[2.5,22,12,10,13,10,13,10,13,15,26]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

r=2
mrg(ws,r,2,r,11); C(ws,r,2,"케어플러스 더마 PDRN 수딩패치 · 8월 인플루언서 시딩 견적 (나노·마이크로 믹스)",b=True,col=WHITE,fill=DGREEN,al="left",sz=13,border=False)
r+=1; mrg(ws,r,2,r,11); C(ws,r,2,"예산 800만 원(±50만) · VAT 별도 · 마크업·운영비 일괄 15% · 숏폼(채널 무관) 바이럴 중심",col=GREY,al="left",sz=9,border=False)
r+=2

# 단가 기준 박스
mrg(ws,r,2,r,11); C(ws,r,2,"■ 고정 단가 기준 (올더베러 최종 픽스 견적 · VAT 별도)",b=True,col=DGREEN,fill=LGRN,al="left")
r+=1
hdr=["구분","채널","정가(원고료)","할인가(A)","할인율","마크업·운영비(B, 15%)","총 견적 건당(C=A+B)"]
cols=[2,3,4,6,8,9,11]
# 헤더 배치 (병합 일부)
C(ws,r,2,"구분",b=True,fill=HFILL); C(ws,r,3,"채널",b=True,fill=HFILL)
C(ws,r,4,"정가(원고료)",b=True,fill=HFILL); mrg(ws,r,4,r,5)
C(ws,r,6,"할인가(A)",b=True,fill=HFILL); C(ws,r,7,"할인율",b=True,fill=HFILL)
C(ws,r,8,"마크업·운영비(B,15%)",b=True,fill=HFILL); mrg(ws,r,8,r,9)
C(ws,r,10,"총 건당(C=A+B)",b=True,fill=HFILL); mrg(ws,r,10,r,11)
r+=1
for nm,ch,li,di,dr,mk,cc in [
    ("나노","IG",NANO_LIST,NANO_DISC,"18.2%",NANO_MK,NANO_C),
    ("마이크로(UGC)","IG·TT",MICR_LIST,MICR_DISC,"22.2%",MICR_MK,MICR_C)]:
    C(ws,r,2,nm,b=True); C(ws,r,3,ch)
    C(ws,r,4,f"{li:,}"); mrg(ws,r,4,r,5)
    C(ws,r,6,f"{di:,}"); C(ws,r,7,dr)
    C(ws,r,8,f"{mk:,}"); mrg(ws,r,8,r,9)
    C(ws,r,10,f"{cc:,}",b=True,col=ACC); mrg(ws,r,10,r,11)
    r+=1
r+=1

# 3안 요약
mrg(ws,r,2,r,11); C(ws,r,2,"■ 예산 배분 3안 (택1) · VAT 별도",b=True,col=DGREEN,fill=LGRN,al="left")
r+=1
for c,h in zip([2,4,6,8,10],["구분","나노(건당 258,750)","마이크로(건당 402,500)","총 인원","합계(VAT별도)"]):
    C(ws,r,c,h,b=True,fill=HFILL)
    if c in (4,6,8,10): mrg(ws,r,c,r,c+1)
mrg(ws,r,2,r,3)
r+=1
scenarios=[
    ("A · 균형형 (추천)",22,6),
    ("B · 도달형 (나노 중심)",26,4),
    ("C · 파워형 (마이크로 중심)",16,10),
]
for nm,n,m in scenarios:
    tot=n*NANO_C+m*MICR_C
    rec = (nm.endswith("(추천)"))
    C(ws,r,2,nm,b=rec,col=(DGREEN if rec else "222222"),fill=(YEL if rec else None),al="left"); mrg(ws,r,2,r,3)
    C(ws,r,4,f"{n}명  ({n*NANO_C:,}원)"); mrg(ws,r,4,r,5)
    C(ws,r,6,f"{m}명  ({m*MICR_C:,}원)"); mrg(ws,r,6,r,7)
    C(ws,r,8,f"{n+m}명",b=True); mrg(ws,r,8,r,9)
    C(ws,r,10,f"{tot:,}원",b=True,col=ACC); mrg(ws,r,10,r,11)
    r+=1
r+=1
mrg(ws,r,2,r,11); C(ws,r,2,"※ VAT(10%) 별도. 예: A안 8,107,500원 + VAT 810,750원 = 8,918,250원.",al="left",sz=8,col=GREY,it=True,border=False)
r+=1; mrg(ws,r,2,r,11); C(ws,r,2,"※ 목표 100인은 상기 고정 단가 기준 유상만으로는 예산 초과. 무가(제품)시딩 병행 시 참여 인원 확대 가능(2번 시트 참고).",al="left",sz=8,col=GREY,it=True,border=False)

# ============ SHEET 2: A안 상세 + A/B ============
ws2=wb.create_sheet("A안_상세_ABtest")
for i,w in enumerate([2.5,20,10,13,13,13,14,15,24],1): ws2.column_dimensions[get_column_letter(i)].width=w
r=2
mrg(ws2,r,2,r,9); C(ws2,r,2,"[추천 A안] 균형형 상세 견적 · 28명 · 8,107,500원 (VAT별도)",b=True,col=WHITE,fill=DGREEN,al="left",sz=12,border=False)
r+=2
for c,h in zip([2,3,4,5,6,7,8,9],["구분","채널","수량","정가","할인가(A)","마크업(B)","건당(C)","소계"]):
    C(ws2,r,c,h,b=True,fill=HFILL)
r+=1
rows=[("나노","IG(Reels)",22,NANO_LIST,NANO_DISC,NANO_MK,NANO_C),
      ("마이크로(UGC)","IG·TT",6,MICR_LIST,MICR_DISC,MICR_MK,MICR_C)]
tot_qty=tot_sum=0
for nm,ch,q,li,di,mk,cc in rows:
    sub=q*cc; tot_qty+=q; tot_sum+=sub
    C(ws2,r,2,nm,b=True,al="left"); C(ws2,r,3,ch); C(ws2,r,4,q)
    C(ws2,r,5,f"{li:,}"); C(ws2,r,6,f"{di:,}"); C(ws2,r,7,f"{mk:,}")
    C(ws2,r,8,f"{cc:,}",b=True); C(ws2,r,9,f"{sub:,}",b=True,col=ACC)
    r+=1
C(ws2,r,2,"합계",b=True,fill=LGRN); C(ws2,r,3,"",fill=LGRN); C(ws2,r,4,tot_qty,b=True,fill=LGRN)
for c in (5,6,7,8): C(ws2,r,c,"",fill=LGRN)
C(ws2,r,8,"",fill=LGRN); C(ws2,r,9,f"{tot_sum:,}",b=True,col=ACC,fill=LGRN)
r+=2

mrg(ws2,r,2,r,9); C(ws2,r,2,"■ 高효율 소재 A/B 테스트 설계 (14 : 14 분할)",b=True,col=DGREEN,fill=LGRN,al="left")
r+=1
C(ws2,r,2,"셀",b=True,fill=HFILL); mrg(ws2,r,2,r,3)
C(ws2,r,4,"USP / 소재 앵글",b=True,fill=HFILL); mrg(ws2,r,4,r,7)
C(ws2,r,8,"구성 / 비중",b=True,fill=HFILL); mrg(ws2,r,8,r,9)
r+=1
ab=[("셀 A · 쿨링/진정","① -6.6℃ 즉각 쿨링 리커버리 — '붙이는 순간 차갑고 촉촉, 관리 후 열감 급속 진정' (시트 앞뒷면 비교 후킹)","나노 11 + 마이크로 3 · 14명(50%)"),
    ("셀 B · 3-PIECE 밀착","② 3-Piece 밀착 설계 — '이마·양볼 굴곡에 딱! 에센스 흘러내림 ZERO' (vs 일반 시트 비교)","나노 11 + 마이크로 3 · 14명(50%)")]
for nm,usp,comp in ab:
    C(ws2,r,2,nm,b=True,al="left"); mrg(ws2,r,2,r,3)
    C(ws2,r,4,usp,al="left",wrap=True); mrg(ws2,r,4,r,7)
    C(ws2,r,8,comp,al="left",wrap=True); mrg(ws2,r,8,r,9)
    ws2.row_dimensions[r].height=30
    r+=1
r+=1
mrg(ws2,r,2,r,9); C(ws2,r,2,"측정지표: 조회수·완주율·저장/공유·댓글 감성 → 위닝 앵글 선정 → 잔여 캠페인·온라인몰 상세 에셋에 집중 재투입",al="left",sz=9,border=False)

wb.save("proposal/케어플러스_8월_시딩견적_나노마이크로믹스.xlsx")
print("saved xlsx")
