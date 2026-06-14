#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""공유 받은 제안 덱(v1.2)에 내가 만든 슬라이드(액션플랜 캘린더·견적서·월별 7장)를 네이티브로 병합"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC="/root/.claude/uploads/bc09908a-8d16-5be3-806b-53899e2cede8/0cae2cb0-BAT______2026________________v1.2.pptx"
OUT="proposal/올더베러_제안덱_병합본.pptx"

BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
FONT="맑은 고딕"

prs=Presentation(SRC)
# pick blank layout
BLANK=None
for l in prs.slide_layouts:
    if l.name=="빈화면": BLANK=l
if BLANK is None: BLANK=prs.slide_layouts[5]
SW,SH=prs.slide_width,prs.slide_height

def cell_style(c,t,size=9,bold=False,color=BLACK,fill=WHITE,align=PP_ALIGN.CENTER,sub=None,subsize=7,subcolor=GRAY):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(4); c.margin_right=Pt(4); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    if sub is not None:
        p2=c.text_frame.add_paragraph(); p2.alignment=align
        r2=p2.add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(subsize); r2.font.color.rgb=subcolor
def title(slide,t,sub=None):
    tb=slide.shapes.add_textbox(Inches(0.5),Inches(0.18),Inches(12.3),Inches(0.7)).text_frame; tb.word_wrap=True
    r=tb.paragraphs[0].add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=BLACK
    if sub:
        r2=tb.add_paragraph().add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(11); r2.font.color.rgb=GRAY
def heat(c,rmax):
    if c==0: return RGBColor(0xFB,0xFB,0xFB),MG
    lv=255-int(168*c/rmax); return RGBColor(lv,lv,lv),(WHITE if lv<135 else BLACK)
def rect(s,x,y,w,h,color):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(x,y,w,h).text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(3)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
def cell(c,t,sz=9.5,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.CENTER):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(5); c.margin_right=Pt(5); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; r=p.add_run(); r.text=t
    r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

# ===== 1) 액션플랜 캘린더 =====
def add_calendar():
    s=prs.slides.add_slide(BLANK)
    title(s,"월별 실행 캘린더 · 액션플랜","8·11·12월 집중 · 항목별 진행 건수(건당 비용) · 잔여 예산은 마이크로/나노 시딩 재배분")
    items=[("매크로","800,000 원",[0,5,0,0,5,5],"15"),("마이크로","400,000 원",[15,35,20,15,30,35],"150"),
     ("나노","250,000 원",[25,55,35,25,40,45],"225"),("KOL 무가시딩","100,000 원",[10,10,10,10,10,10],"60"),
     ("X(트위터) 시딩","500,000 원",[0,5,0,0,5,5],"15"),("__G__","콘텐츠 · 바이럴",None,None),
     ("파워페이지","800,000 원",[0,3,0,0,3,3],"9"),("커뮤니티 체험단","5,000,000 원",[0,1,0,0,1,1],"3"),
     ("네이버 블로그","250,000 원",[0,5,0,0,5,5],"15"),("EGC 콘텐츠","1,500,000 원",[2,2,1,2,1,2],"10"),
     ("__G__","제작",None,None),("KV·디자인·영상","",["●","●","●","●","●","●"],"상시"),
     ("__G__","기타",None,None),("어필리에이트","올영 쇼핑 큐레이터·누적 (명)",[10,20,30,40,50,60],"누적60")]
    months=["7월","8월","9월","10월","11월","12월"]; msub={1:"9월세일 사전부스팅",2:"세일 ★",4:"블프 ★",5:"세일 ★"}
    rowdefs=[("PHASE",),("HEAD",),("G","크리에이터 시딩")]+items+[("SUM",),("BUD",)]
    tbl=s.shapes.add_table(len(rowdefs),8,Inches(0.5),Inches(1.05),Inches(12.33),Inches(6.2)).table
    tbl.columns[0].width=Inches(2.55)
    for i in range(1,7): tbl.columns[i].width=Inches(1.30)
    tbl.columns[7].width=Inches(0.98); ri=0
    for rd in rowdefs:
        k=rd[0]
        if k=="PHASE":
            cell_style(tbl.cell(ri,0),"",fill=WHITE)
            a=tbl.cell(ri,1);b=tbl.cell(ri,3);a.merge(b);cell_style(a,"PHASE 1 · 빠른 반응 확보",10,True,WHITE,BLACK)
            c=tbl.cell(ri,4);d=tbl.cell(ri,6);c.merge(d);cell_style(c,"PHASE 2 · 대표성 확장",10,True,WHITE,GRAY)
            cell_style(tbl.cell(ri,7),"",fill=WHITE)
        elif k=="HEAD":
            cell_style(tbl.cell(ri,0),"항목 (건당 비용)",9.5,True,WHITE,DARK,PP_ALIGN.LEFT)
            for i,m in enumerate(months): cell_style(tbl.cell(ri,i+1),m,11,True,WHITE,DARK,sub=msub.get(i),subsize=7,subcolor=RGBColor(0xDD,0xDD,0xDD))
            cell_style(tbl.cell(ri,7),"합계",9.5,True,WHITE,DARK)
        elif k=="G":
            a=tbl.cell(ri,0);b=tbl.cell(ri,7);a.merge(b);cell_style(a,rd[1],9.5,True,BLACK,LG,PP_ALIGN.LEFT)
        elif k=="SUM":
            cell_style(tbl.cell(ri,0),"월 시딩 총건수",9.5,True,WHITE,BLACK,PP_ALIGN.LEFT)
            for i,v in enumerate([50,110,65,50,90,100]): cell_style(tbl.cell(ri,i+1),str(v),11,True,WHITE,BLACK)
            cell_style(tbl.cell(ri,7),"465건",10,True,WHITE,BLACK)
        elif k=="BUD":
            cell_style(tbl.cell(ri,0),"월 예산 비중 / 금액",9.5,True,BLACK,WHITE,PP_ALIGN.LEFT)
            pct=[10,20,20,10,18,22]; amt=["2,500만","5,000만","5,000만","2,500만","4,500만","5,500만"]
            for i in range(6):
                dk=pct[i]>=20
                cell_style(tbl.cell(ri,i+1),f"{pct[i]}%",12,True,(WHITE if dk else BLACK),(BLACK if dk else LG),sub=amt[i],subsize=8,subcolor=(RGBColor(0xDD,0xDD,0xDD) if dk else GRAY))
            cell_style(tbl.cell(ri,7),"100%",10,True,WHITE,DARK,sub="2.5억",subsize=8,subcolor=RGBColor(0xDD,0xDD,0xDD))
        else:
            name,unit,arr,tot=rd
            if name=="__G__":
                a=tbl.cell(ri,0);b=tbl.cell(ri,7);a.merge(b);cell_style(a,unit,9.5,True,BLACK,LG,PP_ALIGN.LEFT)
            else:
                cell_style(tbl.cell(ri,0),name,9.5,True,BLACK,WHITE,PP_ALIGN.LEFT,sub=(unit or None),subsize=7.5)
                if all(isinstance(x,str) for x in arr):
                    for i,mk in enumerate(arr): cell_style(tbl.cell(ri,i+1),mk,12,False,BLACK,RGBColor(0xF8,0xF8,0xF8))
                else:
                    rmax=max(arr) if max(arr)>0 else 1
                    for i,c in enumerate(arr):
                        fl,fc=heat(c,rmax); cell_style(tbl.cell(ri,i+1),(str(c) if c>0 else "·"),11 if c>0 else 9,c>0,fc,fl)
                cell_style(tbl.cell(ri,7),tot,9.5,True,BLACK,RGBColor(0xEC,0xEC,0xEC))
        ri+=1
    for r in tbl.rows: r.height=Inches(0.34)
    nb=s.shapes.add_textbox(Inches(0.5),Inches(7.28),Inches(12.3),Inches(0.22)).text_frame
    rn=nb.paragraphs[0].add_run(); rn.text="※ 셀 음영이 진할수록 물량 多 · 8·11·12월 집중 운영 · 퍼포먼스 매체는 별도(별첨)"
    rn.font.name=FONT; rn.font.size=Pt(8.5); rn.font.color.rgb=GRAY

# ===== 2) 견적서 =====
def add_quote():
    s=prs.slides.add_slide(BLANK)
    title(s,"견적서 (Quotation)","올리브영 PB 올더베러 브랜드 빌딩 캠페인 · 단위 원(VAT 별도)")
    q=[("G","제작비"),("KV·디자인·영상·콘텐츠 제작 (EGC 제외)","1식","-","67,300,000"),
     ("EGC 콘텐츠","10","1,500,000","15,000,000"),("G","크리에이터 시딩"),
     ("매크로 (인스타·유튜브)","15","800,000","12,000,000"),("마이크로","150","400,000","60,000,000"),
     ("나노","225","250,000","56,250,000"),("KOL 무가시딩","60","100,000","6,000,000"),
     ("X (트위터) 시딩","15","500,000","7,500,000"),("G","콘텐츠 · 바이럴"),
     ("파워페이지 콘텐츠 바이럴","9","800,000","7,200,000"),("커뮤니티 체험단 (탑 3건)","3","5,000,000","15,000,000"),
     ("네이버 블로그","15","250,000","3,750,000"),("G","기타"),
     ("어필리에이트 (올영 쇼핑 큐레이터)","누적 60","-","0"),
     ("T","공급가액 소계","250,000,000"),("T","부가가치세 (10%)","25,000,000"),("TT","합계 금액 (VAT 포함)","275,000,000")]
    t2=s.shapes.add_table(len(q)+1,4,Inches(0.7),Inches(1.05),Inches(11.9),Inches(5.7)).table
    t2.columns[0].width=Inches(6.7); t2.columns[1].width=Inches(1.5); t2.columns[2].width=Inches(1.7); t2.columns[3].width=Inches(2.0)
    cell_style(t2.cell(0,0),"세부 항목",10,True,WHITE,BLACK,PP_ALIGN.LEFT); cell_style(t2.cell(0,1),"수량",10,True,WHITE,BLACK)
    cell_style(t2.cell(0,2),"단가",10,True,WHITE,BLACK); cell_style(t2.cell(0,3),"금액",10,True,WHITE,BLACK,PP_ALIGN.RIGHT)
    ri=1
    for row in q:
        if row[0]=="G":
            a=t2.cell(ri,0);b=t2.cell(ri,3);a.merge(b);cell_style(a,row[1],9.5,True,BLACK,LG,PP_ALIGN.LEFT)
        elif row[0] in ("T","TT"):
            big=row[0]=="TT"; a=t2.cell(ri,0);b=t2.cell(ri,2);a.merge(b)
            cell_style(a,row[1],11 if big else 10,True,(WHITE if big else BLACK),(BLACK if big else WHITE),PP_ALIGN.RIGHT)
            cell_style(t2.cell(ri,3),row[2],13 if big else 10.5,True,(WHITE if big else BLACK),(BLACK if big else WHITE),PP_ALIGN.RIGHT)
        else:
            it,qy,un,am=row
            cell_style(t2.cell(ri,0),it,9.5,False,BLACK,WHITE,PP_ALIGN.LEFT); cell_style(t2.cell(ri,1),qy,9,False,DARK,WHITE)
            cell_style(t2.cell(ri,2),un,9,False,DARK,WHITE); cell_style(t2.cell(ri,3),am,10,am!="0",(BLACK if am!="0" else GRAY),WHITE,PP_ALIGN.RIGHT)
        ri+=1
    for r in t2.rows: r.height=Inches(0.27)
    nb=s.shapes.add_textbox(Inches(0.7),Inches(6.95),Inches(11.9),Inches(0.5)).text_frame; nb.word_wrap=True
    for k,n in enumerate(["· 대행료 항목 제외(나노 시딩 반영) · 파워페이지 건당 800,000원(제작비서 충당) · EGC 건당 1,500,000원 · 공급가액 정확히 2.5억(VAT 별도)",
     "· 퍼포먼스 매체 운영(PA·올영 인앱)은 별도 산정(별첨) · 인플루언서 단가는 인스타·유튜브 기준, X(트위터)는 별도 단가"]):
        p=nb.paragraphs[0] if k==0 else nb.add_paragraph(); r=p.add_run(); r.text=n; r.font.name=FONT; r.font.size=Pt(8.5); r.font.color.rgb=GRAY

# ===== 3) 월별 장표 =====
MONTHS=[7,8,9,10,11,12]
QTY={"매크로":[0,5,0,0,5,5],"마이크로":[15,35,20,15,30,35],"나노":[25,55,35,25,40,45],"KOL 무가시딩":[10,10,10,10,10,10],
 "X(트위터)":[0,5,0,0,5,5],"파워페이지":[0,3,0,0,3,3],"커뮤니티 체험단(탑)":[0,1,0,0,1,1],"네이버 블로그":[0,5,0,0,5,5],
 "EGC 콘텐츠":[2,2,1,2,1,2],"어필리에이트(누적)":[10,20,30,40,50,60]}
QORDER=list(QTY.keys()); SEEDTOTAL=[50,110,65,50,90,100]
BUDpct=[10,20,20,10,18,22]; BUDamt=["25,000,000","50,000,000","50,000,000","25,000,000","45,000,000","55,000,000"]
DATA={
7:("First Shot","착수 · 콘텐츠/메시지 테스트","1차 웨이브 준비","위닝 메시지 발굴 · 시딩 도화선 점화",
   ["콘텐츠·이미지 소재 LMF 테스트 → 위닝 메시지 선별","일반식품 + 심의 불필요 건기식 시딩 즉시 라이브","건기식 심의 파이프라인 착수 (가이드→섭외→콘티→접수)","챌린저스 1차 가동 · 블로그/검색 키워드 선점"],"위닝 메시지 확정 · 후기·콘텐츠 1차 적재"),
8:("Proof Take-off","올영세일 사전 부스팅","★ 1차 웨이브 발생","세일 전 인지·각인 확산 + 선택 증거 확보",
   ["시딩 대량 릴리즈(월 110건) — 세일 1주 전 집중","매크로 합류 · 파워페이지·커뮤니티 체험단(탑)·네이버 블로그 바이럴","건기식 심의 완료분 1차 업로드","‘올영 1위’ 키워드 플레이 · 챌린저스 랭킹 견인"],"1차 웨이브로 사전 부스팅 · 검색·후기 증거 확산"),
9:("Acceleration","올영세일 구매 전환","★ 1차 웨이브 피크","확보한 증거로 올영세일 구매 전환 폭발",
   ["올영세일 집중 — 일반+건기식 콘텐츠 동시 투입","퍼포먼스 매체(별첨)·어필리에이트 가동 전환 유도","건기식 2차 업로드 세일 직전 라이브","랭킹/후기/혜택 결합 소재 운영"],"1차 웨이브 매출 피크 · 카테고리 랭킹 상위/1위"),
10:("Orbit Expansion","유지 · 포트폴리오 확장","유지 (낙수효과)","구매 증거 기반 건기식 라인업 확장",
   ["건기식 6종 본격 콘텐츠 · 확장 메시지 테스트","9월 위닝 소재 낙수 운영(유지)","크로스셀 CRM · 후기 자산 확산"],"매출 유지 · 건기식 라인 인지 형성"),
11:("Boost-up","블프 · 재점화","★ 2차 웨이브 발생","많이 보이고 많이 찾는 브랜드로 재점화",
   ["시딩 재확대(월 90건) · 매크로·체험단(탑)·파워페이지 재가동","블프 프로모션 연계 퍼포먼스 · 건기식 RTB+혜택 결합","챌린저스 2차 · 검색 자산 보강"],"2차 웨이브 시작 · 탐색량·관심층 확대"),
12:("Category Landing","올영세일 · 대표성 착지","★ 2차 웨이브 피크","라인업으로 웰니스 카테고리 대표 브랜드 착지",
   ["12월 올영세일 집중 — 전 라인업 대표성 콘텐츠","세일 1주 전 시딩 대량 릴리즈(월 100건)·바이럴","쟁임/재구매 CRM · 퍼포먼스 2차 피크"],"2차 웨이브 매출 피크 · 대표성 착지")}

def add_month_intro():
    s=prs.slides.add_slide(BLANK)
    rect(s,0,0,SW,Inches(1.4),BLACK)
    txt(s,Inches(0.6),Inches(0.32),Inches(12),Inches(1.0),[("3·4분기 월별 액션플랜",30,True,WHITE),("7월 착수 → 8·9월 1차 웨이브(올영세일) → 10월 유지 → 11·12월 2차 웨이브(블프·올영세일)",14,False,MG)])
    wave=["착수","★1차 웨이브","★1차 웨이브","유지","★2차 웨이브","★2차 웨이브"]; shade=[DARK,BLACK,BLACK,LG,BLACK,BLACK]; tcol=[WHITE,WHITE,WHITE,BLACK,WHITE,WHITE]
    bw=Inches(1.95); gap=Inches(0.12); x0=Inches(0.6); y0=Inches(2.4)
    for i,m in enumerate(MONTHS):
        x=Emu(int(x0)+i*(int(bw)+int(gap))); rect(s,x,y0,bw,Inches(1.5),shade[i])
        txt(s,x,Emu(int(y0)+int(Inches(0.2))),bw,Inches(0.5),[(f"{m}월",22,True,tcol[i])],PP_ALIGN.CENTER)
        txt(s,x,Emu(int(y0)+int(Inches(0.78))),bw,Inches(0.5),[(wave[i],12,True,tcol[i])],PP_ALIGN.CENTER)
        txt(s,x,Emu(int(y0)+int(Inches(1.6))),bw,Inches(0.4),[(f"{BUDpct[i]}%",10,False,GRAY)],PP_ALIGN.CENTER)
    txt(s,Inches(0.6),Inches(5.2),Inches(12),Inches(1.5),[("· Phase 1 (7~9월) — 빠른 반응 확보 : 일반식품·심의불필요 건기식 선행 + 8·9월 1차 웨이브로 올영세일 전환",13,False,DARK),("· Phase 2 (10~12월) — 대표성 확장 : 건기식 라인 확장 + 11·12월 2차 웨이브(블프·올영세일)로 대표 브랜드 착지",13,False,DARK),("· 예산 총 2.5억(공급가) · 퍼포먼스 매체는 별도(별첨)",13,False,GRAY)])

def add_month(idx):
    m=MONTHS[idx]; stage,desc,wavetag,mission,actions,goal=DATA[m]
    s=prs.slides.add_slide(BLANK)
    rect(s,0,0,SW,Inches(1.15),BLACK)
    txt(s,Inches(0.55),Inches(0.16),Inches(7),Inches(0.95),[(f"{m}월  ·  {stage}",26,True,WHITE),(desc,13,False,MG)])
    tagw=Inches(3.1); rect(s,Emu(int(SW)-int(tagw)-int(Inches(0.5))),Inches(0.34),tagw,Inches(0.5),(LG if "유지" in wavetag else WHITE))
    txt(s,Emu(int(SW)-int(tagw)-int(Inches(0.5))),Inches(0.4),tagw,Inches(0.4),[(wavetag,14,True,BLACK)],PP_ALIGN.CENTER)
    rect(s,Inches(0.55),Inches(1.45),Inches(0.08),Inches(0.55),BLACK)
    txt(s,Inches(0.75),Inches(1.45),Inches(11.8),Inches(0.6),[("“"+mission+"”",16,True,DARK)],anchor=MSO_ANCHOR.MIDDLE)
    txt(s,Inches(0.55),Inches(2.35),Inches(7.1),Inches(0.4),[("이 달의 핵심 액션",14,True,BLACK)])
    rect(s,Inches(0.57),Inches(2.78),Inches(1.0),Inches(0.04),BLACK)
    tf=s.shapes.add_textbox(Inches(0.55),Inches(2.95),Inches(7.1),Inches(3.4)).text_frame; tf.word_wrap=True
    for i,a in enumerate(actions):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(9)
        r=p.add_run(); r.text="■  "+a; r.font.name=FONT; r.font.size=Pt(13.5); r.font.color.rgb=DARK
    rect(s,Inches(0.55),Inches(6.35),Inches(7.1),Inches(0.75),LG)
    txt(s,Inches(0.7),Inches(6.45),Inches(6.8),Inches(0.6),[("성과 목표",10,True,GRAY),(goal,13,True,BLACK)])
    txt(s,Inches(7.95),Inches(2.35),Inches(4.8),Inches(0.4),[("이 달의 진행 물량 (건)",14,True,BLACK)])
    rect(s,Inches(7.97),Inches(2.78),Inches(1.0),Inches(0.04),BLACK)
    t=s.shapes.add_table(len(QORDER)+2,2,Inches(7.95),Inches(2.95),Inches(4.85),Inches(3.4)).table
    t.columns[0].width=Inches(3.55); t.columns[1].width=Inches(1.30)
    cell(t.cell(0,0),"항목",9.5,True,WHITE,DARK,PP_ALIGN.LEFT); cell(t.cell(0,1),f"{m}월",9.5,True,WHITE,DARK)
    for ri,name in enumerate(QORDER,1):
        v=QTY[name][idx]
        cell(t.cell(ri,0),name,9,False,BLACK,(WHITE if ri%2 else RGBColor(0xF6,0xF6,0xF6)),PP_ALIGN.LEFT)
        cell(t.cell(ri,1),(str(v) if v>0 else "·"),10,v>0,(BLACK if v>0 else MG),(WHITE if ri%2 else RGBColor(0xF6,0xF6,0xF6)))
    cell(t.cell(len(QORDER)+1,0),"월 시딩 총건수",10,True,WHITE,BLACK,PP_ALIGN.LEFT)
    cell(t.cell(len(QORDER)+1,1),f"{SEEDTOTAL[idx]}",11,True,WHITE,BLACK)
    for r in t.rows: r.height=Inches(0.27)
    rect(s,Inches(7.95),Inches(6.35),Inches(4.85),Inches(0.75),BLACK)
    txt(s,Inches(8.1),Inches(6.43),Inches(4.6),Inches(0.6),[("이 달 예산",10,True,MG),(f"{BUDamt[idx]} 원   ({BUDpct[idx]}%)",16,True,WHITE)])

n_before=len(prs.slides._sldIdLst)
add_calendar(); add_quote(); add_month_intro()
for i in range(6): add_month(i)
n_added=len(prs.slides._sldIdLst)-n_before

# 재배치: '감사' 슬라이드 앞에 삽입 (없으면 끝)
target=None
for i,sl in enumerate(list(prs.slides)[:n_before]):
    for sh in sl.shapes:
        if sh.has_text_frame and "감사" in sh.text_frame.text:
            target=i; break
    if target is not None: break
if target is None: target=n_before
sldIdLst=prs.slides._sldIdLst
ids=list(sldIdLst)
moved=ids[-n_added:]
for el in moved: sldIdLst.remove(el)
for off,el in enumerate(moved): sldIdLst.insert(target+off,el)

prs.save(OUT)
print("merged slides:",len(prs.slides._sldIdLst),"| added:",n_added,"| inserted at:",target)
