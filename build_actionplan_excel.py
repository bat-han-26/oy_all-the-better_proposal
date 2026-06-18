# -*- coding: utf-8 -*-
"""월별 실행 액션플랜 + 월별 견적 비용 (단일 시트 엑셀)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT="맑은 고딕"; ACC='_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'; CNT='#,##0;-;"-"'
DARK="404040"; DGREEN="082A08"; LGRN="EAF3DA"; SUB="D9D9D9"; WHITE="FFFFFF"; INK="262626"; GREY="808080"; GREENF="0E5A30"
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=openpyxl.Workbook(); ws=wb.active; ws.title="월별 액션플랜·견적"
ws.sheet_view.showGridLines=False
for col,w in {1:30,2:14,3:9,4:10,5:10,6:9,7:10,8:10,9:9,10:14}.items():
    ws.column_dimensions[get_column_letter(col)].width=w

def C(r,c,v=None,b=False,col=INK,fill=None,al="center",fmt=None,sz=10,wrap=False,border=True):
    cell=ws.cell(r,c)
    if v is not None: cell.value=v
    cell.font=Font(name=FONT,size=sz,bold=b,color=col)
    if fill: cell.fill=PatternFill("solid",fgColor=fill)
    cell.alignment=Alignment(horizontal=al,vertical="center",wrap_text=wrap)
    if border: cell.border=bd
    if fmt: cell.number_format=fmt
    return cell

# 타이틀
ws.merge_cells("A1:J1")
C(1,1,"월별 실행 타임라인 · 액션플랜 — 견적 포함",b=True,col=DGREEN,al="left",sz=14,border=False)
ws.row_dimensions[1].height=24

# 헤더
HDR=["항목","단가(총 견적)","7월","8월\n세일사전","9월\n세일★","10월","11월\n블프★","12월\n세일★","합계","금액"]
for i,h in enumerate(HDR,1): C(2,i,h,b=True,col=WHITE,fill=DARK,wrap=True)
ws.row_dimensions[2].height=34

ROWS=[
("G","■ 크리에이터 시딩",None,None),
("I","나노",270000,[10,20,12,10,18,20]),
("I","마이크로 (UGC)",420000,[16,30,20,16,28,30]),
("I","매크로",880000,[0,4,0,0,4,4]),
("I","KOL 무가시딩",100000,[5,5,5,5,5,5]),
("I","X (트위터) 시딩",550000,[0,8,0,0,8,8]),
("I","네이버 블로그",220000,[0,5,0,0,5,5]),
("G","■ BATi 앰배서더",None,None),
("I","BATi 마이크로 앰배서더",500000,[10,10,10,10,10,10]),
("S","월 시딩 총건수",None,None),
("G","■ 콘텐츠",None,None),
("I","EGC 콘텐츠",500000,[20,20,20,20,20,20]),
("I","Half EGC 콘텐츠",750000,[4,4,4,4,4,4]),
("I","콘텐츠 제작 (KV·영상·이미지)",1000000,[1,1,1,1,1,1]),
("G","■ 바이럴 · 전환",None,None),
("I","파워페이지",750000,[0,2,0,0,2,2]),
("I","커뮤니티 체험단",4300000,[0,1,0,0,1,0]),
("I","챌린저스",4800000,[0,1,0,0,1,0]),
("G","■ 솔루션",None,None),
("I","스프레이ai 솔루션",23333.3,[1,1,1,1,1,1]),
]
r=3; seed_rows=[]; item_rows=[]; first_item=last_item=None
for kind,label,price,dist in ROWS:
    if kind=="G":
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
        C(r,1,label,b=True,col=DGREEN,fill=LGRN,al="left"); r+=1; continue
    if kind=="S":
        # 월 시딩 총건수 (seeding 합)
        C(r,1,"월 시딩 총건수",b=True,fill=SUB,al="right")
        C(r,2,fill=SUB)
        for mi in range(6):
            cl=get_column_letter(3+mi)
            C(r,3+mi,"="+"+".join(f"{cl}{rr}" for rr in seed_rows),b=True,fill=SUB,fmt=CNT)
        C(r,9,f"=SUM(C{r}:H{r})",b=True,fill=SUB,fmt=CNT); C(r,10,fill=SUB)
        r+=1; continue
    # item row
    C(r,1,label,al="left",sz=9); C(r,2,price,fmt=ACC,sz=9)
    for mi in range(6): C(r,3+mi,dist[mi],fmt=ACC)   # 0은 회계서식으로 "-" 표기
    C(r,9,f"=SUM(C{r}:H{r})",b=True,fmt=CNT)
    C(r,10,f"=I{r}*B{r}",fmt=ACC,b=True)
    item_rows.append(r)
    if first_item is None: first_item=r
    last_item=r
    if label in ("나노","마이크로 (UGC)","매크로","KOL 무가시딩","X (트위터) 시딩","네이버 블로그","BATi 마이크로 앰배서더"):
        seed_rows.append(r)
    r+=1

# ===== 월별 견적 비용 (빨간 박스 위치) =====
C(r,1,"월별 견적 비용 (월 단위)",b=True,col=WHITE,fill=GREENF,al="right")
C(r,2,fill=GREENF)
for mi in range(6):
    cl=get_column_letter(3+mi)
    C(r,3+mi,f"=SUMPRODUCT({cl}{first_item}:{cl}{last_item},$B${first_item}:$B${last_item})",
      b=True,col=WHITE,fill=GREENF,fmt=ACC,sz=9)
C(r,9,"",fill=GREENF)
C(r,10,f"=SUM(C{r}:H{r})",b=True,col=WHITE,fill=GREENF,fmt=ACC)
budget_row=r; r+=1

# 월 예산 비중(%)
C(r,1,"월 예산 비중",b=True,col=DGREEN,fill=LGRN,al="right")
C(r,2,fill=LGRN)
for mi in range(6):
    cl=get_column_letter(3+mi)
    C(r,3+mi,f"={cl}{budget_row}/$J${budget_row}",b=True,col=DGREEN,fill=LGRN,fmt='0.0%',sz=9)
C(r,9,"",fill=LGRN); C(r,10,f"=J{budget_row}/$J${budget_row}",b=True,col=DGREEN,fill=LGRN,fmt='0.0%')
r+=1

# 총 견적
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9)
C(r,1,"총 견적 (VAT 별도)",b=True,col=WHITE,fill=DARK,al="right")
C(r,10,"="+"+".join(f"J{rr}" for rr in item_rows),fmt=ACC,b=True,col=WHITE,fill=DGREEN)

wb.save("proposal/올더베러_월별액션플랜_견적_1p.xlsx")
print("saved. first_item",first_item,"last_item",last_item,"budget_row",budget_row)
