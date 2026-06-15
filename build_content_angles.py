#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""제품·단계·타겟·TPO별 콘텐츠 앵글 — 편집 가능 네이티브 PPT (플라이휠 4단계 정렬)"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

CREAM=RGBColor(0xFC,0xFB,0xF0); BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34)
GRAY=RGBColor(0x78,0x78,0x72); MG=RGBColor(0xB0,0xB0,0xA8); WHITE=RGBColor(0xFF,0xFF,0xFF)
SOFT=RGBColor(0xF1,0xF0,0xE5); LINE=RGBColor(0x22,0x22,0x1F); LGT=RGBColor(0xE6,0xE6,0xDD)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def newslide():
    s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=CREAM; return s
def rrect(s,x,y,w,h,fill,line=None,lw=1.5,rad=0.1,rounded=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded:
        try: sp.adjustments[0]=rad
        except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def tbox(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(3); tf.margin_right=Pt(3); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(1)
        for (t,sz,b,c) in runs:
            r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def head(s,tag,t,sub):
    tbox(s,0.5,0.3,11.5,0.28,[[(tag,12.5,True,GRAY)]])
    tbox(s,0.5,0.58,12.3,0.6,[[(t,26,True,BLACK)]])
    tbox(s,0.5,1.2,12.3,0.34,[[(sub,13,True,DARK)]])
    rrect(s,0.5,1.6,12.33,0.03,BLACK,rounded=False)

s=newslide()
head(s,"02_B 시딩 콘텐츠 앵글","제품·단계·타겟·TPO별 콘텐츠 앵글",
     "VOC가 알려준 ‘소비자가 이미 하는 말’을 훅으로 — 플라이휠 4단계로 굴립니다.")
# 성공 훅 공식 배너
rrect(s,0.5,1.78,12.33,0.5,BLACK,rad=0.1)
tbox(s,0.5,1.78,12.33,0.5,[[("성공 훅 = 명확한 타겟  ×  구체적 상황·에피소드  ×  반전 또는 팩트",15,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
tbox(s,0.5,2.32,12.33,0.26,[[("특정 생활자를 겨냥 → 제품이 필요한 순간을 장면화 → ‘보게 만드는’ 정보·공감·반전 포인트",10.5,True,GRAY)]],PP_ALIGN.CENTER)

# 4 단계 컬럼 (플라이휠 정렬)
cols=[
 ("① 인지·발견","SEED · 첫 바퀴",[
   ("멜라나잇 구미 · 일반식품","밤에 쉽게 못 드는 30대 · 취침 전","“불 끄기 전, 나만의 마무리 한 알”",
    "반전 — 수면제 아닌 ‘무드 리추얼’ · 심의: 수면·숙면 표현 금지"),
   ("웰니스 구미 (에너지) · 일반식품","야근 잦은 직장인 · 오후 슬럼프","“11시 퇴근하는 30대, 오후 3시 버티는 법”",
    "공감 — 직장인 하루 루틴 안착 · TPO: 야근·오후 3시"),
 ]),
 ("② 검증·신뢰","PROVE · 관성",[
   ("올리브오일 캡슐 · 일반식품","공복 섭취가 고민인 입문자 · 아침 공복","“공복에 올리브오일? 괜찮을까 — 한 포로 끝”",
    "답변형 — VOC ‘섭취 타이밍’ 질문에 답 · 편의 55%"),
   ("올리브오일 캡슐 · 일반식품","가성비 따지는 30대 · 정보 탐색","“생오일 한 병 vs 캡슐 한 포, 가성비 팩트체크”",
    "팩트 — VOC 캡슐 가성비·충성(재구매 24.1%)"),
 ]),
 ("③ 전환","BOOST · 가속",[
   ("올리브오일 스틱 · 일반식품","맛 때문에 망설인 사람 · 재구매 고민","“솔직히 오일 맛 나요. 근데 왜 또 살까?”",
    "역설 — VOC 맛 불호 69.8% 선제 인정 → 신뢰 전환"),
   ("히어로 라인업 · 올영세일","장바구니 담아둔 세일 대기족 · 올영세일","“담아둔 거, 이번 올영세일에 사야 하는 이유”",
    "세일 집중 — 위닝 콘텐츠 PA 2차 연계 집행"),
 ]),
 ("④ 재구매·확장","RETAIN · 자가가속",[
   ("구미·캡슐 · 완주 인증","작심삼일러 · 한 통 완주 후","“안 먹고 쌓인 통 vs 싹 비운 통”",
    "비교·완주 인증 — VOC ‘효과 미체감’을 꾸준함으로"),
   ("올리브오일 스틱 · 휴대","K-웰니스 관심 외국인 · K-뷰티 트렌드","“K-뷰티 다음은 K-웰니스 — 가방 속 한 포”",
    "확장 — K-뷰티 언박싱 포맷 차용 · 휴대 앵글"),
 ]),
]
cw=2.96; gap=0.12; x0=0.5; hy=2.66; hh=0.62
ctop=3.36; cbot=6.6; cgap=0.12
ch=(cbot-ctop-cgap)/2
nstrip=0.62
for ci,(stg,en,cards) in enumerate(cols):
    x=x0+ci*(cw+gap)
    rrect(s,x,hy,cw,hh,BLACK,rad=0.12)
    tbox(s,x+0.16,hy+0.07,cw-0.3,0.3,[[(stg,13,True,WHITE)]])
    tbox(s,x+0.16,hy+0.35,cw-0.3,0.24,[[(en,9.5,True,MG)]])
    for j,(prod,tgt,hook,note) in enumerate(cards):
        cy=ctop+j*(ch+cgap)
        rrect(s,x,cy,cw,ch,WHITE,LINE,1.5,rad=0.06)
        tbox(s,x+0.16,cy+0.1,cw-0.3,0.24,[[(prod,8.8,True,GRAY)]])
        tbox(s,x+0.16,cy+0.34,cw-0.3,0.5,[[(tgt,10,True,DARK)]])
        tbox(s,x+0.16,cy+0.8,cw-0.3,0.62,[[(hook,11,True,BLACK)]])
        rrect(s,x,cy+ch-nstrip,cw,nstrip,SOFT,rounded=False)
        rrect(s,x,cy+ch-nstrip,0.06,nstrip,BLACK,rounded=False)
        tbox(s,x+0.16,cy+ch-nstrip+0.05,cw-0.28,nstrip-0.08,[[(note,8.6,True,DARK)]])

# bottom
rrect(s,0.5,6.74,12.33,0.42,SOFT,LINE,1.2,rad=0.2)
tbox(s,0.5,6.74,12.33,0.42,[[("앵글은 모두 VOC 1,104건 딥다이브에서 도출 · 식품 분류별 심의 가드 + 풀 대본·자막 타임코드는 부록 참조",10,True,DARK)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

prs.save("proposal/올더베러_콘텐츠앵글_TPO.pptx")
print("slides:",len(prs.slides._sldIdLst))
