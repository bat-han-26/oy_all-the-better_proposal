#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""병합 덱에 VOC 기반 시딩 콘텐츠 앵글 슬라이드 추가 (시딩 섹션에 삽입)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DECK="proposal/올더베러_제안덱_병합본.pptx"
BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
FONT="맑은 고딕"
prs=Presentation(DECK)
BLANK=None
for l in prs.slide_layouts:
    if l.name=="빈화면": BLANK=l
if BLANK is None: BLANK=prs.slide_layouts[5]
SW,SH=prs.slide_width,prs.slide_height

def rect(s,x,y,w,h,c):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=c; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(x,y,w,h).text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(3)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def title(s,t,sub=None,tag=None):
    rect(s,0,0,SW,Inches(0.14),BLACK)
    tb=s.shapes.add_textbox(Inches(0.5),Inches(0.34),Inches(9.5),Inches(0.95)).text_frame; tb.word_wrap=True
    r=tb.paragraphs[0].add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(23); r.font.bold=True; r.font.color.rgb=BLACK
    if sub:
        r2=tb.add_paragraph().add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(12); r2.font.color.rgb=GRAY
    rect(s,Inches(0.52),Inches(1.4),Inches(1.2),Inches(0.05),BLACK)
    if tag:
        tw=Inches(2.7); rect(s,Emu(int(SW)-int(tw)-int(Inches(0.5))),Inches(0.4),tw,Inches(0.46),LG)
        txt(s,Emu(int(SW)-int(tw)-int(Inches(0.5))),Inches(0.46),tw,Inches(0.36),[(tag,12,True,BLACK)],PP_ALIGN.CENTER)
def cell(c,t,sz=9.5,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.CENTER):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(5); c.margin_right=Pt(5); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; p.word_wrap=True
    r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

def storyboard(s,x,ytop,w,badge,headline,length,items,guard,star=False):
    # heading bar
    rect(s,x,ytop,w,Inches(0.62),BLACK)
    txt(s,Emu(int(x)+int(Inches(0.12))),Emu(int(ytop)+int(Inches(0.06))),Emu(int(w)-int(Inches(0.24))),Inches(0.5),
        [(headline,12.5,True,WHITE),(f"{badge} · {length}",9.5,False,MG)])
    ty=Emu(int(ytop)+int(Inches(0.66)))
    t=s.shapes.add_table(len(items),2,x,ty,w,Inches(0.34*len(items))).table
    t.columns[0].width=Inches(0.95); t.columns[1].width=Emu(int(w)-int(Inches(0.95)))
    for ri,(tm,desc) in enumerate(items):
        cell(t.cell(ri,0),tm,8.5,True,WHITE,DARK)
        cell(t.cell(ri,1),desc,9,False,BLACK,(WHITE if ri%2==0 else RGBColor(0xF5,0xF5,0xF5)),PP_ALIGN.LEFT)
    for r in t.rows: r.height=Inches(0.34)
    gy=Emu(int(ty)+int(Inches(0.34*len(items)))+int(Inches(0.06)))
    rect(s,x,gy,w,Inches(0.5),LG)
    txt(s,Emu(int(x)+int(Inches(0.1))),Emu(int(gy)+int(Inches(0.04))),Emu(int(w)-int(Inches(0.2))),Inches(0.44),
        [(("★ " if star else "")+"심의: "+guard,9,True,BLACK)])

NEW=[]
def newslide():
    s=prs.slides.add_slide(BLANK); NEW.append(s); return s

# ---- S1. 설계 원칙 + VOC 3트랙 ----
s=newslide()
title(s,"시딩 콘텐츠 앵글 — VOC 기반 설계 원칙","리뷰 1,104건 딥다이브 · ‘소비자가 이미 하는 말’을 후킹으로 · 식품 분류별 심의 필터")
txt(s,Inches(0.5),Inches(1.65),Inches(7.1),Inches(0.4),[("설계 원칙",14,True,BLACK)])
txt(s,Inches(0.5),Inches(2.1),Inches(7.1),Inches(3.0),[
 ("① VOC 빈도 = 공감 확률 — 리뷰에 자주 나온 말을 그대로 훅으로 사용",13,False,DARK),
 ("② 식품 분류별 표현 필터링 우선",13,False,DARK),
 ("    · 일반식품(올리브오일·멜라나잇·콜라겐 구미): 건강·기능 효능 표현 일체 불가",11.5,False,GRAY),
 ("    · 건강기능식품 구미(루테인·바나바잎 등): 인정받은 기능성 문구 안에서만",11.5,False,GRAY),
 ("③ 공통 약점 방어 — ‘효과 미체감’·‘가격/할인 의존’을 루틴·기대치 관리로 보완",13,False,DARK),
])
txt(s,Inches(7.9),Inches(1.65),Inches(4.9),Inches(0.4),[("VOC 진단 — 3트랙 전략",14,True,BLACK)])
t=s.shapes.add_table(4,3,Inches(7.9),Inches(2.1),Inches(4.95),Inches(2.3)).table
t.columns[0].width=Inches(1.95); t.columns[1].width=Inches(1.5); t.columns[2].width=Inches(1.5)
for c,h in enumerate(["제품(전략)","재구매율","핵심"]): cell(t.cell(0,c),h,9.5,True,WHITE,DARK)
for ri,row in enumerate([["올리브오일 캡슐 (캐시카우)","24.1%","편의 55% / 목넘김"],
 ["웰니스 구미 (이탈 방어)","9.6%","맛 50.7% / 호불호"],
 ["올리브오일 스틱 (맛 개선)","16.2%","별점 4.65 / 불호 70%"]],1):
    for c,v in enumerate(row): cell(t.cell(ri,c),v,9,c==0,BLACK,(WHITE if ri%2 else RGBColor(0xF5,0xF5,0xF5)),PP_ALIGN.LEFT if c==0 else PP_ALIGN.CENTER)
for r in t.rows: r.height=Inches(0.55)
txt(s,Inches(7.9),Inches(4.55),Inches(4.95),Inches(1.2),[
 ("✅ 살릴 점: ‘먹기 편한 건기식’·캡슐 가성비/충성·데일리 루틴",11,True,BLACK),
 ("⚠️ 막을 점: 효과 미체감(섭취 타이밍 가이드)·맛 호불호(소용량/샘플)·가격 의존",11,False,GRAY),
])

# ---- S2. 8 컨셉 한눈에 ----
s=newslide()
title(s,"제품별 릴스 컨셉 8종 — 한눈에","성공 훅 = 명확한 타겟 × 구체적 상황 × 반전/팩트")
rows=[
 ["캡슐 A","일반식품","생올리브오일 챌린지 → 캡슐 반전","“몸에 좋다길래 생으로 먹어봤다 울었음 😵”"],
 ["캡슐 B","일반식품","‘알약 잘 못 먹는 사람용’ 목넘김 실연","“알약만 먹으면 목에 걸리는 사람… 저요 🙋”"],
 ["구미·루테인","건기식","모니터 앞 직장인의 눈 챙기기","“내 눈 24시간 근무 중인데…”"],
 ["구미·바나바잎","건기식","식후 디저트처럼","“밥 먹고 나면 꼭 단 게 당기죠?”"],
 ["구미·멜라나잇","일반식품","자기 전 무드 리추얼","“불 끄기 전, 나만의 마무리 한 알”"],
 ["구미·콜라겐","일반식품","디저트 대신 한 알","“단 거 당기는데 죄책감은 싫고”"],
 ["스틱 A","일반식품","가방에 쏙, 출근길 한 포","“요즘 가방에 꼭 챙기는 거 하나”"],
 ["스틱 B","일반식품","오일 맛 솔직 리뷰","“솔직히 오일 맛 나요, 근데…”"],
]
t=s.shapes.add_table(len(rows)+1,4,Inches(0.5),Inches(1.6),Inches(12.35),Inches(5.4)).table
t.columns[0].width=Inches(1.7); t.columns[1].width=Inches(1.3); t.columns[2].width=Inches(4.0); t.columns[3].width=Inches(5.35)
for c,h in enumerate(["제품","분류","컨셉 / 훅","인트로 예시 문구"]): cell(t.cell(0,c),h,10,True,WHITE,DARK,PP_ALIGN.LEFT if c in(2,3) else PP_ALIGN.CENTER)
for ri,row in enumerate(rows,1):
    fill=WHITE if ri%2 else RGBColor(0xF5,0xF5,0xF5)
    cell(t.cell(ri,0),row[0],10,True,BLACK,fill)
    cell(t.cell(ri,1),row[1],9,False,(GRAY if row[1]=="일반식품" else BLACK),fill)
    cell(t.cell(ri,2),row[2],9.5,False,BLACK,fill,PP_ALIGN.LEFT)
    cell(t.cell(ri,3),row[3],9.5,False,DARK,fill,PP_ALIGN.LEFT)
for r in t.rows: r.height=Inches(0.52)
txt(s,Inches(0.5),Inches(7.08),Inches(12.3),Inches(0.3),[("※ VOC 키워드를 후킹으로 사용 · 식품 분류별 심의 가드 + 풀 대본/자막 타임코드 보유(부록)",9,False,GRAY)])

# ---- S3. 캡슐 ----
s=newslide()
title(s,"릴스 컨셉 상세 — 올리브오일 캡슐","일반식품 · 효능 표현 금지(원료·맛·편의만) · 캐시카우 충성 강화",tag="일반식품")
storyboard(s,Inches(0.5),Inches(1.6),Inches(6.0),"일반식품","컨셉 A. 생올리브오일 챌린지 → 캡슐 반전","15초·Before/After",
 [("0–2s","생올리브오일 한 입 → ‘윽’ 리액션 (훅)"),("2–6s","자막 ‘몸에 좋다길래 생으로… 너무 매워 😵’"),
  ("6–11s","‘이걸 캡슐로?’ 한 포 물과 함께 슉, 표정 반전"),("11–15s","아침 식탁 루틴 + 제품 클로즈업")],
 "효능 금지 · ‘엑스트라 버진 올리브오일을 한 포 캡슐로’ 원료·맛·편의만")
storyboard(s,Inches(6.83),Inches(1.6),Inches(6.0),"일반식품","컨셉 B. ‘알약 잘 못 먹는 사람용’ 목넘김 실연","12초·공감",
 [("0–3s","‘알약만 먹으면 목에 걸리는 사람… 저요 🙋’ (VOC 29%)"),("3–8s","캡슐 표면 클로즈업 → 물과 함께 ASMR ‘매끄러워서 슉’"),
  ("8–12s","개별포장 ‘들고 다니기도 편함’")],
 "단점1위 목넘김 32.6% 솔직 인정 → ‘매끄럽다’ 실연(과장 ❌)")

# ---- S4. 구미 건기식 ----
s=newslide()
title(s,"릴스 컨셉 상세 — 웰니스 구미 (건강기능식품)","인정 기능성 문구 ‘그대로’만 · 효과 암시·체험 단정 금지",tag="건기식")
storyboard(s,Inches(0.5),Inches(1.6),Inches(6.0),"건기식·루테인","컨셉 A. 모니터 앞 직장인의 눈 챙기기","15초",
 [("0–3s","모니터·폰 일상 타임랩스 ‘내 눈 24시간 근무 중 👀’"),("3–8s","책상 위 구미 톡, 맛있게 ‘맛은 거의 젤리’"),
  ("8–15s","성분 자막 + 유료광고 표시")],
 "‘루테인은 눈 건강에 도움을 줄 수 있어요’ 인정문구 그대로 · ‘눈 좋아진다’ ❌")
storyboard(s,Inches(6.83),Inches(1.6),Inches(6.0),"건기식·바나바잎","컨셉 B. 식후 디저트처럼","15초",
 [("0–3s","식사 마무리 ‘밥 먹고 디저트 대신 이거’"),("3–9s","구미 ASMR + 인정 기능성 자막"),
  ("9–15s","‘식후 루틴’ + 표시광고")],
 "‘식후 혈당상승 억제에 도움’ · ‘혈당 낮춤/당뇨’ 질병·치료 표방 ❌")

# ---- S5. 구미 일반식품 ----
s=newslide()
title(s,"릴스 컨셉 상세 — 웰니스 구미 (일반식품)","효능 표현 전면 배제 · 무드·맛·간식 대체 앵글만",tag="일반식품")
storyboard(s,Inches(0.5),Inches(1.6),Inches(6.0),"일반식품·멜라나잇","컨셉 C. 자기 전 무드 리추얼","12초",
 [("0–3s","조명 끄고 침대 정돈 (무드 훅)"),("3–8s","젤리 입에 쏙 ‘오늘 하루 마무리’"),
  ("8–12s","불 끄는 컷으로 마무리")],
 "수면·숙면·잠 표현 전면 금지 — ‘자기 전 리추얼·달달한 마무리’ 무드/맛만",star=True)
storyboard(s,Inches(6.83),Inches(1.6),Inches(6.0),"일반식품·콜라겐","컨셉 D. 디저트 대신 한 알","12초",
 [("0–3s","달달한 거 당기는 순간 ‘디저트 대신 이거’"),("3–8s","맛·식감 ASMR ‘간식처럼 챙겨 먹는 젤리’"),
  ("8–12s","가방·파우치에서 꺼내는 휴대 컷")],
 "피부·주름·탄력 효능 금지 — 맛·간식 대체·휴대 앵글만")

# ---- S6. 스틱 ----
s=newslide()
title(s,"릴스 컨셉 상세 — 올리브오일 스틱","일반식품 · 휴대·원료·편의 · 맛 불호 69.8% 선제 방어",tag="일반식품")
storyboard(s,Inches(0.5),Inches(1.6),Inches(6.0),"일반식품","컨셉 A. 가방에 쏙, 출근길 한 포","12초·휴대",
 [("0–3s","파우치에서 스틱 발견 ‘가방에 꼭 챙기는 것 👜’"),("3–8s","톡 따서 한 포 ‘한 포 휴대 간편’"),
  ("8–12s","‘엑스트라 버진 올리브오일을 이렇게’ 원료 자막")],
 "효능 금지 · 원료명·휴대·편의만")
storyboard(s,Inches(6.83),Inches(1.6),Inches(6.0),"일반식품","컨셉 B. 오일 맛 솔직 리뷰","15초·내돈내산",
 [("0–4s","‘솔직히 오일 맛 나요. 근데…’ 선제 고지(훅)"),("4–10s","차갑게/물과 함께 먹는 섭취 팁 실연"),
  ("10–15s","‘이 맛 OK면 휴대 진짜 편함’")],
 "맛 약점 먼저 인정 → 신뢰 확보 → 섭취 팁 전환")

# ---- S7. 심의 가드 & 표현 치환표 ----
s=newslide()
title(s,"심의 가드 & 표현 치환표 (❌ → ✅)","식품 분류별 표현 규칙 · 위험 표현을 안전 표현으로")
t=s.shapes.add_table(6,3,Inches(0.5),Inches(1.65),Inches(12.35),Inches(3.4)).table
t.columns[0].width=Inches(3.6); t.columns[1].width=Inches(4.3); t.columns[2].width=Inches(4.45)
for c,h in enumerate(["제품","❌ 위험 표현","✅ 안전 표현"]): cell(t.cell(0,c),h,10,True,WHITE,DARK,PP_ALIGN.LEFT)
sub=[["올리브오일 캡슐·스틱(일반)","혈관에 좋은 올리브오일","엑스트라 버진 올리브오일을 한 포로"],
 ["구미·멜라나잇(일반)","숙면 젤리 / 잠 잘 오는","자기 전 리추얼 / 하루 마무리 젤리"],
 ["구미·콜라겐(일반)","피부·주름 콜라겐","디저트 대신 챙기는 젤리"],
 ["구미·루테인(건기식)","눈이 좋아지는 / 덜 침침","루테인, 눈 건강에 도움을 줄 수 있음"],
 ["구미·바나바잎(건기식)","혈당 낮춰주는 / 당뇨에","식후 혈당상승 억제에 도움을 줄 수 있음"]]
for ri,row in enumerate(sub,1):
    fill=WHITE if ri%2 else RGBColor(0xF5,0xF5,0xF5)
    cell(t.cell(ri,0),row[0],9.5,True,BLACK,fill,PP_ALIGN.LEFT)
    cell(t.cell(ri,1),row[1],9.5,False,GRAY,fill,PP_ALIGN.LEFT)
    cell(t.cell(ri,2),row[2],9.5,True,BLACK,fill,PP_ALIGN.LEFT)
for r in t.rows: r.height=Inches(0.55)
txt(s,Inches(0.5),Inches(5.3),Inches(12.3),Inches(1.6),[
 ("· 일반식품(올리브오일·멜라나잇·콜라겐): 질병 예방·치료, 신체기능·건강 효능 표현 전면 금지",12,True,BLACK),
 ("· 건기식 구미(루테인·바나바잎): 인정 기능성 문구 그대로만, 의약품 오인·과장·체험 단정 금지",12,True,BLACK),
 ("· ‘효과 미체감’ 방어: ‘꾸준히’ 표현 OK, 효과 보장·단정은 금지",12,False,DARK),
])

# ---- S8. 공통 체크리스트 ----
s=newslide()
title(s,"릴스 공통 심의 · 표시광고 체크리스트 (인스타그램)","송출 전 필수 점검 · 건기식 컨셉은 자율심의+법무/RA 검토")
txt(s,Inches(0.6),Inches(1.8),Inches(12),Inches(4.5),[
 ("☑  유료광고/협찬 표시 — 자막+음성+캡션 모두 ‘광고·유료광고·협찬’ 명시 + 인스타 ‘유료 파트너십’ 태그 (공정위 추천·보증 심사지침)",13,True,BLACK),
 ("☑  일반식품(올리브오일·멜라나잇·콜라겐) — 질병 예방·치료, 신체기능·건강 효능 표현 전면 금지",13,True,BLACK),
 ("☑  건기식 구미(루테인·바나바잎) — 인정 기능성 문구 그대로만, 의약품 오인·과장·체험 단정 금지 (한국건강기능식품협회 자율심의 권고)",13,True,BLACK),
 ("☑  ‘효과 미체감’ 방어 — ‘꾸준히’ 표현은 OK, 단 효과 보장·단정은 금지",13,True,BLACK),
 ("☑  건강기능식품 컨셉(루테인·바나바잎)은 송출 전 한국건강기능식품협회 자율심의 및 법무·RA 검토 완료",13,False,DARK),
 ("",6,False,GRAY),
 ("→ AI 심의 가이드 검증 봇으로 사전 필터링(반려↓·기간 단축) 후, 위 체크리스트 최종 점검",12,True,GRAY),
])

# ---- 삽입 위치: '콘텐츠 앵글' 또는 'Seeding' 슬라이드 뒤 ----
n_before=len(prs.slides._sldIdLst)-len(NEW)
target=n_before
for i,sl in enumerate(list(prs.slides)[:n_before]):
    txtall=" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    if ("콘텐츠 앵글" in txtall) or ("Seeding Content" in txtall and "제품별" in txtall):
        target=i+1
for i,sl in enumerate(list(prs.slides)[:n_before]):
    txtall=" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    if "콘텐츠 앵글" in txtall: target=i+1; break

sldIdLst=prs.slides._sldIdLst; ids=list(sldIdLst); moved=ids[-len(NEW):]
for el in moved: sldIdLst.remove(el)
for off,el in enumerate(moved): sldIdLst.insert(target+off,el)
prs.save(DECK)
print("total:",len(prs.slides._sldIdLst),"| added:",len(NEW),"| inserted after idx",target-1)
