#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""올더베러 액션플랜 + 견적서 — 편집 가능한 네이티브 PPT (표/텍스트)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def cell_style(cell,text,size=9,bold=False,color=BLACK,fill=WHITE,align=PP_ALIGN.CENTER,sub=None,subsize=7,subcolor=GRAY):
    cell.fill.solid(); cell.fill.fore_color.rgb=fill
    cell.margin_left=Pt(4); cell.margin_right=Pt(4); cell.margin_top=Pt(1); cell.margin_bottom=Pt(1)
    cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf=cell.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    if sub is not None:
        p2=tf.add_paragraph(); p2.alignment=align
        r2=p2.add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(subsize); r2.font.color.rgb=subcolor

def title(slide,t,sub=None):
    tb=slide.shapes.add_textbox(Inches(0.5),Inches(0.18),Inches(12.3),Inches(0.7)).text_frame
    tb.word_wrap=True; r=tb.paragraphs[0].add_run(); r.text=t
    r.font.name=FONT; r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=BLACK
    if sub:
        r2=tb.add_paragraph().add_run(); r2.text=sub
        r2.font.name=FONT; r2.font.size=Pt(11); r2.font.color.rgb=GRAY

def heat(c,rmax):
    if c==0: return RGBColor(0xFB,0xFB,0xFB), MG
    lv=255-int(168*c/rmax); return RGBColor(lv,lv,lv), (WHITE if lv<135 else BLACK)

# ============================================================
# SLIDE 1 — 타임라인 & 액션플랜 (월별 실행 캘린더)
# ============================================================
s=prs.slides.add_slide(BLANK)
title(s,"월별 실행 캘린더 · 액션플랜","8·11·12월 집중 · 항목별 진행 건수(건당 비용) · 잔여 예산은 마이크로/나노 시딩 재배분")
# data
items=[
 ("매크로","인스타·유튜브 80만",[0,5,0,0,5,5],"15"),
 ("마이크로","40만",[18,35,20,17,28,32],"150"),
 ("나노","25만",[22,42,25,20,30,36],"175"),
 ("KOL 무가시딩","10만 상당·자체 0",[10,10,10,10,10,10],"60"),
 ("X(트위터) 시딩","50만",[0,5,0,0,5,5],"15"),
 ("__G__","콘텐츠 · 바이럴",None,None),
 ("파워페이지","60만~",[0,3,0,0,3,3],"9"),
 ("커뮤니티 체험단","탑500·로200만",[0,1,0,0,1,1],"3"),
 ("네이버 블로그","25만~",[0,5,0,0,5,5],"15"),
 ("EGC 콘텐츠","자체 제작(편)",[2,2,1,2,1,2],"10"),
 ("__G__","제작",None,None),
 ("KV·디자인·영상","제작(●=진행)",["●","●","●","●","●","●"],"상시"),
 ("__G__","기타 (자체 비용 0)",None,None),
 ("어필리에이트","올영 쇼핑 큐레이터·누적",[10,20,30,40,50,60],"누적60"),
]
months=["7월","8월","9월","10월","11월","12월"]
msub={1:"9월세일 사전부스팅",2:"세일 ★",4:"블프 ★",5:"세일 ★"}
# rows: phase + header + 시딩group + 5 + items(14 incl groups) ... build row list
rowdefs=[("PHASE",),("HEAD",),("G","크리에이터 시딩")]+items+[("SUM",),("BUD",)]
nrows=len(rowdefs); ncols=8
tbl=s.shapes.add_table(nrows,ncols,Inches(0.5),Inches(1.05),Inches(12.33),Inches(6.2)).table
tbl.columns[0].width=Inches(2.55)
for i in range(1,7): tbl.columns[i].width=Inches(1.30)
tbl.columns[7].width=Inches(0.98)
# turn off banded style influence by filling every cell
ri=0
for rd in rowdefs:
    kind=rd[0]
    if kind=="PHASE":
        cell_style(tbl.cell(ri,0),"",fill=WHITE)
        a=tbl.cell(ri,1); b=tbl.cell(ri,3); a.merge(b); cell_style(a,"PHASE 1 · 빠른 반응 확보",10,True,WHITE,BLACK)
        c=tbl.cell(ri,4); d=tbl.cell(ri,6); c.merge(d); cell_style(c,"PHASE 2 · 대표성 확장",10,True,WHITE,GRAY)
        cell_style(tbl.cell(ri,7),"",fill=WHITE)
    elif kind=="HEAD":
        cell_style(tbl.cell(ri,0),"항목 (건당 비용)",9.5,True,WHITE,DARK,PP_ALIGN.LEFT)
        for i,m in enumerate(months):
            cell_style(tbl.cell(ri,i+1),m,11,True,WHITE,DARK,sub=msub.get(i),subsize=7,subcolor=RGBColor(0xDD,0xDD,0xDD))
        cell_style(tbl.cell(ri,7),"합계",9.5,True,WHITE,DARK)
    elif kind=="G":
        a=tbl.cell(ri,0); b=tbl.cell(ri,7); a.merge(b); cell_style(a,rd[1],9.5,True,BLACK,LG,PP_ALIGN.LEFT)
    elif kind=="SUM":
        cell_style(tbl.cell(ri,0),"월 시딩 총건수",9.5,True,WHITE,BLACK,PP_ALIGN.LEFT)
        for i,v in enumerate([50,97,55,47,78,88]): cell_style(tbl.cell(ri,i+1),str(v),11,True,WHITE,BLACK)
        cell_style(tbl.cell(ri,7),"415건",10,True,WHITE,BLACK)
    elif kind=="BUD":
        cell_style(tbl.cell(ri,0),"월 예산 비중 / 금액",9.5,True,BLACK,WHITE,PP_ALIGN.LEFT)
        pct=[10,20,20,10,18,22]; amt=["2,500만","5,000만","5,000만","2,500만","4,500만","5,500만"]
        for i in range(6):
            dark=pct[i]>=20
            cell_style(tbl.cell(ri,i+1),f"{pct[i]}%",12,True,(WHITE if dark else BLACK),(BLACK if dark else LG),sub=amt[i],subsize=8,subcolor=(RGBColor(0xDD,0xDD,0xDD) if dark else GRAY))
        cell_style(tbl.cell(ri,7),"100%",10,True,WHITE,DARK,sub="2.5억",subsize=8,subcolor=RGBColor(0xDD,0xDD,0xDD))
    else:
        name,unit,arr,tot=rd
        if name=="__G__":
            a=tbl.cell(ri,0); b=tbl.cell(ri,7); a.merge(b); cell_style(a,unit,9.5,True,BLACK,LG,PP_ALIGN.LEFT)
        else:
            cell_style(tbl.cell(ri,0),name,9.5,True,BLACK,WHITE,PP_ALIGN.LEFT,sub=unit,subsize=7.5)
            if all(isinstance(x,str) for x in arr):  # marker row
                for i,mk in enumerate(arr): cell_style(tbl.cell(ri,i+1),mk,12,False,BLACK,RGBColor(0xF8,0xF8,0xF8))
            else:
                rmax=max(arr) if max(arr)>0 else 1
                for i,c in enumerate(arr):
                    fill,fc=heat(c,rmax); cell_style(tbl.cell(ri,i+1),(str(c) if c>0 else "·"),11 if c>0 else 9,c>0,fc,fill)
            cell_style(tbl.cell(ri,7),tot,9.5,True,BLACK,RGBColor(0xEC,0xEC,0xEC))
    ri+=1
# row heights
for r in tbl.rows: r.height=Inches(0.34)
nb=s.shapes.add_textbox(Inches(0.5),Inches(7.28),Inches(12.3),Inches(0.22)).text_frame
rn=nb.paragraphs[0].add_run(); rn.text="※ 셀 음영이 진할수록 물량 多 · KOL 무가시딩·어필리에이트는 올영 부담/무가로 자체 비용 0 · 퍼포먼스 매체는 별도(별첨)"
rn.font.name=FONT; rn.font.size=Pt(8.5); rn.font.color.rgb=GRAY

# ============================================================
# SLIDE 2 — 견적서
# ============================================================
s2=prs.slides.add_slide(BLANK)
title(s2,"견적서 (Quotation)","올리브영 PB 올더베러 브랜드 빌딩 캠페인 · 작성일 2026.06.14 · 단위 원(VAT 별도)")
q=[
 ("G","프로젝트 대행료"),
 ("기획·운영비 (수수료 · 10%)","1식","-","25,000,000"),
 ("G","제작비"),
 ("KV·디자인·영상·콘텐츠 제작 (EGC 포함)","1식","-","83,600,000"),
 ("G","크리에이터 시딩"),
 ("매크로 (인스타·유튜브)","15","800,000","12,000,000"),
 ("마이크로","150","400,000","60,000,000"),
 ("나노","175","250,000","43,750,000"),
 ("KOL 무가시딩  *자체 비용 0","60","무가","0"),
 ("X (트위터) 시딩","15","500,000","7,500,000"),
 ("G","콘텐츠 · 바이럴"),
 ("파워페이지 콘텐츠 바이럴","9","600,000","5,400,000"),
 ("커뮤니티 체험단 (탑1·로2)","3","500/200만","9,000,000"),
 ("네이버 블로그","15","250,000","3,750,000"),
 ("G","기타"),
 ("어필리에이트 (올영 쇼핑 큐레이터)  *올영 부담","누적 60","-","0"),
 ("T","공급가액 소계","250,000,000"),
 ("T","부가가치세 (10%)","25,000,000"),
 ("TT","합계 금액 (VAT 포함)","275,000,000"),
]
nr=len(q)+1
t2=s2.shapes.add_table(nr,4,Inches(0.7),Inches(1.05),Inches(11.9),Inches(5.7)).table
t2.columns[0].width=Inches(6.7); t2.columns[1].width=Inches(1.5); t2.columns[2].width=Inches(1.7); t2.columns[3].width=Inches(2.0)
# header
cell_style(t2.cell(0,0),"세부 항목",10,True,WHITE,BLACK,PP_ALIGN.LEFT)
cell_style(t2.cell(0,1),"수량",10,True,WHITE,BLACK)
cell_style(t2.cell(0,2),"단가",10,True,WHITE,BLACK)
cell_style(t2.cell(0,3),"금액",10,True,WHITE,BLACK,PP_ALIGN.RIGHT)
ri=1
for row in q:
    if row[0]=="G":
        a=t2.cell(ri,0); b=t2.cell(ri,3); a.merge(b); cell_style(a,row[1],9.5,True,BLACK,LG,PP_ALIGN.LEFT)
    elif row[0] in ("T","TT"):
        big=row[0]=="TT"
        a=t2.cell(ri,0); b=t2.cell(ri,2); a.merge(b)
        cell_style(a,row[1],11 if big else 10,True,(WHITE if big else BLACK),(BLACK if big else WHITE),PP_ALIGN.RIGHT)
        cell_style(t2.cell(ri,3),row[2],13 if big else 10.5,True,(WHITE if big else BLACK),(BLACK if big else WHITE),PP_ALIGN.RIGHT)
    else:
        item,qty,unit,amt=row
        cell_style(t2.cell(ri,0),item,9.5,False,BLACK,WHITE,PP_ALIGN.LEFT)
        cell_style(t2.cell(ri,1),qty,9,False,DARK,WHITE)
        cell_style(t2.cell(ri,2),unit,9,False,DARK,WHITE)
        cell_style(t2.cell(ri,3),amt,10,amt not in("0",),(BLACK if amt!="0" else GRAY),WHITE,PP_ALIGN.RIGHT)
    ri+=1
for r in t2.rows: r.height=Inches(0.27)
nb2=s2.shapes.add_textbox(Inches(0.7),Inches(6.95),Inches(11.9),Inches(0.5)).text_frame
nb2.word_wrap=True
for k,n in enumerate([
 "· 프로젝트 대행료 10%(2,500만), 차액은 제작비(KV·콘텐츠 제작)에 반영 · 공급가액 2.5억은 콘텐츠·시딩·바이럴·제작·대행 운영분",
 "· 퍼포먼스 매체 운영(PA·올영 인앱)은 별도 산정(별첨) · KOL 무가시딩·어필리에이트는 올영 부담/무가로 0원 · 단가는 인스타·유튜브 기준, X 별도",
]):
    p=nb2.paragraphs[0] if k==0 else nb2.add_paragraph()
    r=p.add_run(); r.text=n; r.font.name=FONT; r.font.size=Pt(8.5); r.font.color.rgb=GRAY

prs.save("proposal/올더베러_액션플랜_견적서_편집용.pptx")
print("saved slides:",len(prs.slides._sldIdLst))
