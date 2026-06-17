# -*- coding: utf-8 -*-
"""
example2.pptx(완성형 가이드 덱)의 에어리·에디토리얼 디자인 언어로 제안덱을 재조판.
- 상단 옐로 풀폭 밴드 제거(타이틀 부유형 헤더)
- 검정 풀폭 푸터 바 -> 베이지 F3F2E7 콜아웃 + 그린 좌측 액센트 + 본문 흑색
- 구 베이지 FBF6CE -> 레퍼런스 F3F2E7 통일
원본 보존, 재구성본을 새 파일로 저장.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

SRC = "proposal/BAT_올리브영_2026_올더베러_마케팅제안_디자인적용_HAN.pptx"
DST = "proposal/BAT_올리브영_2026_올더베러_마케팅제안_디자인재구성_HAN.pptx"

EMU = 914400
BEIGE_REF = RGBColor(0xF3, 0xF2, 0xE7)   # 레퍼런스 베이지
BEIGE_OLD = "FBF6CE"
BLACK = RGBColor(0x14, 0x14, 0x12)
GREEN = RGBColor(0x09, 0x58, 0x2F)       # 브랜드 그린
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def hexof(color):
    try: return str(color.rgb)
    except: return None

def is_whiteish(h):
    return h in ("FFFFFF", "FEFFFF", "FFFEFF", "F2F2F2", "FAFAFA")

stats = {"band": 0, "footer": 0, "beige": 0}

def flip_text_to_dark(shape):
    if not shape.has_text_frame: return
    for pa in shape.text_frame.paragraphs:
        for r in pa.runs:
            try:
                if r.font.color is not None and r.font.color.type is not None:
                    h = str(r.font.color.rgb)
                    if is_whiteish(h):
                        r.font.color.rgb = BLACK
            except Exception:
                pass

def process_slide(slide):
    # 1차: 패턴 식별 (그룹 내부는 풀폭 밴드/푸터가 아니므로 최상위만)
    footers = []
    for sh in slide.shapes:
        try:
            solid = sh.fill.type == 1
        except Exception:
            solid = False
        L = (sh.left or 0) / EMU
        T = (sh.top or 0) / EMU
        W = (sh.width or 0) / EMU
        H = (sh.height or 0) / EMU
        if solid:
            h = hexof(sh.fill.fore_color)
            # 상단 옐로 풀폭 밴드 -> 제거(화이트, 라인 없음)
            if h in ("FEFFA1", "FFFF00", "FBF6CE") and W > 11 and T < 0.6 and H < 0.7:
                sh.fill.solid(); sh.fill.fore_color.rgb = WHITE
                try: sh.line.fill.background()
                except Exception: pass
                stats["band"] += 1
                continue
            # 검정 풀폭 푸터 바 -> 베이지 콜아웃
            if h in ("141412", "111111") and W > 10 and T > 6.2:
                sh.fill.solid(); sh.fill.fore_color.rgb = BEIGE_REF
                try: sh.line.fill.background()
                except Exception: pass
                flip_text_to_dark(sh)
                footers.append((sh.left, sh.top, sh.height))
                stats["footer"] += 1
                continue
    # 푸터 영역에 겹치는 별도 텍스트 박스의 흰 글씨도 흑색 반전
    for (l, t, hgt) in footers:
        top = t / EMU
        bot = (t + hgt) / EMU
        for sh in slide.shapes:
            if not sh.has_text_frame: continue
            cy = ((sh.top or 0) + (sh.height or 0) / 2) / EMU
            if top - 0.15 <= cy <= bot + 0.15:
                flip_text_to_dark(sh)
    # 푸터 콜아웃 좌측 그린 액센트 추가
    for (l, t, hgt) in footers:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Emu(int(0.07 * EMU)), hgt)
        bar.fill.solid(); bar.fill.fore_color.rgb = GREEN
        bar.line.fill.background()
        bar.shadow.inherit = False

    # 전역 베이지 통일 (그룹 내부 포함)
    def recolor(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                recolor(sh.shapes); continue
            try:
                if sh.fill.type == 1 and hexof(sh.fill.fore_color) == BEIGE_OLD:
                    sh.fill.fore_color.rgb = BEIGE_REF
                    stats["beige"] += 1
            except Exception:
                pass
            # 표 셀 베이지
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        try:
                            if c.fill.type == 1 and hexof(c.fill.fore_color) == BEIGE_OLD:
                                c.fill.fore_color.rgb = BEIGE_REF
                                stats["beige"] += 1
                        except Exception:
                            pass
    recolor(slide.shapes)

def main():
    prs = Presentation(SRC)
    for s in prs.slides:
        process_slide(s)
    prs.save(DST)
    print("saved:", DST)
    print("stats:", stats)

if __name__ == "__main__":
    main()
