#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""올더베러 플라이휠 3p — 편집 가능 네이티브 PPT (1 프레임워크 / 2 왜 시딩 1순위 / 3 왜 마이크로·나노)"""
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

CREAM=RGBColor(0xFC,0xFB,0xF0); BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34)
GRAY=RGBColor(0x78,0x78,0x72); MG=RGBColor(0xB0,0xB0,0xA8); WHITE=RGBColor(0xFF,0xFF,0xFF)
SOFT=RGBColor(0xF1,0xF0,0xE5); LINE=RGBColor(0x22,0x22,0x1F); LGT=RGBColor(0xE6,0xE6,0xDD)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def newslide():
    s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=CREAM; return s
def rrect(s,x,y,w,h,fill,line=None,lw=1.5,rad=0.1,rounded=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded:
        try: sp.adjustments[0]=rad
        except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def oval(s,x,y,w,h,fill,line=None,lw=2):
    sp=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def tri(s,cxp,cyp,w,h,color,rot=0):
    sp=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(cxp-w/2),Inches(cyp-h/2),Inches(w),Inches(h))
    sp.rotation=rot; sp.fill.solid(); sp.fill.fore_color.rgb=color; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def tbox(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(3); tf.margin_right=Pt(3); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(2)
        for (t,sz,b,c) in runs:
            r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def head(s,tag,t,sub):
    tbox(s,0.5,0.3,11.5,0.28,[[(tag,12.5,True,GRAY)]])
    tbox(s,0.5,0.58,12.0,0.6,[[(t,27,True,BLACK)]])
    tbox(s,0.5,1.2,12.3,0.34,[[(sub,13.5,True,DARK)]])
    rrect(s,0.5,1.62,12.33,0.03,BLACK,rounded=False)
def cell(c,t,sz=10,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.LEFT):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(5); c.margin_right=Pt(4); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; r=p.add_run(); r.text=t
    r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

# ===== SLIDE 1 — 플라이휠 프레임워크 =====
s=newslide()
head(s,"02_A 올더베러 캠페인 프레임워크","플라이휠 (Flywheel)","처음엔 무겁지만, 한번 돌면 팍팍 돌아갑니다 — 올더베러를 ‘플라이휠’로 굴립니다.")
# left concept
rrect(s,0.5,1.95,3.05,4.5,WHITE,LINE,2,0.06)
tbox(s,0.72,2.12,2.7,0.3,[[("개념",11,True,GRAY)]])
tbox(s,0.72,2.4,2.7,0.34,[[("플라이휠이란?",15,True,BLACK)]])
tbox(s,0.72,2.84,2.72,3.5,[
 [("멈춰 있는 바퀴는 ‘처음’ 돌리기가",11,False,DARK)],[("가장 무겁습니다.",11,False,DARK)],
 [("",5,False,DARK)],
 [("그러나 한번 돌기 시작하면 관성과",11,False,DARK)],[("가속이 붙어, 점점 적은 힘으로",11,False,DARK)],[("더 빠르게 돕니다.",11,False,DARK)],
 [("",5,False,DARK)],
 [("올더베러 마케팅도 같은 원리입니다.",11,True,BLACK)],[("가장 무거운 첫 바퀴를 다수의 힘으",11,False,DARK)],[("로 밀어 돌리면, 후기·신뢰·전환·재",11,False,DARK)],[("구매가 연료가 되어 스스로 가속합니다.",11,False,DARK)]])
# flywheel
cx,cy,R=7.35,4.4,1.78
oval(s,cx-R,cy-R,2*R,2*R,None,MG,3)
for a in (45,135,225,315):
    px=cx+R*math.cos(math.radians(a)); py=cy+R*math.sin(math.radians(a))
    tri(s,px,py,0.32,0.34,BLACK,rot=a+180)
oval(s,cx-0.78,cy-0.78,1.56,1.56,BLACK)
tbox(s,cx-0.78,cy-0.42,1.56,0.4,[[("올더베러",16,True,WHITE)]],PP_ALIGN.CENTER)
tbox(s,cx-0.78,cy+0.04,1.56,0.3,[[("FLYWHEEL",10.5,True,MG)]],PP_ALIGN.CENTER)
nodes=[(cx,cy-R,"① 시딩","콘텐츠·후기 볼륨 (첫 바퀴)"),
 (cx+R,cy,"② 신뢰","다수 반복 노출 → 익숙함"),
 (cx,cy+R,"③ 전환·랭킹","위닝·PA·올영세일"),
 (cx-R,cy,"④ 재구매","후기·CRM (다음 연료)")]
nw,nh=2.0,0.92
for nx,ny,t,sub in nodes:
    rrect(s,nx-nw/2,ny-nh/2,nw,nh,WHITE,LINE,2,0.1)
    tbox(s,nx-nw/2,ny-0.32,nw,0.34,[[(t,12.5,True,BLACK)]],PP_ALIGN.CENTER)
    tbox(s,nx-nw/2,ny+0.02,nw,0.3,[[(sub,9.5,True,GRAY)]],PP_ALIGN.CENTER)
# heavy first push label
tbox(s,cx-1.5,cy-R-0.66,3.0,0.3,[[("↑ 가장 무거운 첫 바퀴",10.5,True,DARK)]],PP_ALIGN.CENTER)
# right operation
rrect(s,10.65,1.95,2.18,4.5,SOFT,LINE,2,0.06)
tbox(s,10.82,2.12,1.9,0.3,[[("작동 원리",11,True,GRAY)]])
rrect(s,10.82,2.45,1.84,1.02,BLACK,rad=0.1)
tbox(s,10.82,2.55,1.84,0.9,[[("첫 바퀴 (7월)",11.5,True,WHITE)],[("가장 무겁게,",10,False,LGT)],[("강하게 민다",10,False,LGT)]],PP_ALIGN.CENTER)
rrect(s,10.82,3.62,1.84,1.02,WHITE,LINE,2,0.1)
tbox(s,10.82,3.72,1.84,0.9,[[("가속",11.5,True,BLACK)],[("관성이 붙어",10,False,DARK)],[("점점 빨리",10,False,DARK)]],PP_ALIGN.CENTER)
rrect(s,10.82,4.79,1.84,1.05,BLACK,rad=0.1)
tbox(s,10.82,4.9,1.84,0.92,[[("정상",11.5,True,WHITE)],[("올영세일",10,False,LGT)],[("랭킹 1위",10,False,LGT)]],PP_ALIGN.CENTER)
# bottom connector
rrect(s,0.5,6.62,12.33,0.6,BLACK,rad=0.12)
tbox(s,0.5,6.62,12.33,0.6,[[("가장 무거운 ‘첫 바퀴’ = 시딩  ·  그 바퀴를 돌리는 ‘힘’ = 다수의 마이크로·나노  —  다음 두 장에서 증명합니다.",14.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# ===== SLIDE 2 — 왜 시딩 1순위 (첫 바퀴) =====
s=newslide()
head(s,"02_A 플라이휠 ①","왜 인플루언서 시딩이 1순위인가","플라이휠의 ‘첫 바퀴’ — 가장 무겁지만, 가장 먼저 밀어야 하는 곳이 시딩입니다.")
cards=[
 ("01","‘올더베러’의 존재를 시장에 선보여야 합니다","브랜드 검색 월 약 3,000 · 해시태그 #올더베러 +100여 건","네이버·구글 검색량 + 인스타그램 해시태그 집계",
  "다수의 진성(나노·마이크로) UGC 콘텐츠를 전략 일자에 노출해 ‘익숙함과 신뢰’를 먼저 쌓습니다."),
 ("02","소비자는 사기 전에 검증합니다","제품 검색 → 올영 이동 → 후기·리뷰 확인 후 구매(검증형) · VOC 재구매율 15.5% · 구매이유 1위 ‘가격·할인’","리스닝마인드 검색 경로 분석 + 올더베러 리뷰 1,104건 VOC 딥다이브(BAT 자체 크롤링)",
  "가격 외 ‘선택 근거’(후기·리뷰·해시태그·콘텐츠 볼륨)를 먼저 쌓아 잘 보이는 제품으로 전환을 견인합니다."),
 ("03","‘어떻게 먹는지’는 인플루언서가 보여줍니다","RFP 요구 ‘언제·왜·어떻게 먹는지 생활 맥락’에 따른 영상으로 학습하고 디토(따라 구매)하는 소비자들","올리브영 VOC 딥다이브 + 인스타 릴스 콘텐츠 분석",
  "SNS 인플루언서가 실제 사용 상황(아침 공복·야근·취침 전)을 자기 일상으로 가져오는 소비자 CDJ 패턴에 들어섭니다."),
 ("04","이미 올영 PB에서 고성과로 검증된 방식입니다","우리가 운영 중인 올영 PB 시딩+PA가 고성과: 브링그린 검색 9K→140K(15배) · 산리오 파트너십 ROAS 127→356→589% · 올영세일 ROAS 324→515%","BAT 올리브영 PB 운영 캠페인 실적(레퍼런스)",
  "‘시딩 → 위닝 콘텐츠 → PA 2차’ 성공 공식을 올더베러에 그대로 적용합니다. (한정 비용 내 최고 효율)"),
]
cw=2.96;gap=0.12;x0=0.5;cy=1.92;ch=4.32;astrip=1.4
for i,(num,t,evd,src,ans) in enumerate(cards):
    x=x0+i*(cw+gap)
    rrect(s,x,cy,cw,ch,WHITE,LINE,2,0.06)
    tbox(s,x+0.2,cy+0.12,cw-0.36,0.4,[[(num,18,True,MG)]])
    tbox(s,x+0.2,cy+0.5,cw-0.36,0.7,[[(t,12.5,True,BLACK)]])
    tbox(s,x+0.2,cy+1.28,cw-0.36,2.0,[[("데이터 근거",8.5,True,GRAY)],[(evd,9.3,False,DARK)],[("",4,False,GRAY)],[("출처 ",8,True,GRAY),(src,8.3,False,GRAY)]])
    rrect(s,x,cy+ch-astrip,cw,astrip,SOFT,rounded=False)
    rrect(s,x,cy+ch-astrip,0.07,astrip,BLACK,rounded=False)
    tbox(s,x+0.2,cy+ch-astrip+0.08,cw-0.36,astrip-0.14,[[("→ 첫 바퀴를 시딩으로 미는 이유",9,True,BLACK)],[(ans,9.8,True,BLACK)]])
rrect(s,0.5,6.5,12.33,0.56,BLACK,rad=0.12)
tbox(s,0.5,6.5,12.33,0.56,[[("그래서 인플루언서 시딩이 1순위입니다 — 플라이휠을 돌리는, 가장 무겁고 가장 먼저인 ‘첫 바퀴’.",15,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
tbox(s,0.5,7.16,12.3,0.3,[[("※ 근거: 검색량·해시태그 집계 / 리스닝마인드 경로분석 / VOC 1,104건 딥다이브 / 인스타 릴스 분석 / BAT 올영 PB 운영 실적",9,False,GRAY)]])

# ===== SLIDE 3 — 왜 마이크로·나노 (미는 힘) =====
s=newslide()
head(s,"02_A 플라이휠 ②","왜 마이크로·나노 중심인가","그 무거운 첫 바퀴를 돌리는 ‘힘’ — 한 명의 큰 힘이 아니라, 다수의 마이크로·나노입니다.")
tbox(s,0.5,1.78,11.5,0.3,[[("한눈 비교 — 메가·매크로 vs 마이크로·나노 (2026 기준)",12,True,BLACK)]])
colw=[2.5,4.85,4.98]; colx=[0.5,3.0,7.85]
ty=2.12; rows=[("지표","메가·매크로 (넓고 얕음)","마이크로·나노 (좁고 깊음) ★"),
 ("인게이지먼트(IG)","평균 약 1%대에 머뭅니다","나노가 평균 6.23%로 전 티어 중 가장 높습니다"),
 ("전환율","매크로 캠페인은 상대적으로 낮습니다","나노·마이크로가 매크로 대비 2~3배 높습니다"),
 ("진성도·신뢰","광고로 인식되어 신뢰가 낮습니다","나노가 전 티어 중 신뢰가 가장 높습니다"),
 ("비용·ROI","고비용·일회성에 그칩니다","73% 브랜드가 우선 선택하며 ROI도 더 우수합니다")]
g=s.shapes.add_table(5,3,Inches(0.5),Inches(ty),Inches(12.33),Inches(1.95)).table
g.columns[0].width=Inches(2.5); g.columns[1].width=Inches(4.85); g.columns[2].width=Inches(4.98)
for ri,row in enumerate(rows):
    for ci,val in enumerate(row):
        if ri==0: cell(g.cell(ri,ci),val,11,True,WHITE,BLACK)
        else:
            fill=(RGBColor(0xE9,0xE8,0xDB) if ci==2 else (WHITE if ri%2 else RGBColor(0xF6,0xF5,0xEC)))
            cell(g.cell(ri,ci),val,10,(ci!=1),(BLACK if ci!=1 else GRAY),fill)
for r in g.rows: r.height=Inches(0.39)
# 3 cards
c2=[("진성도","광고가 아니라 ‘내 주변의 후기’입니다","대형 인플루언서의 일회성 광고보다, 니치한 생활자가 실제 루틴처럼 말하는 콘텐츠가 웰니스 제품의 신뢰를 만듭니다.","소비자는 작은 크리에이터를 더 신뢰합니다 (나노 인스타 ER 6.23%로 전 티어 최고)."),
 ("반복 노출","다수가 같은 시점·다양한 TPO로 말합니다","다수의 마이크로/나노가 같은 시점에 다양한 TPO로 말할 때 ‘요즘 올영에서 많이 보이는 제품’ 인식이 형성됩니다.","73%의 브랜드가 마이크로·미드티어 크리에이터를 우선 선택합니다(2026)."),
 ("자산화","저단가·다수 → 후기·검색·퍼포의 원천입니다","저단가·다수 콘텐츠를 통해 후기·검색·퍼포먼스 소재로 이어지는, 활용 가능한 원천 콘텐츠를 확보합니다.","유료 광고로 쓰면 ER 2~3배·CPA 절감 (위닝 콘텐츠 PA 2차·BAT 실적).")]
cw=3.97;gap=0.2;cy=4.3;ch=1.95
for i,(t,sub,body,evd) in enumerate(c2):
    x=0.5+i*(cw+gap)
    rrect(s,x,cy,cw,ch,WHITE,LINE,2,0.08)
    tbox(s,x+0.22,cy+0.13,cw-0.44,0.34,[[(t,15,True,BLACK)]])
    tbox(s,x+0.22,cy+0.47,cw-0.44,0.5,[[(sub,11,True,DARK)]])
    tbox(s,x+0.22,cy+0.92,cw-0.44,0.7,[[(body,9.5,False,DARK)]])
    rrect(s,x,cy+ch-0.5,cw,0.5,SOFT,rounded=False); rrect(s,x,cy+ch-0.5,0.06,0.5,BLACK,rounded=False)
    tbox(s,x+0.2,cy+ch-0.46,cw-0.34,0.44,[[("데이터 ",8,True,GRAY),(evd,9.2,True,BLACK)]])
rrect(s,0.5,6.46,12.33,0.56,BLACK,rad=0.12)
tbox(s,0.5,6.46,12.33,0.56,[[("메가는 ‘광고’가 되고, 마이크로·나노는 ‘후기’가 됩니다 — 다수의 작은 힘이 플라이휠을 굴립니다.",15.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
tbox(s,0.5,7.16,12.3,0.3,[[("※ 근거: 2026 인플루언서 마케팅 벤치마크(인게이지먼트·전환·신뢰·ROI) + 브랜드 선호 조사 + BAT 올영 PB 운영 실적",9,False,GRAY)]])

prs.save("proposal/올더베러_플라이휠_3p.pptx")
print("slides:",len(prs.slides._sldIdLst))
