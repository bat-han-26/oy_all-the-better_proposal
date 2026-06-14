#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""병합 덱에 크리에이터 시딩 전략 2장 추가(역할×채널 / KOL 무가시딩), 시딩 앵글 앞에 삽입"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DECK="proposal/올더베러_제안덱_병합본.pptx"
BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
F5=RGBColor(0xF5,0xF5,0xF5)
FONT="맑은 고딕"
prs=Presentation(DECK)
BLANK=None
for l in prs.slide_layouts:
    if l.name=="빈화면": BLANK=l
if BLANK is None: BLANK=prs.slide_layouts[5]
SW,SH=prs.slide_width,prs.slide_height

def rect(s,x,y,w,h,c):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=c; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(x,y,w,h).text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(3)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def title(s,t,sub=None,tag=None):
    rect(s,0,0,SW,Inches(0.14),BLACK)
    tb=s.shapes.add_textbox(Inches(0.5),Inches(0.34),Inches(9.5),Inches(0.95)).text_frame; tb.word_wrap=True
    r=tb.paragraphs[0].add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(23); r.font.bold=True; r.font.color.rgb=BLACK
    if sub:
        r2=tb.add_paragraph().add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(12); r2.font.color.rgb=GRAY
    rect(s,Inches(0.52),Inches(1.4),Inches(1.2),Inches(0.05),BLACK)
    if tag:
        tw=Inches(2.7); rect(s,Emu(int(SW)-int(tw)-int(Inches(0.5))),Inches(0.4),tw,Inches(0.46),LG)
        txt(s,Emu(int(SW)-int(tw)-int(Inches(0.5))),Inches(0.46),tw,Inches(0.36),[(tag,12,True,BLACK)],PP_ALIGN.CENTER)
def cell(c,t,sz=9.5,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.CENTER):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(5); c.margin_right=Pt(5); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; p.word_wrap=True
    r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

NEW=[]
def newslide():
    s=prs.slides.add_slide(BLANK); NEW.append(s); return s

# ============ S1. 크리에이터 전략 (역할 × 채널) ============
s=newslide()
title(s,"크리에이터 전략 — 역할 × 채널",
      "진성 다수(나노·마이크로) = 코어 / 매크로 소수 = 점화 · 인스타그램 절대 비중")

# ① 역할별
txt(s,Inches(0.5),Inches(1.6),Inches(7),Inches(0.3),[("① 역할별 구분 (비중 = 견적서 기준)",14,True,BLACK)])
txt(s,Inches(7.0),Inches(1.66),Inches(5.85),Inches(0.3),
    [("※ ER(인게이지먼트율) = (좋아요+댓글+저장+공유) ÷ 팔로워, 게시물당 평균 · 인스타그램 기준",9,False,GRAY)],PP_ALIGN.RIGHT)
t=s.shapes.add_table(4,5,Inches(0.5),Inches(1.98),Inches(12.35),Inches(2.0)).table
t.columns[0].width=Inches(1.4); t.columns[1].width=Inches(1.5); t.columns[2].width=Inches(1.3)
t.columns[3].width=Inches(4.9); t.columns[4].width=Inches(3.25)
for c,h in enumerate(["구분","팔로워","ER(IG)","역할","물량 · 예산"]):
    cell(t.cell(0,c),h,10,True,WHITE,DARK,PP_ALIGN.LEFT if c in(3,) else PP_ALIGN.CENTER)
rows=[
 ["나노","1만 이하","4~6%","진성 후기·UGC 볼륨 — 동일 메시지 동시다발 = ‘익숙함’","225건 / 56,250,000원 (44%)"],
 ["마이크로","1만~5만","2~4%","카테고리 전문·전환 코어 — 위닝 소재 발굴 → PA 2차","150건 / 60,000,000원 (47%)"],
 ["매크로","5만 이상","1~2%","도달·화제성 점화 — 세일/블프 집중","15건 / 12,000,000원 (9%)"],
]
for ri,row in enumerate(rows,1):
    fill=WHITE if ri%2 else F5
    cell(t.cell(ri,0),row[0],11,True,BLACK,fill)
    cell(t.cell(ri,1),row[1],10,False,DARK,fill)
    cell(t.cell(ri,2),row[2],10,False,GRAY,fill)
    cell(t.cell(ri,3),row[3],9.5,False,BLACK,fill,PP_ALIGN.LEFT)
    cell(t.cell(ri,4),row[4],9.5,True,BLACK,fill,PP_ALIGN.LEFT)
for r in t.rows: r.height=Inches(0.5)
txt(s,Inches(0.5),Inches(3.98),Inches(12.3),Inches(0.28),
    [("크리에이터 코어(나노·마이크로·매크로) 합계 = 390건 / 128,250,000원",10,True,GRAY)])

# ② 채널별
txt(s,Inches(0.5),Inches(4.32),Inches(7),Inches(0.3),[("② 채널별 구분 (인스타그램 절대 비중)",14,True,BLACK)])
t=s.shapes.add_table(6,3,Inches(0.5),Inches(4.68),Inches(8.4),Inches(2.4)).table
t.columns[0].width=Inches(2.2); t.columns[1].width=Inches(1.3); t.columns[2].width=Inches(4.9)
for c,h in enumerate(["채널","포맷","역할"]):
    cell(t.cell(0,c),h,10,True,WHITE,DARK,PP_ALIGN.LEFT if c==2 else PP_ALIGN.CENTER)
ch=[
 ["★ 인스타그램","릴스","시딩 코어·전환 허브 → 메타 파트너십 광고 2차 활용",True],
 ["유튜브","쇼츠","확산 보조 + 제품 사용 장면",False],
 ["틱톡","쇼츠","별도 시딩 상품 — 숏폼 바이럴·확산 (독립 물량·소재)",False],
 ["X(트위터)","피드","실시간 트렌드 + 바이럴 · 15건 / 7,500,000원",False],
 ["네이버 블로그","정보성 리뷰","후기 신뢰·정보 탐색 대응 · 15건 / 3,750,000원",False],
]
for ri,row in enumerate(ch,1):
    fill=RGBColor(0xEA,0xEA,0xEA) if row[3] else (WHITE if ri%2 else F5)
    cell(t.cell(ri,0),row[0],10,True,BLACK,fill,PP_ALIGN.LEFT)
    cell(t.cell(ri,1),row[1],9.5,True,DARK,fill)
    cell(t.cell(ri,2),row[2],9.5,row[3],BLACK,fill,PP_ALIGN.LEFT)
for r in t.rows: r.height=Inches(0.4)

# 웰니스 풀 (우측)
rect(s,Inches(9.1),Inches(4.68),Inches(3.75),Inches(2.4),F5)
txt(s,Inches(9.3),Inches(4.82),Inches(3.4),Inches(0.4),[("크리에이터 풀 = 웰니스 접점 타겟",11.5,True,BLACK)])
txt(s,Inches(9.3),Inches(5.28),Inches(3.4),Inches(1.8),[
 ("· 홈트 · 헬스 / 피트니스",10,False,DARK),
 ("· 요가 / 필라테스 · 러닝크루",10,False,DARK),
 ("· 헬시 레시피 / 식단 기록",10,False,DARK),
 ("· 다이어트 · 바디프로필",10,False,DARK),
 ("· 직장인 자기관리 · 모닝루틴 브이로그",10,False,DARK),
 ("· 헬시플레저 라이프스타일",10,False,DARK),
])

# ============ S2. KOL 무가시딩 운영 전략 ============
s=newslide()
title(s,"KOL 무가시딩 운영 전략 — 대상 × 방식",
      "광고비 지불 없이 제품·키트만 제공 → 자발 게시 유도 · 관계로 지속 발화")
# intro band
rect(s,Inches(0.5),Inches(1.62),Inches(12.35),Inches(0.6),BLACK)
txt(s,Inches(0.7),Inches(1.72),Inches(12),Inches(0.42),
    [("무가시딩 = 노출을 사는 게 아니라 ‘관계를 빌드업’ 하는 활동  ·  60건 / 6,000,000원",12.5,True,WHITE)],anchor=MSO_ANCHOR.MIDDLE)

# ① 누구에게
rect(s,Inches(0.5),Inches(2.45),Inches(6.0),Inches(0.5),DARK)
txt(s,Inches(0.7),Inches(2.52),Inches(5.6),Inches(0.4),[("① 누구에게 — 대상 선정",13.5,True,WHITE)])
txt(s,Inches(0.6),Inches(3.1),Inches(5.85),Inches(3.6),[
 ("웰니스 접점이 강한 마이크로~매크로",12,True,BLACK),
 ("브랜드 무드(‘채움’) 적합도 우선",10.5,False,GRAY),
 ("",4,False,GRAY),
 ("올더베러 타깃과 라이프스타일이 겹치는 결",12,True,BLACK),
 ("30대 직장인 · 웰니스 라이프 · 부모",10.5,False,GRAY),
 ("",4,False,GRAY),
 ("이미 자생적으로 헬시플레저·자기관리",12,True,BLACK),
 ("콘텐츠를 올리는 진성 계정 (협찬 티 안 나는 결)",10.5,False,GRAY),
])

# ② 어떻게
rect(s,Inches(6.85),Inches(2.45),Inches(6.0),Inches(0.5),DARK)
txt(s,Inches(7.05),Inches(2.52),Inches(5.6),Inches(0.4),[("② 어떻게 — 운영 방식",13.5,True,WHITE)])
txt(s,Inches(6.95),Inches(3.1),Inches(5.85),Inches(3.6),[
 ("시그니처 키트 — 브랜드 무드 담은 패키지 + 큐레이션 (단순 택배 ❌)",11,True,BLACK),
 ("커스텀 레터 — 개인화 메시지 (보도자료식 ❌)",11,True,BLACK),
 ("등급별 차등 — 매크로급=커스텀 키트+정기 접촉 / 마이크로·나노=제품 제공",11,True,BLACK),
 ("게시 강요 ❌ — 자발 게시 유도, 게시해 준 분은 본품 유지 + 주기 발송으로 관계 빌드업",11,True,BLACK),
 ("자산화 — 오가닉 후기·UGC, 우수 콘텐츠는 퍼포 소재 2차 활용",11,True,BLACK),
])

# KPI band
rect(s,Inches(0.5),Inches(6.55),Inches(12.35),Inches(0.62),F5)
txt(s,Inches(0.7),Inches(6.64),Inches(12),Inches(0.46),
    [("KPI (게런티 X) — 키트 발송 수  ·  오가닉 게시 전환율  ·  우호 후기 수  ·  도달 / 저장",11.5,True,BLACK)],anchor=MSO_ANCHOR.MIDDLE)

# ---- 삽입: '시딩 콘텐츠 앵글' 슬라이드 앞 ----
n_before=len(prs.slides._sldIdLst)-len(NEW)
target=n_before
for i,sl in enumerate(list(prs.slides)[:n_before]):
    txtall=" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    if "크리에이터 콘텐츠 기획" in txtall and "콘텐츠 앵글" in txtall:
        target=i; break

sldIdLst=prs.slides._sldIdLst; ids=list(sldIdLst); moved=ids[-len(NEW):]
for el in moved: sldIdLst.remove(el)
for off,el in enumerate(moved): sldIdLst.insert(target+off,el)
prs.save(DECK)
print("total:",len(prs.slides._sldIdLst),"| added:",len(NEW),"| inserted at idx",target)
