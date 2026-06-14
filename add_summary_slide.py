#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""공유 덱 41페이지(한판 요약)에 구조 인포그래픽을 네이티브 도형으로 추가 (FOUNDATION 제외)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC="/root/.claude/uploads/bc09908a-8d16-5be3-806b-53899e2cede8/aa4e42b6-_____________1.pptx"
OUT="proposal/올더베러_제안덱_병합본.pptx"
BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34); GRAY=RGBColor(0x78,0x78,0x72)
MG=RGBColor(0xB4,0xB4,0xAC); WHITE=RGBColor(0xFF,0xFF,0xFF); SOFT=RGBColor(0xF1,0xF0,0xE6); LINE=RGBColor(0x20,0x20,0x1E)
FONT="맑은 고딕"
prs=Presentation(SRC)
s=list(prs.slides)[40]  # page 41

def rr(x,y,w,h,fill,line=None,lw=1.5,rad=0.10):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try: sp.adjustments[0]=rad
    except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def tb(x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(2); tf.margin_right=Pt(2); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(2)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
def chevron(x,y):
    tb(x,y,0.5,0.4,[("▶",18,True,BLACK)],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# Row1 PROBLEM -> STRATEGY
rr(0.5,1.5,5.55,0.92,None,LINE,2)
tb(0.7,1.58,5.2,0.3,[("PROBLEM  ·  진단",10.5,True,GRAY)])
tb(0.7,1.82,5.2,0.4,[("팔리는데 기억되지 않는다",16,True,BLACK)])
tb(0.7,2.16,5.2,0.3,[("브랜드 검색 월 3천 · 시그니처 키워드 0",11,False,DARK)])
chevron(6.12,1.78)
rr(6.7,1.5,6.13,0.92,BLACK)
tb(6.9,1.58,5.8,0.3,[("STRATEGY  ·  전략",10.5,True,MG)])
tb(6.9,1.8,5.8,0.36,[("‘우연 → 의도’ 각인  ·  시그니처 키워드 「채움」",14.5,True,WHITE)])
tb(6.9,2.16,5.8,0.3,[("EAT — 매일(Everyday) · 쉽게(Accessible) · 맛있게(Tasty)",10.5,False,RGBColor(0xE2,0xE2,0xDB))])

# MISSION bar
rr(0.5,2.56,12.33,0.5,DARK)
tb(0.72,2.64,1.6,0.34,[("MISSION",11,True,MG)],anchor=MSO_ANCHOR.MIDDLE)
tb(2.05,2.6,10.6,0.42,[("7월 중 올영 카테고리 랭킹 1위 달성  →  9·12월 올영세일 확대  →  웰니스 대표 PB 브랜드 착지",13.5,True,WHITE)],anchor=MSO_ANCHOR.MIDDLE)

# ENGINE label
tb(0.5,3.16,12.33,0.3,[("ENGINE   ·   SEED → PROVE → BOOST → RETAIN   (시딩–콘텐츠–퍼포 연결)",12,True,DARK)])

# 4 step boxes
steps=[
 ("1","SEED","경험·후기 생산",["마이크로/나노 시딩 (월 대량)","챌린저스 · KOL 무가시딩","EGC · 제품 콘텐츠 제작"]),
 ("2","PROVE","선택 근거 증명",["리뷰·해시태그·UGC 축적","커뮤니티 체험단(탑)·파워페이지","네이버 블로그 SEO·검색 자산"]),
 ("3","BOOST","전환 증폭",["퍼포 매체 (PA·올영 인앱)","LMF 위닝 소재 · 어필리에이트","올영세일 집중 집행"]),
 ("4","RETAIN","재구매·확장",["CRM(플친)·루틴 리마인드","위닝 소재 재활용","라인 확장 (크로스셀)"]),
]
bw=2.95; gap=0.14; x0=0.5; ey=3.5; hh=0.44; bodyh=1.62
for i,(num,en,ko,tac) in enumerate(steps):
    x=x0+i*(bw+gap)
    rr(x,ey,bw,hh,BLACK,rad=0.16)
    tb(x+0.14,ey+0.06,0.4,0.32,[(num,15,True,MG)],anchor=MSO_ANCHOR.MIDDLE)
    tb(x+0.5,ey+0.05,1.4,0.34,[(en,14,True,WHITE)],anchor=MSO_ANCHOR.MIDDLE)
    tb(x+bw-1.55,ey+0.08,1.45,0.3,[(ko,10.5,True,RGBColor(0xE2,0xE2,0xDB))],PP_ALIGN.RIGHT,MSO_ANCHOR.MIDDLE)
    rr(x,ey+hh,bw,bodyh,WHITE,LINE,1.5,rad=0.06)
    tb(x+0.16,ey+hh+0.06,bw-0.3,0.24,[("전술",9.5,True,GRAY)])
    lines=[]
    for tt in tac: lines.append(("•  "+tt,11,False,DARK))
    tf=s.shapes.add_textbox(Inches(x+0.16),Inches(ey+hh+0.32),Inches(bw-0.3),Inches(bodyh-0.4)).text_frame
    tf.word_wrap=True
    for j,(t,sz,b,c) in enumerate(lines):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.space_after=Pt(8)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    if i<3: chevron(x+bw+gap/2-0.25,ey+hh+0.5)

# Timeline
tb(0.5,5.62,12.33,0.26,[("6개월 운영 · 1·2차 웨이브",11,True,GRAY)])
months=[("7월","착수",False),("8월","★1차 웨이브",True),("9월","★1차 피크·세일",True),("10월","유지",False),("11월","★2차·블프",True),("12월","★2차 피크·세일",True)]
tw=2.0; tg=0.066; tx=0.5; ty=5.9
for i,(m,lab,dark) in enumerate(months):
    x=tx+i*(tw+tg)
    rr(x,ty,tw,0.6,(BLACK if dark else SOFT),(None if dark else LINE),1.3,rad=0.12)
    tb(x+0.16,ty+0.07,tw-0.3,0.28,[(m,14,True,(WHITE if dark else BLACK))])
    tb(x+0.16,ty+0.34,tw-0.3,0.24,[(lab,10,True,(RGBColor(0xE6,0xE6,0xDF) if dark else DARK))])

# KPI chips + 차별점
chips=["브랜드 검색량 3,000 → 10,000","올영세일 랭킹 1위","시그니처 키워드 노출률 90%+","예산 2.5억 (퍼포 별첨)"]
cw=2.95; cg=0.13; cx=0.5; ky=6.62
for i,c in enumerate(chips):
    x=cx+i*(cw+cg)
    rr(x,ky,cw,0.38,WHITE,LINE,1.5,rad=0.5)
    tb(x,ky+0.05,cw,0.3,[(c,11,True,BLACK)],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
rr(0.5,7.06,12.33,0.36,SOFT,LINE,1.2,rad=0.2)
tb(0.7,7.1,12.0,0.3,[("차별점   웰니스 심의 다이나믹스 대응 + AI 심의 가이드 검증  ·  콘텐츠 × 퍼포먼스 × 운영을 하나의 구매 흐름으로 설계 = BAT",11,True,DARK)],anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved. page41 shapes now:",len(list(prs.slides)[40].shapes))
