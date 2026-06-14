#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""병합 덱에 '크리에이터 관계 빌드업 프로세스' 1장 추가 — KOL 무가시딩 다음에 삽입"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DECK="proposal/올더베러_제안덱_병합본.pptx"
BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
F5=RGBColor(0xF5,0xF5,0xF5)
FONT="맑은 고딕"
prs=Presentation(DECK)
BLANK=None
for l in prs.slide_layouts:
    if l.name=="빈화면": BLANK=l
if BLANK is None: BLANK=prs.slide_layouts[5]
SW,SH=prs.slide_width,prs.slide_height

def rect(s,x,y,w,h,c):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=c; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(x,y,w,h).text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(2)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def title(s,t,sub=None):
    rect(s,0,0,SW,Inches(0.14),BLACK)
    tb=s.shapes.add_textbox(Inches(0.5),Inches(0.34),Inches(12),Inches(0.95)).text_frame; tb.word_wrap=True
    r=tb.paragraphs[0].add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(23); r.font.bold=True; r.font.color.rgb=BLACK
    if sub:
        r2=tb.add_paragraph().add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(12); r2.font.color.rgb=GRAY
    rect(s,Inches(0.52),Inches(1.4),Inches(1.2),Inches(0.05),BLACK)

NEW=[]
def newslide():
    s=prs.slides.add_slide(BLANK); NEW.append(s); return s

# ============ 관계 빌드업 프로세스 ============
s=newslide()
title(s,"크리에이터 관계 빌드업 프로세스",
      "1회성 협찬이 아니라, ‘자발적으로 계속 올리는’ 우호 크리에이터를 누적해 가는 5단계 관계 설계")

stages=[
 ("STEP 1","발굴 · 컨택","낯선 제안 → 호감",
  ["버티컬·등급별 롱리스트에서 진성·심의 스크리닝","개인화 첫 컨택(보도자료식 ❌)","브랜드 무드·취지 공유로 호감 형성"],
  "응답·수락 → 다음 단계"),
 ("STEP 2","첫 시딩","제품 경험 제공",
  ["시그니처 키트(‘채움’ 무드) + 커스텀 레터","제품 큐레이션·사용 가이드 동봉","게시 강요 ❌ — 자발 경험 유도"],
  "오가닉 첫 게시 → 다음 단계"),
 ("STEP 3","관계 형성","소통 · 반응 누적",
  ["게시물 반응·DM 소통·진심 피드백","반응 좋은 분에게 신제품 우선 발송","콘텐츠 결·반응 데이터 기록"],
  "재게시·우호 반응 → 다음 단계"),
 ("STEP 4","정론화","우호 관계 고정",
  ["게시해 준 분 본품 유지 + 정기 발송","시즌·신제품마다 재컨택(관계 지속)","등급 상향 시 커스텀 키트·정기 접촉"],
  "지속 발화·신뢰 → 다음 단계"),
 ("STEP 5","자산화 · 앰배서더","반복 발화 · 소재 전환",
  ["오가닉 반복 게시 = 검색·리뷰 자산","우수 콘텐츠 → 퍼포(PA·파트너십) 2차 활용","장기 우호풀 = 세일 시점 즉시 가동"],
  "브랜드 우호 앰배서더 풀 누적"),
]
n=len(stages); gap=Inches(0.18); x0=Inches(0.5)
total_w=int(SW)-2*int(Inches(0.5)); cw=Emu((total_w-(n-1)*int(gap))//n)
y=Inches(1.75); ch=Inches(4.55)
for i,(no,name,goal,acts,sig) in enumerate(stages):
    x=Emu(int(x0)+i*(int(cw)+int(gap)))
    # header
    rect(s,x,y,cw,Inches(0.92),BLACK)
    txt(s,Emu(int(x)+int(Inches(0.1))),Emu(int(y)+int(Inches(0.08))),Emu(int(cw)-int(Inches(0.2))),Inches(0.8),
        [(no,9.5,True,MG),(name,12.5,True,WHITE),(goal,9,False,RGBColor(0xCF,0xCF,0xCF))])
    # body
    by=Emu(int(y)+int(Inches(0.98)))
    rect(s,x,by,cw,Inches(2.75),F5)
    runs=[("• "+a,9,False,DARK) for a in acts]
    txt(s,Emu(int(x)+int(Inches(0.12))),Emu(int(by)+int(Inches(0.1))),Emu(int(cw)-int(Inches(0.24))),Inches(2.6),runs)
    # signal
    sy=Emu(int(by)+int(Inches(2.81)))
    rect(s,x,sy,cw,Inches(0.78),LG)
    txt(s,Emu(int(x)+int(Inches(0.1))),Emu(int(sy)+int(Inches(0.06))),Emu(int(cw)-int(Inches(0.2))),Inches(0.66),
        [("▶ "+sig,8.5,True,BLACK)])
    # arrow between
    if i<n-1:
        ax=Emu(int(x)+int(cw)+int(Inches(0.02)))
        txt(s,ax,Emu(int(y)+int(Inches(1.9))),Inches(0.18),Inches(0.4),[("›",16,True,GRAY)],PP_ALIGN.CENTER)

# 하단 핵심 메시지 밴드
rect(s,Inches(0.5),Inches(6.55),Inches(12.35),Inches(0.62),DARK)
txt(s,Inches(0.7),Inches(6.64),Inches(12),Inches(0.46),
    [("핵심 — 무가시딩의 ROI는 ‘한 번의 노출’이 아니라 ‘반복 발화하는 우호 크리에이터 풀’ · 세일 시점에 비용 없이 즉시 가동되는 자산",11,True,WHITE)],
    anchor=MSO_ANCHOR.MIDDLE)

# ---- 삽입: 'KOL 무가시딩 운영 전략' 다음 ----
n_before=len(prs.slides._sldIdLst)-len(NEW)
target=n_before
for i,sl in enumerate(list(prs.slides)[:n_before]):
    txtall=" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    if "KOL 무가시딩 운영 전략" in txtall:
        target=i+1; break

sldIdLst=prs.slides._sldIdLst; ids=list(sldIdLst); moved=ids[-len(NEW):]
for el in moved: sldIdLst.remove(el)
for off,el in enumerate(moved): sldIdLst.insert(target+off,el)
prs.save(DECK)
print("total:",len(prs.slides._sldIdLst),"| added:",len(NEW),"| inserted at idx",target)
