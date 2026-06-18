# -*- coding: utf-8 -*-
"""올더베러 견적 v2 + 월별 액션플랜 — 수정 가능 Excel.
모델: 투입비(집행) 2.5억 고정 + 마크업(수수료) 별도(15~17.5%) + 무상 서비스 0원.
제작비(KV)는 목표 투입비 자동정산(잔액)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BROWN="3D2B1B"; BEIGE="F3F2E7"; GREEN="0E5A30"; GRAYF="EFEFEC"; WHITE="FFFFFF"; LIME="C7CE3E"
thin=Side(style="thin",color="D8D5CC"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def C(ws,r,c,v=None,bold=False,color="141412",fill=None,align="left",fmt=None,size=10,wrap=False):
    cc=ws.cell(row=r,column=c)
    if v is not None: cc.value=v
    cc.font=Font(name="Pretendard",size=size,bold=bold,color=color)
    if fill: cc.fill=PatternFill("solid",fgColor=fill)
    cc.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap); cc.border=border
    if fmt: cc.number_format=fmt
    return cc
NUM="#,##0"

wb=openpyxl.Workbook()
ws=wb.active; ws.title="견적서"
for i,w in enumerate([24,40,11,14,16,34]): ws.column_dimensions[get_column_letter(i+1)].width=w
C(ws,1,1,"올더베러 캠페인 견적서 (Quotation)",bold=True,size=16)
C(ws,2,1,"올리브영 PB 올더베러 브랜드 빌딩 · 단위 원(VAT 별도) · 견적방식: 전체 투입비 대비 수수료율(마크업)",color="868680",size=10)
# 컨트롤 셀
C(ws,2,5,"목표 투입비",bold=True,align="right"); C(ws,2,6,250000000,bold=True,fmt=NUM,color=GREEN)
C(ws,3,5,"마크업율(15~17.5%)",bold=True,align="right"); C(ws,3,6,0.15,bold=True,fmt="0.0%",color=GREEN)
# header
for c,h in enumerate(["구분","세부 항목","수량","단가","금액","비고"],1):
    C(ws,5,c,h,bold=True,color=WHITE,fill=BROWN,align="center" if c>2 else "left")
def sec(r,t):
    for c in range(1,7): C(ws,r,c,fill=BEIGE)
    C(ws,r,1,t,bold=True,fill=BEIGE)
def item(r,name,qty,price,note="",formula=True):
    C(ws,r,2,name); C(ws,r,3,qty,align="center")
    C(ws,r,4,price,fmt=NUM,align="right")
    if formula: C(ws,r,5,f"=C{r}*D{r}",fmt=NUM,align="right",bold=True)
    if note: C(ws,r,6,note,color="868680",size=9,wrap=True)

# 크리에이터 시딩
sec(6,"■ 크리에이터 시딩")
C(ws,7,1,"시딩")
item(7,"매크로 (인스타·틱톡)",15,800000,"인스타·틱톡 동일 견적")
item(8,"마이크로 (UGC 기준·인스타·틱톡)",150,500000,"UGC=인플루언서 시딩 기준 단가")
item(9,"나노",225,250000)
item(10,"KOL 무가시딩",60,100000,"제품 제공형")
item(11,"X (트위터) 시딩",15,500000)
item(12,"네이버 블로그",15,300000,"시딩 하단으로 위치 이동 · 단가 30만")
# 콘텐츠 제작
sec(13,"■ 콘텐츠 제작")
item(14,"EGC 콘텐츠 (UGC 마이크로 ×2)",12,1000000,"월 2개 × 6개월")
item(15,"Half EGC 콘텐츠 (UGC 마이크로 ×2.5)",6,1250000,"월 1개 × 6개월")
C(ws,16,2,"KV·디자인·영상 제작"); C(ws,16,3,"1식",align="center")
C(ws,16,4,"=E16",fmt=NUM,align="right")
C(ws,16,5,"=$F$2-(SUM(E7:E12)+E14+E15+SUM(E19:E22))",fmt=NUM,align="right",bold=True)
C(ws,16,6,"※ 목표 투입비(2.5억) 자동 정산(잔액). 추가비용은 제작비에서 흡수",color="868680",size=9,wrap=True)
# 바이럴·전환
sec(18,"■ 바이럴 · 전환")
item(19,"파워페이지 콘텐츠 바이럴",9,800000)
item(20,"커뮤니티 체험단 (탑)",3,5000000)
item(21,"챌린저스 (랭킹 견인)",2,5000000,"1회 500만 × 2회 (신규 반영)")
item(22,"어필리에이트 (올영 쇼핑 큐레이터)",60,100000,"건당 10만 (신규 반영)")
# 무상 제공
sec(23,"■ 무상 제공 (Value-add · 0원)")
freebies=["고성과 리포트 대시보드 (실시간 성과·랭킹 트래킹)",
          "올리비아 — AI 콘텐츠·리포팅 어시스턴트",
          "퍼포먼스 소재 제작 지원 (위닝 소재 PA 2차)",
          "VOC 딥다이브 리포트 (리뷰 크롤링·키워드 분석)",
          "Claude AI 심의 사전검수 (반려율↓·기간 단축)",
          "월간 성과 리포팅 (일·주·월간)",
          "크리에이터 풀 매칭·관리 / 위기 대응 모니터링"]
for i,fb in enumerate(freebies):
    r=24+i; C(ws,r,2,fb); C(ws,r,3,"-",align="center"); C(ws,r,4,"-",align="right")
    C(ws,r,5,0,fmt=NUM,align="right",bold=True,color=GREEN); C(ws,r,6,"무상 제공",color=GREEN,size=9)
# totals
tr=31
C(ws,tr,2,"투입비(집행) 소계",bold=True,fill=GRAYF,align="right")
for c in (1,3,4,6): C(ws,tr,c,fill=GRAYF)
C(ws,tr,5,"=SUM(E7:E12)+E14+E15+E16+SUM(E19:E22)",fmt=NUM,align="right",bold=True,fill=GRAYF)
C(ws,tr+1,2,"마크업 (수수료, 투입비×율)",bold=True,fill=PatternFill('solid',fgColor=LIME) and None,align="right")
C(ws,tr+1,2,"마크업 (수수료 = 투입비 × 마크업율)",bold=True,align="right",fill="FBFBE0")
for c in (1,3,4,6): C(ws,tr+1,c,fill="FBFBE0")
C(ws,tr+1,5,f"=E{tr}*$F$3",fmt=NUM,align="right",bold=True,fill="FBFBE0",color=GREEN)
C(ws,tr+2,2,"공급가액 (투입비 + 마크업)",bold=True,fill=GRAYF,align="right")
for c in (1,3,4,6): C(ws,tr+2,c,fill=GRAYF)
C(ws,tr+2,5,f"=E{tr}+E{tr+1}",fmt=NUM,align="right",bold=True,fill=GRAYF)
C(ws,tr+3,2,"부가가치세 (10%)",bold=True,fill=GRAYF,align="right")
for c in (1,3,4,6): C(ws,tr+3,c,fill=GRAYF)
C(ws,tr+3,5,f"=E{tr+2}*0.1",fmt=NUM,align="right",bold=True,fill=GRAYF)
C(ws,tr+4,2,"합계 금액 (VAT 포함)",bold=True,color=WHITE,fill=BROWN,align="right")
for c in (1,3,4,6): C(ws,tr+4,c,fill=BROWN)
C(ws,tr+4,5,f"=E{tr+2}*1.1",fmt=NUM,align="right",bold=True,color=WHITE,fill=BROWN)
C(ws,tr+6,4,"목표 대비 차액",bold=True,align="right"); C(ws,tr+6,5,f"=$F$2-E{tr}",fmt=NUM,align="right",color=GREEN)
C(ws,tr+6,6,"0 이면 투입비 2.5억 정확히 일치",color="868680",size=9)
for r in range(5,tr+5): ws.row_dimensions[r].height=21
ws.row_dimensions[16].height=30

# ===================== 월별 액션플랜 =====================
ms=wb.create_sheet("월별 액션플랜")
for i,w in enumerate([24,13,8,11,10,9,11,11,9,15]): ms.column_dimensions[get_column_letter(i+1)].width=w
C(ms,1,1,"월별 실행 타임라인 · 액션플랜 (투입비 연동)",bold=True,size=15)
C(ms,2,1,"단가/금액은 견적서와 동일 · 합계 금액 = 투입비 2.5억 (마크업·VAT 별도)",color="868680",size=10)
for c,h in enumerate(["항목","단가","7월","8월\n세일사전","9월\n세일★","10월","11월\n블프★","12월\n세일★","합계","금액"],1):
    C(ms,4,c,h,bold=True,color=WHITE,fill=BROWN,align="center",wrap=True)
ms.row_dimensions[4].height=34
def msec(r,t):
    for c in range(1,11): C(ms,r,c,fill=BEIGE)
    C(ms,r,1,t,bold=True,fill=BEIGE)
def mrow(r,name,unit,months,kv=False,note=None):
    C(ms,r,1,name)
    C(ms,r,2,"상시" if kv else unit, align="center" if kv else "right", fmt=None if kv else NUM)
    for i,v in enumerate(months): C(ms,r,3+i,v if v else None,align="center")
    if kv:
        C(ms,r,9,"상시",align="center"); C(ms,r,10,"=견적서!E16",fmt=NUM,align="right",bold=True)
    else:
        C(ms,r,9,f"=SUM(C{r}:H{r})",align="center",bold=True)
        C(ms,r,10,f"=I{r}*B{r}",fmt=NUM,align="right",bold=True)
msec(5,"■ 크리에이터 시딩")
mrow(6,"매크로",800000,[0,5,0,0,5,5])
mrow(7,"마이크로 (UGC)",500000,[15,35,20,15,30,35])
mrow(8,"나노",250000,[25,55,35,25,40,45])
mrow(9,"KOL 무가시딩",100000,[10,10,10,10,10,10])
mrow(10,"X (트위터) 시딩",500000,[0,5,0,0,5,5])
mrow(11,"네이버 블로그",300000,[0,5,0,0,5,5])
msec(12,"■ 콘텐츠 제작")
mrow(13,"EGC 콘텐츠",1000000,[2,2,2,2,2,2])
mrow(14,"Half EGC 콘텐츠",1250000,[1,1,1,1,1,1])
mrow(15,"KV·디자인·영상 (상시)",None,[None]*6,kv=True)
msec(16,"■ 바이럴 · 전환")
mrow(17,"파워페이지",800000,[0,3,0,0,3,3])
mrow(18,"커뮤니티 체험단",5000000,[0,1,0,0,1,1])
mrow(19,"챌린저스",5000000,[0,1,0,0,1,0])
# 어필리에이트 (누적표기, 금액 고정)
C(ms,20,1,"어필리에이트 (누적)"); C(ms,20,2,100000,fmt=NUM,align="right")
for i,v in enumerate([10,20,30,40,50,60]): C(ms,20,3+i,v,align="center")
C(ms,20,9,"누적 60",align="center",bold=True); C(ms,20,10,"=60*B20",fmt=NUM,align="right",bold=True)
# 합계행
C(ms,22,1,"월 시딩 총건수",bold=True,fill=GRAYF); C(ms,22,2,"",fill=GRAYF)
for i in range(6):
    col=get_column_letter(3+i); C(ms,22,3+i,f"=SUM({col}6:{col}11)",align="center",bold=True,fill=GRAYF)
C(ms,22,9,"=SUM(C22:H22)",align="center",bold=True,fill=GRAYF); C(ms,22,10,"",fill=GRAYF)
C(ms,23,9,"투입비 합계",bold=True,align="right")
C(ms,23,10,"=SUM(J6:J20)",fmt=NUM,align="right",bold=True,color=GREEN)
C(ms,24,9,"견적 투입비 일치?",align="right",size=9)
C(ms,24,10,'=IF(J23=견적서!E31,"일치","불일치")',align="right",size=9,color=GREEN)
for r in range(5,21): ms.row_dimensions[r].height=20

wb.save("proposal/올더베러_견적_액션플랜.xlsx")
print("saved")
