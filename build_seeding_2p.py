#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""시딩 전략 2p — 편집 가능 네이티브 PPT (1p 빌드업 피라미드/도화선, 2p 왜 시딩 1순위)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

CREAM=RGBColor(0xFC,0xFB,0xF0); BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34)
GRAY=RGBColor(0x78,0x78,0x72); MG=RGBColor(0xB0,0xB0,0xA8); WHITE=RGBColor(0xFF,0xFF,0xFF)
SOFT=RGBColor(0xF1,0xF0,0xE5); LINE=RGBColor(0x22,0x22,0x1F); LGT=RGBColor(0xEC,0xEC,0xDF)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def newslide():
    s=prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=CREAM
    return s
def rrect(s,x,y,w,h,fill,line=None,lw=1.5,rad=0.1,rounded=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded:
        try: sp.adjustments[0]=rad
        except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def tri(s,x,y,w,h,color,rot=0):
    sp=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.rotation=rot; sp.fill.solid(); sp.fill.fore_color.rgb=color; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def tbox(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,wrap=True):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=Pt(3); tf.margin_right=Pt(3); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(2)
        for (t,sz,b,c) in runs:
            r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def title(s,t):
    tbox(s,0.5,0.32,11.5,0.7,[[(t,26,True,BLACK)]])
    rrect(s,0.5,1.16,12.33,0.035,BLACK,rounded=False)

# ============ PAGE 1 — 빌드업 피라미드 / 도화선 ============
s=newslide()
title(s,"올더베러 캠페인 — 도화선에 불을 붙이는 일")
tbox(s,0.5,0.9,11.5,0.3,[[("Moonshot Rocket Launch · 토대조차 세워지지 않은 지금, 7월 First Shot으로 점화하여 단계적으로 쌓아 정상으로 발사합니다",11.5,False,GRAY)]])
cx=4.05
# apex
gw=4.7; rrect(s,cx-gw/2,1.5,gw,0.5,BLACK,rad=0.4)
tbox(s,cx-gw/2,1.56,gw,0.4,[[("▲  발사 — 올영세일 랭킹 1위 · 웰니스 대표 PB",12.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
tri(s,cx-0.16,2.05,0.32,0.22,BLACK,rot=180)
bands=[("BOOST","전환 증폭","메타 PA · 올영 인앱광고 · LMF 위닝소재 · 미디어커머스 · 어필리에이트 · CRM",3.1),
 ("PROVE","사회적 증거","후기·리뷰·해시태그 · 커뮤니티 체험단 · 파워페이지 · 챌린저스 랭킹 견인",4.4),
 ("SEED","콘텐츠·시딩 볼륨","나노/마이크로 시딩 · UGC·EGC·Half-EGC · 인스타/유튜브/틱톡/X/블로그",5.7),
 ("토대","브랜드 인지·각인","시그니처 키워드 「채움」 · KV · 상품단 메시지 체계 — 현재 미구축",7.0)]
y=2.34; bh=0.88; gap=0.1
for stage,role,tac,w in bands:
    x=cx-w/2
    rrect(s,x,y,w,bh,SOFT,LINE,1.5,0.08)
    tbox(s,x+0.22,y+0.1,w-1.0,0.6,[[(stage,15,True,BLACK),("   · "+role,11,True,GRAY)],[(tac,10,False,DARK)]])
    tbox(s,x+w-0.85,y+0.1,0.7,0.3,[[("과제",8.5,True,GRAY)]],PP_ALIGN.RIGHT)
    y+=bh+gap
# 도화선 base
bw=7.3; bx=cx-bw/2; by=y+0.04
rrect(s,bx,by,bw,0.6,BLACK,rad=0.1)
tri(s,bx+0.26,by+0.16,0.26,0.28,WHITE)
tbox(s,bx+0.62,by+0.06,bw-0.8,0.5,[[("도화선 점화 · First Shot (7월)",13,True,WHITE)],[("제품 경쟁력·올리브영 채널이라는 ‘연료’에 불을 붙입니다 — 위닝 메시지 점화",9.5,False,LGT)]])
# fuse line + flame + up arrow
rrect(s,0.45,2.1,0.05,by+0.3-2.1,BLACK,rounded=False)
tri(s,0.36,1.86,0.22,0.24,BLACK)
tri(s,0.36,by+0.18,0.22,0.26,BLACK)
tbox(s,0.08,3.9,0.35,2.0,[[("B",11,True,BLACK)],[("A",11,True,BLACK)],[("T",11,True,BLACK)],[("쌓",10,True,GRAY)],[("아",10,True,GRAY)],[("올",10,True,GRAY)],[("림",10,True,GRAY)]],PP_ALIGN.CENTER)
# story panel
px=8.05; pw=4.78
tbox(s,px,1.55,pw,0.3,[[("제안의 출발점",11.5,True,GRAY)]])
def note(y,h,head,body,dark=False):
    rrect(s,px,y,pw,h,(BLACK if dark else WHITE),(None if dark else LINE),2,0.08)
    tbox(s,px+0.25,y+0.14,pw-0.4,h-0.2,[[(head,12.5,True,(WHITE if dark else BLACK))]]+[[(ln,10.3,False,(LGT if dark else DARK))] for ln in body])
note(1.95,1.45,"01  현재 진단",["올더베러는 제품 경쟁력과 올리브영 채널","이라는 ‘연료’를 갖추었으나, 브랜드 토대","(인지·각인)는 아직 세워지지 않았습니다."])
note(3.55,1.45,"02  핵심 과제",["토대·SEED·PROVE·BOOST 전 단계가 비어","있으며, 이를 순차적으로 축적해 정상에","도달해야 합니다."])
note(5.15,1.7,"03  우리가 하는 일",["본 캠페인은 그 도화선에 불을 붙이는 일","입니다. 7월 First Shot으로 점화하여 단계","를 쌓고, 올영세일 시점에 정상으로 발사","합니다."],dark=True)
tbox(s,0.5,7.16,12.3,0.3,[[("※ 현재 활동 = 클라이언트 확인 기준 · 제품·채널 외 시딩·바이럴·퍼포·CRM 및 브랜드 토대 모두 미구축 → 본 제안의 실행 범위",9.5,False,GRAY)]])

# ============ PAGE 2 — 왜 시딩 1순위 ============
s=newslide()
title(s,"왜 인플루언서 시딩 콘텐츠가 1순위인가")
rrect(s,0.5,1.3,12.33,0.58,BLACK,rad=0.12)
tbox(s,0.5,1.3,12.33,0.58,[[("광고로 ‘알리기’ 전에, 먼저 ‘선택받는 증거’를 쌓습니다 — 그 출발점이 인플루언서 시딩입니다.",16,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
cards=[
 ("01","이름만으론 끌리지 않습니다","브랜드 검색 월 약 3,000 · 해시태그 #올더베러 +100여 건","네이버·구글 검색량 + 인스타그램 해시태그 집계",
  "다수 진성(나노·마이크로)이 동시다발로 노출해 ‘익숙함 → 신뢰’를 먼저 만듭니다. (나노 ER 4~6% 최고 · 브랜드 44%가 나노 선호)"),
 ("02","소비자는 사기 전에 검증합니다","‘구미 영양제’ 검색 → 올영 이동 → 후기·매장 확인 후 구매(검증형) · VOC 재구매율 15.5% · 구매이유 1위 ‘가격·할인’","리스닝마인드 검색 경로(path) 분석 + 올더베러 리뷰 1,104건 VOC 딥다이브(자체 크롤링)",
  "가격 외 ‘선택 근거’(후기·리뷰·해시태그·콘텐츠 볼륨)를 먼저 쌓아야 전환됩니다."),
 ("03","‘어떻게 먹는지’는 인플루언서가 보여줍니다","RFP 요구 ‘언제·왜·어떻게 먹는지 생활 맥락’ · VOC ‘효과 미체감’ 상위 단점 · 올리브오일 릴스는 ‘공복 루틴+성분’ 조합이 최고 참여","올리브영 RFP + VOC 딥다이브 + 인스타 릴스 콘텐츠 분석",
  "SNS 인플루언서가 실제 사용 상황(아침 공복·야근·취침 전)을 자기 일상으로 보여줘야 소비자가 ‘나도’라고 느낍니다 — 기능 나열로는 안 됩니다."),
 ("04","이미 올영 PB에서 고성과로 검증된 방식입니다","우리가 운영 중인 올영 PB 시딩+PA가 고성과: 브링그린 검색 9K→140K(15배) · 산리오 파트너십 ROAS 127→356→589% · 올영세일 ROAS 324→515%","BAT 올리브영 PB 운영 캠페인 실적(레퍼런스)",
  "동일한 ‘시딩 → 위닝 콘텐츠 → PA 2차’ 방법론을 올더베러에 그대로 적용합니다. (한정 2.5억 내 최고 효율)"),
]
cw=2.96; gap=0.12; x0=0.5; cy=2.26; ch=4.06; astrip=1.34
for i,(num,t,evd,src,ans) in enumerate(cards):
    x=x0+i*(cw+gap)
    rrect(s,x,cy,cw,ch,WHITE,LINE,2,0.06)
    tbox(s,x+0.22,cy+0.14,cw-0.4,0.4,[[(num,19,True,MG)]])
    tbox(s,x+0.22,cy+0.55,cw-0.4,0.6,[[(t,13,True,BLACK)]])
    tbox(s,x+0.22,cy+1.25,cw-0.4,2.0,[
        [("데이터 근거",9,True,GRAY)],[(evd,9.8,False,DARK)],
        [("",4,False,GRAY)],[("출처 ",8.5,True,GRAY),(src,8.8,False,GRAY)]])
    rrect(s,x,cy+ch-astrip,cw,astrip,SOFT,rounded=False)
    rrect(s,x,cy+ch-astrip,0.07,astrip,BLACK,rounded=False)
    tbox(s,x+0.22,cy+ch-astrip+0.08,cw-0.4,astrip-0.16,[[("→ 시딩이 1순위인 이유",9.5,True,BLACK)],[(ans,10.2,True,BLACK)]])
rrect(s,0.5,6.5,12.33,0.56,BLACK,rad=0.12)
tbox(s,0.5,6.5,12.33,0.56,[[("그래서 인플루언서 시딩이 1순위입니다 — Moonshot Rocket Launch의 ‘도화선’, 가장 먼저 불을 붙이는 일.",15,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
tbox(s,0.5,7.16,12.3,0.3,[[("※ 근거: 검색량·해시태그 집계 / 리스닝마인드 경로분석 / VOC 1,104건 딥다이브 / 인스타 릴스 분석 / BAT 올영 PB 운영 실적",9.5,False,GRAY)]])

prs.save("proposal/올더베러_시딩전략_2p.pptx")
print("slides:",len(prs.slides._sldIdLst))
