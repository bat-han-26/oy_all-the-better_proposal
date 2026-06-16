#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""바이럴 파트 5p — 편집 가능 네이티브 PPT
1 바이럴 역할(확신) / 2 챌린저스 / 3 커뮤니티·파워페이지 / 4 GEO·AEO / 5 주요 KPI"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

CREAM=RGBColor(0xFC,0xFB,0xF0); BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34)
GRAY=RGBColor(0x78,0x78,0x72); MG=RGBColor(0xB0,0xB0,0xA8); WHITE=RGBColor(0xFF,0xFF,0xFF)
SOFT=RGBColor(0xF1,0xF0,0xE5); LINE=RGBColor(0x22,0x22,0x1F); LGT=RGBColor(0xE6,0xE6,0xDD)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def newslide():
    s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=CREAM; return s
def rrect(s,x,y,w,h,fill,line=None,lw=1.5,rad=0.1,rounded=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded:
        try: sp.adjustments[0]=rad
        except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def chev(s,x,y,col=BLACK,sz=15):
    tbox(s,x,y,0.5,0.4,[[("▶",sz,True,col)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
def tbox(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(3); tf.margin_right=Pt(3); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(2)
        for (t,sz,b,c) in runs:
            r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def head(s,tag,t,sub):
    tbox(s,0.5,0.3,11.5,0.28,[[(tag,12.5,True,GRAY)]])
    tbox(s,0.5,0.58,12.5,0.6,[[(t,24,True,BLACK)]])
    tbox(s,0.5,1.18,12.4,0.5,[[(sub,12.5,True,DARK)]])
    rrect(s,0.5,1.66,12.33,0.03,BLACK,rounded=False)
def cell(c,t,sz=10,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.LEFT):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(6); c.margin_right=Pt(5); c.margin_top=Pt(2); c.margin_bottom=Pt(2); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; p.word_wrap=True
    r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

# ===== S1 (p42) — 바이럴 역할 = 구매 직전의 확신 =====
s=newslide()
head(s,"02_C 바이럴 파트의 역할","바이럴의 역할은 ‘입소문’이 아니라 구매 직전의 ‘확신’을 만드는 것입니다",
     "올더베러는 아직 브랜드명 검색·자발 탐색이 약합니다. 노출을 늘리기보다, 소비자가 제품을 본 뒤 검색·비교하는 순간에 충분한 후기와 추천 맥락이 보이도록 설계합니다.")
# 구매 여정 띠
flow=["제품을 본다","검색·비교한다","‘선택해도 된다’ 확신","구매·전환"]
fw=2.78;fg=0.36;fx=0.7;fy=1.86
for i,t in enumerate(flow):
    x=fx+i*(fw+fg)
    rrect(s,x,fy,fw,0.5,(BLACK if i==3 else SOFT),(None if i==3 else LINE),1.3,rad=0.5)
    tbox(s,x,fy+0.05,fw,0.36,[[(t,11.5,True,(WHITE if i==3 else BLACK))]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    if i<3: chev(s,x+fw+fg/2-0.25,fy+0.05)
# 3 cards
cards=[
 ("비교 시 ‘선택 근거’ 제공","후기·장바구니·랭킹 맥락을 만들어 “선택해도 되는 제품”이라는 확신을 형성합니다.",
  "VOC상 다수가 ‘검색→올영 이동→후기 확인 후 구매’(검증형). 구매이유 1위가 ‘가격·할인’인 만큼, 가격 외 선택 근거를 채웁니다.","→ 커뮤니티·파워페이지 (다음 장)"),
 ("구매 ‘전환’으로 연결","챌린저스·어필리에이트·쇼핑 큐레이터로 실제 구매와 리뷰·랭킹 견인까지 연결합니다.",
  "단순 노출이 아니라 진성 구매·리뷰·랭킹이라는 ‘행동’으로 전환. 위닝 콘텐츠는 PA 2차로 재투입합니다.","→ 챌린저스 (다음 장)"),
 ("검색 시 ‘비어 있지 않게’","제품명·카테고리·상황 키워드 검색 시 정보성 콘텐츠가 노출되도록 자산화합니다.",
  "검색·AI 답변 결과에 올더베러가 ‘근거 있는 추천’으로 등장하도록 GEO/AEO 자산을 축적합니다.","→ GEO·AEO (다음 장)"),
]
cw=3.97;gap=0.21;x0=0.5;cy=2.62;ch=3.7;nstrip=0.5
for i,(t,body,evd,link) in enumerate(cards):
    x=x0+i*(cw+gap)
    rrect(s,x,cy,cw,ch,WHITE,LINE,2,0.06)
    rrect(s,x+0.22,cy+0.22,0.5,0.5,BLACK,rad=0.16)
    tbox(s,x+0.22,cy+0.22,0.5,0.5,[[(str(i+1),18,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    tbox(s,x+0.84,cy+0.24,cw-1.0,0.5,[[(t,13.5,True,BLACK)]],anchor=MSO_ANCHOR.MIDDLE)
    tbox(s,x+0.24,cy+0.92,cw-0.46,1.0,[[(body,11,True,DARK)]])
    tbox(s,x+0.24,cy+2.0,cw-0.46,1.1,[[("근거  ",9,True,GRAY),(evd,9.8,False,DARK)]])
    rrect(s,x,cy+ch-nstrip,cw,nstrip,SOFT,rounded=False); rrect(s,x,cy+ch-nstrip,0.06,nstrip,BLACK,rounded=False)
    tbox(s,x+0.22,cy+ch-nstrip+0.07,cw-0.4,nstrip-0.1,[[(link,10,True,BLACK)]])
rrect(s,0.5,6.6,12.33,0.56,BLACK,rad=0.12)
tbox(s,0.5,6.6,12.33,0.56,[[("바이럴 = 노출량이 아니라, 구매 직전 ‘확신’의 밀도 — 플라이휠 ② 검증·신뢰 단계를 채웁니다.",14.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# ===== S2 (p43) — 챌린저스 =====
s=newslide()
head(s,"02_C 바이럴 ① 전환","챌린저스 — ‘실시간 랭킹 견인’을 위한 구매 인증형 바이럴 엔진",
     "유저가 제품가 일부를 부담하고 구매·사용·리뷰 미션을 수행하는 ‘인증 기반’ 챌린지 — 진성 구매자 리뷰를 동시 게시해 랭킹·검색 증거를 만듭니다.")
# 메커니즘 4단계
steps=[("01","챌린지 참여 모집","목표·습관 앱 챌린저스 내 미션 개설 (실구매 의사 유저 타겟)"),
 ("02","제품가 일부 부담 구매","유저가 직접 비용 일부 부담 → ‘진성 구매자’만 참여"),
 ("03","사용·리뷰 인증 미션","구매·사용·올영 리뷰·릴스 업로드를 ‘인증’으로 수행"),
 ("04","동시·분산 게시","피크 시간대에 분산 집중 게시 → 랭킹·인기검색 견인")]
tbox(s,0.5,1.78,11.0,0.3,[[("작동 메커니즘",13,True,BLACK)]])
bw=2.93;gap=0.16;x0=0.5;ey=2.12;hh=0.5;bh=1.34
for i,(num,t,desc) in enumerate(steps):
    x=x0+i*(bw+gap)
    rrect(s,x,ey,bw,hh,BLACK,rad=0.14)
    tbox(s,x+0.14,ey+0.05,0.5,0.4,[[(num,15,True,MG)]],anchor=MSO_ANCHOR.MIDDLE)
    tbox(s,x+0.62,ey+0.05,bw-0.7,0.4,[[(t,11,True,WHITE)]],anchor=MSO_ANCHOR.MIDDLE)
    rrect(s,x,ey+hh,bw,bh,WHITE,LINE,1.5,rad=0.06)
    tbox(s,x+0.16,ey+hh+0.12,bw-0.3,bh-0.2,[[(desc,10,False,DARK)]])
    if i<3: chev(s,x+bw+gap/2-0.25,ey+hh+0.4)
# 목표 + 시간대 전략
tbox(s,0.5,4.34,11.0,0.3,[[("목표 — 진성 구매자 리뷰 확보 + 동시 게시로 랭킹·검색 증거 형성",12.5,True,BLACK)]])
rrect(s,0.5,4.7,6.0,1.66,WHITE,LINE,2,0.06)
tbox(s,0.7,4.82,5.6,0.3,[[("랭킹 견인 공략법 — 피크 시간대 분산 집중",11.5,True,BLACK)]])
# timeline 10-12-14-16-18
times=["10시","12시","14시","16시","18시"]
twx=0.7;tww=0.86;twg=0.16;twy=5.2
for i,tm in enumerate(times):
    x=twx+i*(tww+twg)
    rrect(s,x,twy,tww,0.46,SOFT,LINE,1.2,rad=0.12)
    tbox(s,x,twy+0.05,tww,0.36,[[(tm,11,True,BLACK)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    if i<4: chev(s,x+tww+twg/2-0.22,twy+0.03,GRAY,12)
tbox(s,0.7,5.78,5.6,0.5,[[("진성 구매·리뷰 인증을 시간대별로 분산 집중 → 누적 리뷰·인기검색이 ",10,False,DARK),("저녁 피크에 랭킹 상위로 견인",10,True,BLACK),(". 적은 비용으로 고효율.",10,False,DARK)]])
rrect(s,6.66,4.7,6.17,1.66,SOFT,LINE,2,0.06)
tbox(s,6.84,4.82,5.8,0.3,[[("왜 효과적인가",11.5,True,BLACK)]])
tbox(s,6.84,5.16,5.85,1.2,[
 [("· 진성 구매자 자발 리뷰 — ‘어뷰징’이 아닌 실제 사용 후기로 신뢰·검색 자산 확보",10,False,DARK)],
 [("· 동시 게시 = 랭킹·인기검색어·후기 볼륨의 ‘동시 점화’ → 플라이휠 첫 바퀴 가속",10,False,DARK)],
 [("· 표시광고 가이드 준수(체험·인증 고지) · 진성성 우선",10,False,DARK)]])
tbox(s,0.5,6.5,12.3,0.3,[[("※ 챌린저스: 목표·습관 기반 인증 앱 — 제품가 일부 부담으로 실구매 의사 유저만 참여, 구매·리뷰·릴스 등 맞춤 인증 미션 운영(공식 자료 기반).",9,False,GRAY)]])

# ===== S3 (p46) — 커뮤니티·파워페이지 =====
s=newslide()
head(s,"02_C 바이럴 ② 검증","커뮤니티·파워페이지 — ‘추천템’ 맥락에 올더베러를 진입시키는 장치",
     "소비자는 세일 전후로 “무엇을 살지” 검색·비교합니다. 이때 올더베러가 추천 리스트와 장바구니 맥락 안에 자연스럽게 들어가야 합니다.")
two=[
 ("A","뷰티 정보 플랫폼 · 대형 커뮤니티 체험단","카테고리 관여가 높은 타겟에게 ‘검증된 후기’를 시딩",
  [("타겟","화해 등 뷰티앱·대형 커뮤니티의 카테고리 관여 高 사용자"),
   ("실행","체험단·어워드·랭킹 참여로 ‘검증된 후기’ 볼륨 확보"),
   ("근거","화장품 구매 시 64.2%가 화해 어워드·랭킹 참고 · 4단계 어뷰징 필터로 신뢰도 높음")]),
 ("B","파워페이지 (정보성 블로그·전문 채널)","정보성 검색 결과를 점유",
  [("타겟","‘올리브오일 효능’·‘웰니스 구미 추천’ 등 정보성 검색 사용자"),
   ("실행","정보성 파워 블로그·전문 채널로 검색 상위·추천 리스트 점유"),
   ("근거","세일 전 ‘무엇을 살지’ 탐색 트래픽을 선점 → 장바구니 진입 유도")]),
]
cw=6.06;gap=0.21;x0=0.5;cy=1.86;ch=4.62
for i,(tag,t,role,rows) in enumerate(two):
    x=x0+i*(cw+gap)
    rrect(s,x,cy,cw,ch,WHITE,LINE,2,0.05)
    rrect(s,x,cy,cw,0.86,BLACK,rad=0.05)
    rrect(s,x+0.22,cy+0.2,0.46,0.46,WHITE,rad=0.2)
    tbox(s,x+0.22,cy+0.2,0.46,0.46,[[(tag,16,True,BLACK)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    tbox(s,x+0.82,cy+0.14,cw-1.0,0.4,[[(t,13,True,WHITE)]])
    tbox(s,x+0.82,cy+0.5,cw-1.0,0.32,[[(role,10.5,True,MG)]])
    yy=cy+1.06
    for lab,val in rows:
        rrect(s,x+0.24,yy,1.0,0.34,SOFT,rad=0.2)
        tbox(s,x+0.24,yy+0.03,1.0,0.28,[[(lab,10,True,BLACK)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
        tbox(s,x+1.36,yy-0.04,cw-1.6,0.9,[[(val,10.5,(lab=="근거"),DARK)]])
        yy+=1.08
rrect(s,0.5,6.62,12.33,0.54,BLACK,rad=0.12)
tbox(s,0.5,6.62,12.33,0.54,[[("‘추천템’ 검색·비교 맥락을 선점 — 세일 전 탐색이 곧장 장바구니로 이어지게 만듭니다.",14.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# ===== S4 (p48) — GEO/AEO =====
s=newslide()
head(s,"02_C 바이럴 ③ 검색 자산","GEO/AEO — ‘검색 답변 시대’의 Pull 자산으로 설계",
     "이제 검색은 ‘AI가 어떤 정보를 골라 답하는가’의 싸움입니다. 언론보도·정보성 콘텐츠를 AI가 인용하기 좋은 구조로 발행해 검색·AI 답변을 선점합니다.")
# 정의 2칩
rrect(s,0.5,1.82,6.06,0.92,SOFT,LINE,1.5,0.06)
tbox(s,0.7,1.9,5.7,0.3,[[("GEO  ",12,True,BLACK),("생성형 엔진 최적화",11,True,GRAY)]])
tbox(s,0.7,2.2,5.7,0.5,[[("ChatGPT·Perplexity·Gemini가 답변을 만들 때 우리 콘텐츠를 ‘근거·출처’로 고르게 설계",10,False,DARK)]])
rrect(s,6.77,1.82,6.06,0.92,SOFT,LINE,1.5,0.06)
tbox(s,6.97,1.9,5.7,0.3,[[("AEO  ",12,True,BLACK),("답변 엔진 최적화",11,True,GRAY)]])
tbox(s,6.97,2.2,5.7,0.5,[[("구글 AI Overview·검색 요약 답변에 ‘정확한 출처’로 반영되도록 최적화",10,False,DARK)]])
# 5단계 플로우
tbox(s,0.5,2.94,11.0,0.3,[[("운영 플로우",13,True,BLACK)]])
flow=["키워드 설계","콘텐츠 발행\n(언론보도·정보성)","검색·AI 답변 점유","세일 전 탐색 대응","자연 유입 증가"]
fw=2.3;fg=0.13;fx=0.5;fy=3.3;fh=0.86
for i,t in enumerate(flow):
    x=fx+i*(fw+fg)
    last=(i==4)
    rrect(s,x,fy,fw,fh,(BLACK if last else WHITE),(None if last else LINE),2,0.1)
    lines=t.split("\n"); paras=[[(lines[0],11.5,True,(WHITE if last else BLACK))]]
    if len(lines)>1: paras.append([(lines[1],9,True,(MG if last else GRAY))])
    tbox(s,x,fy+0.1,fw,fh-0.16,paras,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    if i<4: chev(s,x+fw+fg/2-0.2,fy+0.22,GRAY,13)
# 원칙 + 왜 지금
rrect(s,0.5,4.5,6.06,1.86,WHITE,LINE,2,0.06)
tbox(s,0.7,4.6,5.7,0.3,[[("‘AI가 읽기 좋은’ 콘텐츠 원칙",12,True,BLACK)]])
tbox(s,0.7,4.96,5.75,1.36,[
 [("· 질문–답변(Q&A)형 구조로 작성 (검색 의도 직답)",10.5,False,DARK)],
 [("· 출처·수치·정의를 명시해 ‘인용 가능’하게",10.5,False,DARK)],
 [("· 언론보도 초안부터 AI 인용 형식으로 설계·발행",10.5,False,DARK)],
 [("· 발행 전 AI가 실제로 어떻게 읽는지 구성 점검 후 작성",10.5,False,DARK)]])
rrect(s,6.77,4.5,6.06,1.86,SOFT,LINE,2,0.06)
tbox(s,6.97,4.6,5.7,0.3,[[("왜 지금인가 (2026)",12,True,BLACK)]])
tbox(s,6.97,4.96,5.75,1.36,[
 [("· 일상 속 생성형 AI 사용 급증 · 구글 AI Overviews 확산",10.5,False,DARK)],
 [("· AI 추천 기반 쇼핑 기능까지 확대",10.5,False,DARK)],
 [("· ‘검색 결과 1위’보다 ‘AI 답변에 인용’이 새 경쟁력",10.5,True,BLACK)],
 [("· Push(광고)가 닿지 않는 탐색 순간을 Pull로 선점",10.5,False,DARK)]])
tbox(s,0.5,6.5,12.3,0.3,[[("※ GEO=생성형 AI 답변 출처 선택 최적화 / AEO=AI 요약 답변 출처 반영 최적화 (2026 검색 트렌드 자료 기반).",9,False,GRAY)]])

# ===== S5 (p72) — 주요 KPI =====
s=newslide()
head(s,"04 성과 관리","주요 KPI — 플라이휠 단계별 성과 지표",
     "‘노출’이 아니라 플라이휠이 도는지를 봅니다. 단계별 선행지표 → 최종 ‘올영세일 랭킹 1위’로 연결합니다.")
rows=[("플라이휠 단계","핵심 KPI","목표 (6개월)","측정·근거"),
 ("① 인지·시딩","UGC·시딩 콘텐츠 수량 · #올더베러 해시태그","월 대량 시딩 · 해시태그 100여 → 대폭 확대","인스타 해시태그·콘텐츠 집계"),
 ("② 검증·신뢰","리뷰 수·평균 별점 · 브랜드 검색량","브랜드 검색 3,000 → 10,000 · 리뷰 볼륨↑","네이버·구글 검색량 · 올영 리뷰"),
 ("③ 전환","올영 카테고리 랭킹 · ROAS","7월 중 카테고리 1위 · 올영세일 ROAS 개선","올영 랭킹 · 매체 리포트"),
 ("④ 재구매·확장","재구매율 · CRM(플친) 유입","캡슐 충성 유지(24.1%) · 구미 재구매 개선","VOC 딥다이브 · CRM"),
 ("종합","시그니처 키워드 노출률 · EMV·SoV","시그니처 키워드 노출률 90%+ · SoV 확대","검색·리스닝 · EMV 환산")]
ty=1.84; g=s.shapes.add_table(6,4,Inches(0.5),Inches(ty),Inches(12.33),Inches(3.0)).table
g.columns[0].width=Inches(2.2); g.columns[1].width=Inches(4.0); g.columns[2].width=Inches(3.65); g.columns[3].width=Inches(2.48)
for ri,row in enumerate(rows):
    for ci,val in enumerate(row):
        if ri==0: cell(g.cell(ri,ci),val,11,True,WHITE,BLACK)
        else:
            fill=(WHITE if ri%2 else RGBColor(0xF6,0xF5,0xEC))
            cell(g.cell(ri,ci),val,9.8,(ci in(0,2)),(BLACK if ci in(0,2) else DARK),fill)
for r in g.rows: r.height=Inches(0.5)
# ER 벤치마크 하단 패널
tbox(s,0.5,5.04,11.0,0.3,[[("인게이지먼트(ER) 벤치마크 — ‘진성도’ 가드 (티어별)",12,True,BLACK)]])
chips=[("나노 (1만 이하)","ER 5–15%"),("마이크로 (1만–5만)","ER 3–8%"),("매크로 (5만 이상)","ER 1–3%")]
cw=4.0;cg=0.16;cx=0.5;ky=5.4
for i,(t,v) in enumerate(chips):
    x=cx+i*(cw+cg)
    rrect(s,x,ky,cw,0.6,(BLACK if i==0 else WHITE),(None if i==0 else LINE),1.5,rad=0.1)
    tbox(s,x+0.18,ky+0.07,cw-0.3,0.46,[[(t+"   ",10.5,True,(WHITE if i==0 else BLACK)),(v,12,True,(MG if i==0 else BLACK))]],anchor=MSO_ANCHOR.MIDDLE)
rrect(s,0.5,6.16,12.33,0.5,SOFT,LINE,1.2,rad=0.1)
tbox(s,0.7,6.2,12.0,0.42,[[("ER = (좋아요+댓글+저장+공유) ÷ 도달 — 다수의 마이크로·나노로 ‘진성 ER’ 확보가 플라이휠을 굴리는 핵심 선행지표.",10.5,True,DARK)]],anchor=MSO_ANCHOR.MIDDLE)
tbox(s,0.5,6.74,12.3,0.3,[[("※ KPI 체계: 도달·SoV / 인게이지먼트 / 전환·ROAS / 진성도 / 브랜드 상기·EMV (2026 인플루언서 마케팅 벤치마크 + BAT 올영 PB 실적).",9,False,GRAY)]])

prs.save("proposal/올더베러_바이럴_5p.pptx")
print("slides:",len(prs.slides._sldIdLst))
