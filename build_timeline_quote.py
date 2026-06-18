# -*- coding: utf-8 -*-
"""월별 실행 타임라인 · 액션플랜 — 견적 포함 (1p).
고정 액션플랜 수치 기준. BATi = 시딩(크리에이터 시딩 그룹, 월 시딩 총건수에 포함)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

FONT="맑은 고딕"
INK=RGBColor(0x1A,0x1A,0x1A); WHITE=RGBColor(0xFF,0xFF,0xFF)
GBAND=RGBColor(0xEC,0xEF,0xE4)   # 그룹밴드 연그린그레이
GBORDER=RGBColor(0xD2,0xD2,0xCC)
BLACK=RGBColor(0x20,0x20,0x20)
PH1=RGBColor(0x33,0x33,0x33); PH2=RGBColor(0x9C,0x9C,0x97)
ACCENT=RGBColor(0x0E,0x5A,0x30)

prs=Presentation(); prs.slide_width=Inches(13.33); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])

def setfont(run,sz,col,bold=False):
    run.font.name=FONT; run.font.size=Pt(sz); run.font.bold=bold; run.font.color.rgb=col
    rPr=run._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
    if ea is None: ea=rPr.makeelement(qn('a:ea'),{}); rPr.append(ea)
    ea.set('typeface',FONT)

def tb(x,y,w,h,lines,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,wrap=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=wrap
    tf.vertical_anchor=anchor
    tf.margin_left=Pt(2);tf.margin_right=Pt(2);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0)
    first=True
    for spec in lines:
        t,sz,col,bold=spec[0],spec[1],spec[2],spec[3]
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.alignment=align; p.line_spacing=1.0
        r=p.add_run(); r.text=t; setfont(r,sz,col,bold)
    return box

def rect(x,y,w,h,fill=None,line=None,lw=0.75):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False
    return sp

def dashed_vline(x,y1,y2,color=RGBColor(0xB5,0xB5,0xB0),w=1.4):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y1),Inches(x),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w); c.shadow.inherit=False
    ln=c.line._get_or_add_ln(); pd=ln.makeelement(qn('a:prstDash'),{'val':'dash'}); ln.append(pd)

# ---------- geometry ----------
LX,LW=0.55,2.75
MX0,MW=3.30,1.30
SX,SW=11.10,1.68
def mx(i): return MX0+i*MW
months=["7월","8월","9월","10월","11월","12월"]
msub=["","세일 사전","세일 ★","","블프 ★","세일 ★"]

# ---------- title ----------
tb(0.55,0.28,12.2,0.6,[("월별 실행 타임라인 · 액션플랜  —  견적 포함",25,INK,True)],align=PP_ALIGN.LEFT)
rect(0.57,0.92,3.4,0.045,fill=ACCENT)

# ---------- phase bands ----------
py,ph=1.04,0.32
rect(MX0,py,MW*3,ph,fill=PH1)
tb(MX0,py,MW*3,ph,[("PHASE 1 · 빠른 반응 확보",11.5,WHITE,True)])
rect(MX0+MW*3,py,MW*3,ph,fill=PH2)
tb(MX0+MW*3,py,MW*3,ph,[("PHASE 2 · 대표성 확장",11.5,WHITE,True)])

# ---------- header row ----------
hy,hh=1.40,0.40
rect(LX,hy,LW,hh,fill=BLACK)
tb(LX,hy,LW,hh,[("항목  (건당 비용)",11,WHITE,True)],align=PP_ALIGN.LEFT)
for i in range(6):
    rect(mx(i),hy,MW,hh,fill=BLACK)
    if msub[i]:
        tb(mx(i),hy,MW,hh,[(months[i],11.5,WHITE,True),(msub[i],8,RGBColor(0xCF,0xCF,0xCB),False)])
    else:
        tb(mx(i),hy,MW,hh,[(months[i],11.5,WHITE,True)])
rect(SX,hy,SW,hh,fill=BLACK)
tb(SX,hy,SW,hh,[("합계",11.5,WHITE,True)])

# ---------- data ----------
SEED=[
 ("나노","270,000원",[10,20,12,10,18,20],90),
 ("마이크로 (UGC)","420,000원",[16,30,20,16,28,30],140),
 ("매크로","880,000원",[0,4,0,0,4,4],12),
 ("KOL 무가시딩","100,000원",[5,5,5,5,5,5],30),
 ("X (트위터) 시딩","550,000원",[0,8,0,0,8,8],24),
 ("네이버 블로그","220,000원",[0,5,0,0,5,5],15),
 ("BATi 마이크로 앰배서더","500,000원",[10,10,10,10,10,10],60),
]
CONTENT=[
 ("EGC 콘텐츠","500,000원",[20,20,20,20,20,20],120),
 ("Half EGC 콘텐츠","750,000원",[4,4,4,4,4,4],24),
 ("콘텐츠 제작 (KV·디자인·영상)","1,000,000원",None,"6 (상시)"),
]
VIRAL=[
 ("파워페이지","750,000원",[0,2,0,0,2,2],6),
 ("커뮤니티 체험단","4,300,000원",[0,1,0,0,1,0],2),
 ("챌린저스","4,800,000원",[0,1,0,0,1,0],2),
]
SOLUTION=[
 ("스프레이ai 솔루션","23,333원",None,"6 (상시)"),
]
SEED_TOTAL=[41,82,47,41,78,82]   # 월 시딩 총건수(BATi 포함)
BUD_PCT=["12%","23%","12%","12%","22%","19%"]
BUD_AMT=["2,894만","5,714만","3,116만","2,894만","5,576만","4,804만"]

def cell_gray(q,rmax):
    t=q/rmax if rmax else 0
    v=int(228-t*158)
    return RGBColor(v,v,v), (WHITE if v<150 else INK)

IRH=0.240; GRH=0.240; y=hy+hh
FULLW=LW+MW*6+SW

def group_band(label):
    global y
    rect(LX,y,FULLW,GRH,fill=GBAND,line=GBORDER,lw=0.5)
    tb(LX+0.08,y,FULLW-0.1,GRH,[(label,10.5,RGBColor(0x2C,0x3A,0x2C),True)],align=PP_ALIGN.LEFT)
    y+=GRH

def item_row(name,cost,dist,total):
    global y
    rect(LX,y,LW,IRH,fill=WHITE,line=GBORDER,lw=0.5)
    tbx=tb(LX+0.08,y,LW-0.1,IRH,[(name,9.5,INK,True),(cost,7.5,RGBColor(0x8A,0x8A,0x86),False)],align=PP_ALIGN.LEFT)
    if dist is None:
        for i in range(6):
            rect(mx(i),y,MW,IRH,fill=RGBColor(0x6E,0x6E,0x6A),line=GBORDER,lw=0.5)
            tb(mx(i),y,MW,IRH,[("●",10,WHITE,False)])
    else:
        rmax=max(dist) or 1
        for i in range(6):
            q=dist[i]
            if q>0:
                fill,tc=cell_gray(q,rmax)
                rect(mx(i),y,MW,IRH,fill=fill,line=GBORDER,lw=0.5)
                tb(mx(i),y,MW,IRH,[(str(q),10,tc,True)])
            else:
                rect(mx(i),y,MW,IRH,fill=RGBColor(0xFB,0xFB,0xFA),line=GBORDER,lw=0.5)
                tb(mx(i),y,MW,IRH,[("-",9,RGBColor(0xBD,0xBD,0xB8),False)])
    rect(SX,y,SW,IRH,fill=RGBColor(0xF2,0xF4,0xEC),line=GBORDER,lw=0.5)
    tb(SX,y,SW,IRH,[(str(total),10.5,INK,True)])
    y+=IRH

# 크리에이터 시딩 (BATi 포함)
group_band("■ 크리에이터 시딩    (BATi 마이크로 앰배서더 포함 · 모두 '시딩')")
for it in SEED: item_row(*it)
# 월 시딩 총건수
srh=0.28
rect(LX,y,LW,srh,fill=BLACK)
tb(LX,y,LW-0.1,srh,[("월 시딩 총건수",10.5,WHITE,True)],align=PP_ALIGN.RIGHT)
for i in range(6):
    rect(mx(i),y,MW,srh,fill=BLACK)
    tb(mx(i),y,MW,srh,[(str(SEED_TOTAL[i]),11,WHITE,True)])
rect(SX,y,SW,srh,fill=ACCENT)
tb(SX,y,SW,srh,[("371 건",11,WHITE,True)])
y+=srh

group_band("■ 콘텐츠")
for it in CONTENT: item_row(*it)
group_band("■ 바이럴 · 전환")
for it in VIRAL: item_row(*it)
group_band("■ 솔루션")
for it in SOLUTION: item_row(*it)

# 월 예산 비중 / 금액
bry=y+0.05; brh=0.44
rect(LX,bry,LW,brh,fill=BLACK)
tb(LX,bry,LW-0.1,brh,[("월 예산 비중 / 금액",11,WHITE,True)],align=PP_ALIGN.RIGHT)
for i in range(6):
    pv=int(BUD_PCT[i].replace("%",""))
    g=int(120-pv*1.4); g=max(20,g); fill=RGBColor(g,g,g)
    rect(mx(i),bry,MW,brh,fill=fill)
    tb(mx(i),bry,MW,brh,[(BUD_PCT[i],13,WHITE,True),(BUD_AMT[i],8.5,RGBColor(0xDA,0xDA,0xD6),False)])
rect(SX,bry,SW,brh,fill=ACCENT)
tb(SX,bry,SW,brh,[("100% · 2.5억",11.5,WHITE,True)])

# phase divider (between 9월 and 10월)
dashed_vline(MX0+MW*3, py, bry+brh)

# footer
tb(0.55,bry+brh+0.10,7,0.30,[("총 견적 250,000,000원 (VAT 별도)  ·  건당 비용은 집행 단가 기준  ·  © 2026 BAT",9.5,RGBColor(0x9A,0x9A,0x95),False)],align=PP_ALIGN.LEFT)

prs.save("proposal/올더베러_월별액션플랜_타임라인_1p.pptx")
print("saved", "rows end y=",round(bry+brh,2))
