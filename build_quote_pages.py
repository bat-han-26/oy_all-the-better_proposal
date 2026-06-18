# -*- coding: utf-8 -*-
"""고정된 견적상세(종합)에 맞춰 견적제안·월별 액션플랜만 재작성.
견적상세 시트는 절대 수정하지 않음. 마진 표기는 어디에도 넣지 않음. 마진 계산 탭 제거."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC="/root/.claude/uploads/bc09908a-8d16-5be3-806b-53899e2cede8/ee549bea-BAT_________260619____CGO_F.BF.xlsx"
OUT="proposal/올더베러_견적서_BAT.xlsx"

FONT="맑은 고딕"; ACC='_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
DARK="404040"; DGREEN="082A08"; LGRN="EAF3DA"; SUB="D9D9D9"; WHITE="FFFFFF"; INK="262626"; GREY="808080"
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=openpyxl.load_workbook(SRC)

# 마진 계산 탭 제거
for nm in list(wb.sheetnames):
    if nm.startswith("마진 계산"):
        del wb[nm]

def C(ws,r,c,v=None,b=False,col=INK,fill=None,al="center",fmt=None,sz=10,wrap=False,border=True,it=False):
    cell=ws.cell(r,c)
    if v is not None: cell.value=v
    cell.font=Font(name=FONT,size=sz,bold=b,color=col,italic=it)
    if fill: cell.fill=PatternFill("solid",fgColor=fill)
    cell.alignment=Alignment(horizontal=al,vertical="center",wrap_text=wrap)
    if border: cell.border=bd
    if fmt: cell.number_format=fmt
    return cell

def wipe(ws):
    for mc in list(ws.merged_cells.ranges): ws.unmerge_cells(str(mc))
    for row in ws.iter_rows():
        for c in row:
            c.value=None
            c.fill=PatternFill(fill_type=None); c.border=Border()
            c.alignment=Alignment(); c.font=Font(name=FONT,size=10)

# ============================================================ 견적제안
qp=wb["견적제안"]; wipe(qp); qp.sheet_view.showGridLines=False
for col,w in {1:2.5,2:2,3:16,4:2,5:118}.items(): qp.column_dimensions[get_column_letter(col)].width=w
def QL(r,label,val,head=False):
    if label is not None:
        C(qp,r,3,label,b=True,col=(WHITE if head else DGREEN),fill=(DARK if head else None),al="left",border=False)
    if val is not None:
        C(qp,r,5,val,al="left",wrap=True,sz=10,border=False)
qp.merge_cells("C3:E3"); C(qp,3,3,"BAT  |  견적 제안서",b=True,col=WHITE,fill=DGREEN,al="left",sz=15,border=False)
QL(5,"· 프로젝트명","올리브영 PB 올더베러 브랜드 빌딩 캠페인")
QL(7,"· 집행 기간","2026년 7월 1일 ~ 2026년 12월 31일 (6개월) · 인플루언서 시딩·바이럴 통합 운영")
QL(9,"· 집행 내용","① 인플루언서 시딩(나노·마이크로 코어) 중심 진성 UGC 대량 확산")
QL(10,None,"② BATi 마이크로 앰배서더(10명 × 6회) 다회차 협업으로 고퀄리티 콘텐츠 지속 발행")
QL(11,None,"③ 바이럴(챌린저스·커뮤니티·파워페이지) + EGC / Half EGC 신뢰 콘텐츠")
QL(12,None,"④ 위닝 콘텐츠·시딩 자산을 올영세일 구매 전환으로 직접 연결")
QL(14,"· 견적 요약","· 총 견적 : KRW 250,000,000원 (VAT 별도)")
QL(15,None,"· 목표 수량 : 크리에이터 시딩 311건 + BATi 60건 + EGC 120 · Half EGC 24 + 콘텐츠 제작 6식 + 바이럴 10 + 솔루션 6식")
QL(17,"· A/C 견적 상세","· 크리에이터 시딩 : 113,160,000원  (나노 90 · 마이크로 140 · 매크로 12 · KOL 30 · X 24 · 블로그 15)")
QL(18,None,"· BATi 마이크로 앰배서더 : 30,000,000원  (10명 × 6회 = 60건 × 500,000원)")
QL(19,None,"· 콘텐츠 EGC : 78,000,000원  (EGC 120 × 500,000 · Half EGC 24 × 750,000)")
QL(20,None,"· 콘텐츠 제작(KV·영상·이미지) : 6,000,000원  (6식 × 1,000,000)")
QL(21,None,"· 바이럴 : 22,700,000원  (파워페이지 6 · 커뮤니티 2 · 챌린저스 2)")
QL(22,None,"· 스프레이ai 솔루션 : 140,000원  (6식)")
QL(23,None,"· 총 합계 : 250,000,000원 (VAT 별도)")
QL(25,"· 단가 안내","· 나노 27만 / 마이크로 42만 / 매크로 88만 / KOL 10만 / X 55만 / 블로그 22만 (건당 총 견적)")
QL(26,None,"· EGC 50만 / Half EGC 75만 / 콘텐츠 제작 100만 / BATi 50만 (건당 총 견적)")
QL(28,"· BATi 안내","· BAT 보유 우수 마이크로 인플루언서 실(實)리드 풀 + 마이크로 MCN 모듈 기반 '마이크로 앰배서더' 6개월 다회차 협업")
QL(29,None,"· 1명당 월 1회 이상(총 6회) 협업 · e.g. 일본 마이크로 인플루언서 SK-II 다회차 협업 사례")
QL(31,"· 무상 지원","· 고성과 리포트 대시보드 / 올리비아 AI 어시스턴트 / VOC 딥다이브 / Claude AI 심의 검수 / 월간 리포팅 / 위닝 2차 활용")
QL(33,"· 특이사항","· 시장 단가 변동 및 모집 기간에 따라 총 원고료는 소폭 변동될 수 있습니다.")
QL(34,None,"· 인플루언서 섭외 & 콘텐츠 업로드 & 원고료 지급까지 BAT 담당 ※ 제품 발송은 올리브영 PB")
QL(35,None,"· 항목별 상세는 '견적상세'(월별), 집행 일정은 '월별 액션플랜' 시트를 확인 바랍니다.")

# ============================================================ 월별 액션플랜
ap=wb["월별 액션플랜"]; wipe(ap); ap.sheet_view.showGridLines=False
for col,w in {1:30,2:11,3:7,4:9,5:9,6:7,7:9,8:9,9:8,10:14}.items(): ap.column_dimensions[get_column_letter(col)].width=w
C(ap,1,1,"월별 실행 타임라인 · 액션플랜",b=True,col=DGREEN,sz=13,border=False)
aph=["항목","단가(총 견적)","7월","8월\n세일사전","9월\n세일★","10월","11월\n블프★","12월\n세일★","합계","금액"]
for i,h in enumerate(aph,1): C(ap,3,i,h,b=True,col=WHITE,fill=DARK,wrap=True)
# 단가 = 총견적건당(원고료+마크업) / 월별 분배 [7,8,9,10,11,12]
ROWS=[
("■ 크리에이터 시딩",None,None),
("나노",270000,[10,20,12,10,18,20]),
("마이크로 (UGC)",420000,[16,30,20,16,28,30]),
("매크로",880000,[0,4,0,0,4,4]),
("KOL 무가시딩",100000,[5,5,5,5,5,5]),
("X (트위터) 시딩",550000,[0,8,0,0,8,8]),
("네이버 블로그",220000,[0,5,0,0,5,5]),
("■ BATi 앰배서더",None,None),
("BATi 마이크로 앰배서더",500000,[10,10,10,10,10,10]),
("■ 콘텐츠",None,None),
("EGC 콘텐츠",500000,[20,20,20,20,20,20]),
("Half EGC 콘텐츠",750000,[4,4,4,4,4,4]),
("콘텐츠 제작 (KV·영상·이미지)",1000000,[1,1,1,1,1,1]),
("■ 바이럴 · 전환",None,None),
("파워페이지",750000,[0,2,0,0,2,2]),
("커뮤니티 체험단",4300000,[0,1,0,0,1,0]),
("챌린저스",4800000,[0,1,0,0,1,0]),
("■ 솔루션",None,None),
("스프레이ai 솔루션",23333.3,[1,1,1,1,1,1]),
]
r=4; data_rows=[]; seed_rows=[]
for label,price,dist in ROWS:
    if price is None:
        ap.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
        C(ap,r,1,label,b=True,col=DGREEN,fill=LGRN,al="left"); r+=1; continue
    C(ap,r,1,label,al="left",sz=9); C(ap,r,2,price,fmt=ACC,sz=9)
    for mi in range(6):
        q=dist[mi]; C(ap,r,3+mi,(q if q>0 else "-"))
    C(ap,r,9,f"=SUM(C{r}:H{r})",b=True)
    C(ap,r,10,f"=I{r}*B{r}",fmt=ACC,b=True)
    data_rows.append(r)
    if label in ("나노","마이크로 (UGC)","매크로","KOL 무가시딩","X (트위터) 시딩","네이버 블로그"): seed_rows.append(r)
    r+=1
# 월 시딩 총건수
C(ap,r,1,"월 시딩 총건수",b=True,fill=SUB,al="right")
for mi,col in enumerate(range(3,9)):
    cl=get_column_letter(col)
    C(ap,r,col,"="+"+".join(f"{cl}{rr}" for rr in seed_rows),b=True,fill=SUB)
C(ap,r,9,f"=SUM(C{r}:H{r})",b=True,fill=SUB); C(ap,r,2,"",fill=SUB); C(ap,r,10,"",fill=SUB)
r+=1
# 총 견적
ap.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9)
C(ap,r,1,"총 견적 (VAT 별도)",b=True,col=WHITE,fill=DARK,al="right")
C(ap,r,10,"="+"+".join(f"J{rr}" for rr in data_rows),fmt=ACC,b=True,col=WHITE,fill=DGREEN)

# 시트 순서: 견적제안, 견적상세, 월별 액션플랜
order=["견적제안","견적상세","월별 액션플랜"]
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)

wb.save(OUT)
print("saved ->", OUT, "| sheets:", wb.sheetnames)
