# -*- coding: utf-8 -*-
"""콘텐츠 유형 비교 + 마케팅 트렌드 변천사 2장."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from template_lib import *

prs=Presentation(); prs.slide_width=Inches(13.33); prs.slide_height=Inches(7.5)
BL=prs.slide_layouts[6]

def conc(s,text):
    rect(s,0.9,6.72,11.53,0.5,fill=BEIGE,rounded=False); rect(s,0.9,6.72,0.09,0.5,fill=GREEN)
    textbox(s,1.2,6.72,11.0,0.5,[(text,12.5,INK,True)],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
def footnote(s,t,y=7.02):
    textbox(s,0.9,y,11.53,0.3,[(t,8.5,GRAY,False)])
def imgbox(s,x,y,w,h,label):
    rect(s,x,y,w,h,fill=BEIGE,line=LGRAY,rounded=True)
    oval(s,x+w/2-0.3,y+h/2-0.42,0.6,fill=WHITE,line=GREEN,line_w=1.5)
    textbox(s,x+w/2-0.16,y+h/2-0.42,0.32,0.6,[("▶",16,GREEN,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    textbox(s,x,y+h/2+0.28,w,0.34,[(label,10.5,GRAY,False)],align=PP_ALIGN.CENTER)

# ============ S1: UGC vs EGC vs Half EGC ============
s=prs.slides.add_slide(BL)
content_header(s,"콘텐츠","CONTENT","UGC vs EGC vs Half EGC","콘텐츠 유형별 특징 · 제작기간 · 제작비 비교")
cols=[("UGC","리얼함 · 현장감","3~4주","50만원",""),
      ("EGC","미디어 커머스 (정보·전문성)","1주","100만원",""),
      ("Half EGC","소셜 에비던스 (소비자 신뢰 형성)","2주","150만원","제작 100만 + 출연 섭외 50만(기본)")]
w=3.66; xs=[0.9,4.83,8.76]
for (name,feat,term,cost,note),x in zip(cols,xs):
    hl = (name=="Half EGC")
    rect(s,x,2.4,w,4.02,fill=WHITE,line=LGRAY,rounded=True)
    rect(s,x,2.4,w,0.66,fill=GREEN)
    textbox(s,x,2.4,w,0.66,[(name,17,WHITE,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    rect(s,x+0.22,3.18,w-0.44,0.5,fill=BEIGE,rounded=True)
    textbox(s,x+0.22,3.18,w-0.44,0.5,[("+ "+feat,10.5,GREEN,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    imgbox(s,x+0.3,3.84,w-0.6,1.45,"영상 예시")
    # 제작기간 / 제작비
    textbox(s,x+0.3,5.42,1.6,0.3,[("제작기간",10,GRAY,True)])
    textbox(s,x+0.3,5.42,w-0.6,0.3,[(term,12,INK,True)],align=PP_ALIGN.RIGHT)
    rect(s,x+0.3,5.8,w-0.6,0.5,fill=BEIGE,rounded=True)
    textbox(s,x+0.5,5.8,1.5,0.5,[("제작비",10,GRAY,True)],anchor=MSO_ANCHOR.MIDDLE)
    textbox(s,x+0.3,5.8,w-0.7,0.5,[(cost,15,GREEN,True)],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
    if note: textbox(s,x+0.32,6.32,w-0.6,0.24,[(note,8,GRAY,False)],align=PP_ALIGN.CENTER)
footnote(s,"※ 제작비·제작기간은 콘텐츠 난이도·출연자 섭외에 따라 상이하며, 본 표는 예상 금액 기준입니다.")
# page num
textbox(s,13.33/2-0.5,7.04,1.0,0.3,[("1",10,LGRAY,False)],align=PP_ALIGN.CENTER)

# ============ S2: 마케팅 트렌드 변천사 ============
s=prs.slides.add_slide(BL)
content_header(s,"트렌드","TREND","리테일 마케팅 트렌드 변천사","레퍼런스 → 인플루언서 → Micro-IMC : 2026, 그 다음은?")
# Phase 1
x1,w1=0.9,3.2
rect(s,x1,2.55,w1,3.85,fill=WHITE,line=LGRAY,rounded=True)
rect(s,x1,2.55,w1,0.62,fill=GREEN)
textbox(s,x1,2.55,w1,0.62,[("STEP 1",11,YELLOW,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
textbox(s,x1+0.2,3.32,w1-0.4,0.5,[("레퍼런스 마케팅 Era",13.5,INK,True)],align=PP_ALIGN.CENTER)
textbox(s,x1+0.2,3.84,w1-0.4,0.34,[("자사몰 중심",10.5,GREEN,True)],align=PP_ALIGN.CENTER)
oval(s,x1+w1/2-0.45,4.35,0.9,fill=BEIGE,line=GREEN,line_w=1.5)
textbox(s,x1+w1/2-0.45,4.35,0.9,0.9,[("미디어\n커머스",11,GREEN,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
textbox(s,x1+0.25,5.45,w1-0.5,0.85,[("자사몰·리테일 레퍼런스 중심으로 성장하던 시기",9.8,GRAY,False)],align=PP_ALIGN.CENTER)
arrow(s,x1+w1+0.04,4.2,0.34)
# Phase 2
x2,w2=4.6,3.9
rect(s,x2,2.55,w2,3.85,fill=WHITE,line=LGRAY,rounded=True)
rect(s,x2,2.55,w2,0.62,fill=GREEN)
textbox(s,x2,2.55,w2,0.62,[("STEP 2",11,YELLOW,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
textbox(s,x2+0.2,3.32,w2-0.4,0.5,[("인플루언서 마케팅 Era",13.5,INK,True)],align=PP_ALIGN.CENTER)
textbox(s,x2+0.2,3.84,w2-0.4,0.34,[("B2C 광고·콘텐츠 그로스 중심",10.5,GREEN,True)],align=PP_ALIGN.CENTER)
maps=[("①","MCN","공동구매"),("②","UGC","어필리에이트"),("③","EGC","Half EGC")]
for i,(n,a,b) in enumerate(maps):
    yy=4.3+i*0.5
    textbox(s,x2+0.3,yy,0.4,0.34,[(n,11,GREEN,True)])
    textbox(s,x2+0.72,yy,1.4,0.34,[(a,11,INK,True)])
    textbox(s,x2+1.95,yy,0.4,0.34,[("→",11,GREEN_B,True)])
    textbox(s,x2+2.35,yy,1.4,0.34,[(b,11,INK,True)])
rect(s,x2+0.3,5.86,w2-0.6,0.42,fill=BEIGE,rounded=True)
textbox(s,x2+0.3,5.86,w2-0.6,0.42,[("협업 플랫폼 ·  Meta  ·  TikTok",10,INK,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
arrow(s,x2+w2+0.04,4.2,0.34)
# Phase 3 (highlight)
x3,w3=8.9,3.53
rect(s,x3,2.55,w3,3.85,fill=GREEN,rounded=True)
textbox(s,x3,2.7,w3,0.34,[("2026.06 · WHAT’S NEXT?",10.5,YELLOW,True,1)],align=PP_ALIGN.CENTER)
textbox(s,x3,3.06,w3,0.6,[("Micro-IMC",22,WHITE,True)],align=PP_ALIGN.CENTER)
# venn
oval(s,x3+0.5,3.75,1.35,fill=None,line=YELLOW,line_w=2.0)
oval(s,x3+1.55,3.75,1.35,fill=None,line=WHITE,line_w=2.0)
textbox(s,x3+0.4,4.18,1.2,0.4,[("인플루언서",9.5,WHITE,True)],align=PP_ALIGN.CENTER)
textbox(s,x3+1.9,4.18,1.2,0.4,[("퍼포먼스",9.5,WHITE,True)],align=PP_ALIGN.CENTER)
for i,b in enumerate(["+ 퍼포먼스 광고","+ 히어로컷","+ 미디어커머스 (EGC/UGC)"]):
    textbox(s,x3+0.34,5.4+i*0.33,w3-0.6,0.3,[(b,10,YELLOW,True)])
conc(s,"인플루언서 × 퍼포먼스의 교집합 — Micro-IMC가 올더베러의 다음 해답")
textbox(s,13.33/2-0.5,7.04,1.0,0.3,[("2",10,LGRAY,False)],align=PP_ALIGN.CENTER)

prs.save("proposal/올더베러_콘텐츠유형_트렌드_2p.pptx")
print("saved, slides:",len(prs.slides._sldIdLst))
