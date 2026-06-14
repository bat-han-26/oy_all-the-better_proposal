#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""병합 덱에 '크리에이터 리스트업' 1장 추가 — 역할×채널 다음, KOL 무가시딩 앞에 삽입"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DECK="proposal/올더베러_제안덱_병합본.pptx"
BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
F5=RGBColor(0xF5,0xF5,0xF5); EA=RGBColor(0xEA,0xEA,0xEA)
FONT="맑은 고딕"
prs=Presentation(DECK)
BLANK=None
for l in prs.slide_layouts:
    if l.name=="빈화면": BLANK=l
if BLANK is None: BLANK=prs.slide_layouts[5]
SW,SH=prs.slide_width,prs.slide_height

def rect(s,x,y,w,h,c):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=c; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(x,y,w,h).text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(3)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def title(s,t,sub=None):
    rect(s,0,0,SW,Inches(0.14),BLACK)
    tb=s.shapes.add_textbox(Inches(0.5),Inches(0.34),Inches(11.5),Inches(0.95)).text_frame; tb.word_wrap=True
    r=tb.paragraphs[0].add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(23); r.font.bold=True; r.font.color.rgb=BLACK
    if sub:
        r2=tb.add_paragraph().add_run(); r2.text=sub; r2.font.name=FONT; r2.font.size=Pt(12); r2.font.color.rgb=GRAY
    rect(s,Inches(0.52),Inches(1.4),Inches(1.2),Inches(0.05),BLACK)
def cell(c,t,sz=9.5,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.CENTER):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(5); c.margin_right=Pt(5); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; p.word_wrap=True
    r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

NEW=[]
def newslide():
    s=prs.slides.add_slide(BLANK); NEW.append(s); return s

# ============ 크리에이터 리스트업 ============
s=newslide()
title(s,"크리에이터 리스트업 — 선정 기준 × 예시 풀",
      "‘많이’가 아니라 ‘맞는’ 크리에이터 · 웰니스 버티컬별로 제품·앵글 매칭해 리스트업")

# 좌측: 선정 기준 4축
txt(s,Inches(0.5),Inches(1.6),Inches(4.0),Inches(0.3),[("선정 4기준",13.5,True,BLACK)])
crit=[
 ("① 브랜드 적합도","올더베러 타깃(30대 직장인·웰니스·부모) 라이프스타일 일치"),
 ("② 진성도","팔로워 대비 ER·진짜 댓글·저장 — 가짜 팔로워·구매 후기 NO"),
 ("③ 콘텐츠 결","협찬 티 안 나고 일상에 녹는 톤 · 제품 무드(‘채움’) 부합"),
 ("④ 심의 안전성","과장·효능 단정 이력 없음 · 식품 표현 가이드 준수 가능"),
]
y=2.0
for h,d in crit:
    rect(s,Inches(0.5),Emu(int(Inches(y))),Inches(4.0),Inches(1.02),F5)
    rect(s,Inches(0.5),Emu(int(Inches(y))),Inches(0.08),Inches(1.02),BLACK)
    txt(s,Inches(0.68),Emu(int(Inches(y))+int(Inches(0.08))),Inches(3.7),Inches(0.9),
        [(h,11.5,True,BLACK),(d,9.5,False,GRAY)])
    y+=1.12

# 우측: 예시 풀 table (웰니스 버티컬 × 채널 × 등급 × 제품·앵글)
txt(s,Inches(4.85),Inches(1.6),Inches(8),Inches(0.3),[("웰니스 버티컬별 예시 풀 (매칭 가이드)",13.5,True,BLACK)])
rows=[
 ["홈트 · 피트니스","인스타·유튜브","마이크로·나노","구미(루테인·올인원) — 운동 루틴 속 챙김"],
 ["요가 · 필라테스","인스타","마이크로·나노","멜라나잇·콜라겐 — 자기관리 무드 리추얼"],
 ["러닝크루","인스타·틱톡","나노","스틱·캡슐 — 휴대·아침 에너지 채움"],
 ["헬시 레시피 · 식단","인스타·유튜브","마이크로","올리브오일 — 요리·식단 속 자연 노출"],
 ["다이어트 · 바디프로필","인스타·틱톡","마이크로·나노","구미·스틱 — 식단 관리 간식 대체"],
 ["직장인 루틴 브이로그","인스타·유튜브","마이크로","캡슐·구미 — 아침/점심 루틴, 직장인 공감"],
 ["헬시플레저 라이프","인스타·X","매크로·마이크로","전 라인 — ‘건강하게 즐기는’ 무드 점화"],
]
t=s.shapes.add_table(len(rows)+1,4,Inches(4.85),Inches(2.0),Inches(8.0),Inches(4.0)).table
t.columns[0].width=Inches(2.0); t.columns[1].width=Inches(1.5); t.columns[2].width=Inches(1.5); t.columns[3].width=Inches(3.0)
for c,h in enumerate(["웰니스 버티컬","주력 채널","적합 등급","매칭 제품 · 앵글"]):
    cell(t.cell(0,c),h,9.5,True,WHITE,DARK,PP_ALIGN.LEFT if c==3 else PP_ALIGN.CENTER)
for ri,row in enumerate(rows,1):
    fill=WHITE if ri%2 else F5
    cell(t.cell(ri,0),row[0],9.5,True,BLACK,fill,PP_ALIGN.LEFT)
    cell(t.cell(ri,1),row[1],9,False,DARK,fill)
    cell(t.cell(ri,2),row[2],9,False,GRAY,fill)
    cell(t.cell(ri,3),row[3],9,False,BLACK,fill,PP_ALIGN.LEFT)
for r in t.rows: r.height=Inches(0.5)

# 하단 프로세스 밴드
rect(s,Inches(0.5),Inches(6.55),Inches(12.35),Inches(0.62),BLACK)
txt(s,Inches(0.7),Inches(6.64),Inches(12),Inches(0.46),
    [("리스트업 프로세스 — ① 버티컬·등급별 롱리스트 → ② 진성·심의 스크리닝 → ③ 숏리스트 컨택 → ④ 시즌별 배분(세일엔 매크로 가중)",11,True,WHITE)],
    anchor=MSO_ANCHOR.MIDDLE)

# ---- 삽입: 'KOL 무가시딩 운영 전략' 슬라이드 앞 ----
n_before=len(prs.slides._sldIdLst)-len(NEW)
target=n_before
for i,sl in enumerate(list(prs.slides)[:n_before]):
    txtall=" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    if "KOL 무가시딩 운영 전략" in txtall:
        target=i; break

sldIdLst=prs.slides._sldIdLst; ids=list(sldIdLst); moved=ids[-len(NEW):]
for el in moved: sldIdLst.remove(el)
for off,el in enumerate(moved): sldIdLst.insert(target+off,el)
prs.save(DECK)
print("total:",len(prs.slides._sldIdLst),"| added:",len(NEW),"| inserted at idx",target)
