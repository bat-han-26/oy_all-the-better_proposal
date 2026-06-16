#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""심의 준수 체계 + Claude AI 사전검수 — 편집 가능 네이티브 PPT (2p)"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

CREAM=RGBColor(0xFC,0xFB,0xF0); BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34)
GRAY=RGBColor(0x78,0x78,0x72); MG=RGBColor(0xB0,0xB0,0xA8); WHITE=RGBColor(0xFF,0xFF,0xFF)
SOFT=RGBColor(0xF1,0xF0,0xE5); LINE=RGBColor(0x22,0x22,0x1F); LGT=RGBColor(0xE6,0xE6,0xDD)
RED=RGBColor(0xB0,0x3A,0x2E); GRN=RGBColor(0x2E,0x5A,0x3A)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def newslide():
    s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=CREAM; return s
def rrect(s,x,y,w,h,fill,line=None,lw=1.5,rad=0.1,rounded=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded:
        try: sp.adjustments[0]=rad
        except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def chev(s,x,y,col=BLACK):
    tbox(s,x,y,0.5,0.4,[[("▶",16,True,col)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
def tbox(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(3); tf.margin_right=Pt(3); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(2)
        for (t,sz,b,c) in runs:
            r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def head(s,tag,t,sub):
    tbox(s,0.5,0.3,11.5,0.28,[[(tag,12.5,True,GRAY)]])
    tbox(s,0.5,0.58,12.3,0.6,[[(t,26,True,BLACK)]])
    tbox(s,0.5,1.2,12.3,0.34,[[(sub,13,True,DARK)]])
    rrect(s,0.5,1.6,12.33,0.03,BLACK,rounded=False)
def cell(c,t,sz=10,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.LEFT):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(6); c.margin_right=Pt(5); c.margin_top=Pt(2); c.margin_bottom=Pt(2); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; p.word_wrap=True
    r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

# ===== SLIDE 1 — 준수 체계 + AI 사전검수 =====
s=newslide()
head(s,"03 심의·표시광고 준수 체계","심의 리스크는 ‘앞단’에서 차단합니다",
     "BAT는 식품·건기식 심의 가이드를 준수하고, 송출 前 Claude AI로 1차 검수합니다.")
# 배너
rrect(s,0.5,1.78,12.33,0.56,BLACK,rad=0.1)
tbox(s,0.5,1.78,12.33,0.56,[[("콘텐츠 송출 前, Claude AI가 먼저 읽습니다 — 사람이 놓치는 위험 표현을 앞단에서 거릅니다.",15.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# AI 사전검수 4단계 플로우
tbox(s,0.5,2.56,11.0,0.3,[[("Claude AI 사전 검수 워크플로우",13,True,BLACK)]])
steps=[("01","대본·자막·캡션 입력","릴스/쇼츠 대본, 자막 타임코드, 캡션·해시태그를 검수 큐에 등록"),
 ("02","Claude AI 1차 스크리닝","제품 분류(일반식품/건기식) 자동 분기 → 효능·과장·질병 표현 자동 탐지·플래그"),
 ("03","안전 표현 치환 제안","분류별 규칙 기반으로 위험 문구 → 안전 문구 대안 제시 (근거 조항 함께)"),
 ("04","법무·RA 최종 검토 · 송출","건기식은 자율심의·법무 확인까지 완료 후 송출")]
bw=2.93; gap=0.16; x0=0.5; ey=2.92; hh=0.5; bh=1.42
for i,(num,t,desc) in enumerate(steps):
    x=x0+i*(bw+gap)
    rrect(s,x,ey,bw,hh,BLACK,rad=0.14)
    tbox(s,x+0.14,ey+0.05,0.5,0.4,[[(num,15,True,MG)]],anchor=MSO_ANCHOR.MIDDLE)
    tbox(s,x+0.62,ey+0.05,bw-0.7,0.4,[[(t,11.5,True,WHITE)]],anchor=MSO_ANCHOR.MIDDLE)
    rrect(s,x,ey+hh,bw,bh,WHITE,LINE,1.5,rad=0.06)
    tbox(s,x+0.16,ey+hh+0.12,bw-0.3,bh-0.2,[[(desc,10.5,False,DARK)]])
    if i<3: chev(s,x+bw+gap/2-0.25,ey+hh+0.45)

# 효과 chips
chips=["심의 반려율 ↓","검토 기간 단축","표현 일관성 확보","분류별 자동 분기"]
cw=2.93;cg=0.16;cx=0.5;ky=5.04
for i,c in enumerate(chips):
    x=cx+i*(cw+cg)
    rrect(s,x,ky,cw,0.42,SOFT,LINE,1.3,rad=0.5)
    tbox(s,x,ky+0.05,cw,0.32,[[("✓  "+c,11,True,BLACK)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# 준수 프레임 — 2 분류
tbox(s,0.5,5.66,11.0,0.3,[[("제품 분류별 심의 원칙",13,True,BLACK)]])
rrect(s,0.5,5.98,6.06,1.06,WHITE,LINE,1.5,rad=0.06)
tbox(s,0.68,6.06,5.8,0.3,[[("일반식품  ",11,True,BLACK),("올리브오일 · 멜라나잇 · 콜라겐 구미",9.5,True,GRAY)]])
tbox(s,0.68,6.36,5.8,0.64,[[("질병 예방·치료, 신체조직·기능 효능 표현 ",9.8,False,DARK),("전면 금지",9.8,True,RED),(" → 원료·맛·편의·무드 앵글만",9.8,False,DARK)]])
rrect(s,6.77,5.98,6.06,1.06,WHITE,LINE,1.5,rad=0.06)
tbox(s,6.95,6.06,5.8,0.3,[[("건강기능식품  ",11,True,BLACK),("루테인 · 바나바잎 구미",9.5,True,GRAY)]])
tbox(s,6.95,6.36,5.8,0.64,[[("인정 기능성 문구 ",9.8,False,DARK),("그대로만",9.8,True,GRN),(" · 의약품 오인·과장·체험 단정 금지 · 자율심의 대상",9.8,False,DARK)]])
tbox(s,0.5,7.12,12.3,0.3,[[("근거: 식품 등의 표시·광고에 관한 법률 · 한국건강기능식품협회 자율심의 · 공정위 추천·보증 심사지침(유료광고·협찬 표시)",9,False,GRAY)]])

# ===== SLIDE 2 — 표현 치환표 + 체크리스트 =====
s=newslide()
head(s,"03 심의·표시광고 준수 체계","표현 치환표 & 송출 전 체크리스트",
     "위험 표현(❌)을 안전 표현(✅)으로 — Claude AI가 1차로, 사람이 최종으로 검수합니다.")
# 치환표
rows=[("제품 (분류)","❌ 위험 표현","✅ 안전 표현"),
 ("올리브오일 캡슐·스틱 (일반)","혈관·콜레스테롤에 좋은","엑스트라 버진 올리브오일을 한 포로"),
 ("멜라나잇 구미 (일반)","숙면 젤리 / 잠 잘 오는","자기 전 리추얼 / 하루 마무리 한 알"),
 ("콜라겐 구미 (일반)","피부·주름·탄력 콜라겐","디저트 대신 챙기는 젤리"),
 ("루테인 구미 (건기식)","눈이 좋아지는 / 덜 침침해진","루테인, 눈 건강에 도움을 줄 수 있음"),
 ("바나바잎 구미 (건기식)","혈당 낮추는 / 당뇨에 좋은","식후 혈당 상승 억제에 도움을 줄 수 있음")]
ty=1.78; g=s.shapes.add_table(6,3,Inches(0.5),Inches(ty),Inches(12.33),Inches(2.7)).table
g.columns[0].width=Inches(3.5); g.columns[1].width=Inches(4.4); g.columns[2].width=Inches(4.43)
for ri,row in enumerate(rows):
    for ci,val in enumerate(row):
        if ri==0: cell(g.cell(ri,ci),val,11,True,WHITE,BLACK)
        else:
            fill=(WHITE if ri%2 else RGBColor(0xF6,0xF5,0xEC))
            col=(BLACK if ci==0 else (RED if ci==1 else GRN))
            cell(g.cell(ri,ci),val,10,(ci!=1),col,fill)
for r in g.rows: r.height=Inches(0.45)

# 체크리스트
tbox(s,0.5,4.66,11.0,0.3,[[("송출 전 필수 체크리스트 (인스타그램 릴스 기준)",13,True,BLACK)]])
checks=[
 "유료광고·협찬 표시 — 자막 + 음성 + 캡션 + 인스타 ‘유료 파트너십’ 태그 모두 적용",
 "일반식품(올리브오일·멜라나잇·콜라겐) — 질병·신체기능·건강 효능 표현 전면 배제 확인",
 "건기식(루테인·바나바잎) — 인정 기능성 문구 그대로, 의약품 오인·과장·체험 단정 없음",
 "‘효과 미체감’ 방어 — ‘꾸준히’ 표현 OK, 효과 보장·단정 금지",
 "건기식 컨셉은 송출 前 한국건강기능식품협회 자율심의 + 법무·RA 검토 완료",
]
cy=4.98
for i,c in enumerate(checks):
    rrect(s,0.5,cy,12.33,0.4,(SOFT if i%2 else WHITE),LINE,1.2,rad=0.1)
    tbox(s,0.66,cy+0.05,0.4,0.3,[[("☑",13,True,BLACK)]])
    tbox(s,1.06,cy+0.05,11.6,0.32,[[(c,11,True,DARK)]],anchor=MSO_ANCHOR.MIDDLE)
    cy+=0.46
# AI footer bar
rrect(s,0.5,cy+0.04,12.33,0.42,BLACK,rad=0.1)
tbox(s,0.5,cy+0.04,12.33,0.42,[[("→ Claude AI 사전 검수로 1차 필터링(반려↓·기간 단축) 후, 위 체크리스트로 사람이 최종 점검합니다.",11.5,True,WHITE)]],PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

prs.save("proposal/올더베러_심의준수_AI검수_2p.pptx")
print("slides:",len(prs.slides._sldIdLst))
