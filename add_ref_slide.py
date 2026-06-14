#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""페이지 11(올영 1등 케이스)에 역량별 레퍼런스 카드를 네이티브 도형으로 추가"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DECK="proposal/올더베러_제안덱_병합본.pptx"
BLACK=RGBColor(0x14,0x14,0x12); DARK=RGBColor(0x37,0x37,0x34); GRAY=RGBColor(0x78,0x78,0x72)
MG=RGBColor(0xB2,0xB2,0xAA); WHITE=RGBColor(0xFF,0xFF,0xFF); SOFT=RGBColor(0xF3,0xF2,0xE7); LINE=RGBColor(0x22,0x22,0x1F)
FONT="맑은 고딕"
prs=Presentation(DECK)
s=list(prs.slides)[10]  # page 11

def rrect(x,y,w,h,fill,line=None,lw=1.5,rad=0.10,rounded=True):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded:
        try: shp.adjustments[0]=rad
        except: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    shp.shadow.inherit=False; return shp
def tbox(x,y,w,h,paras,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(2); tf.margin_right=Pt(2); tf.margin_top=Pt(1); tf.margin_bottom=Pt(1)
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(2)
        for (t,sz,b,c) in runs:
            r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf

# subtitle under existing title
tbox(0.5,1.5,12.4,0.3,[[("우리가 올리브영에서 직접 만든 성과 — 매체·데이터·크리에이티브·시딩 네 역량 모두 올더베러에 그대로 적용",11.5,False,GRAY)]])

cols=[
 ("매체 · 퍼포먼스","올영세일 운영·타겟·IP 매체로 ROAS",[
   ("올영세일 메타 CPM 위기 대응","세일 경쟁으로 특정 세트 CPM 210,158원 폭등","고CPM OFF→구매잠재(기프트카드·쿠폰)·어드밴티지 재구성","CPM 210,158→12,033원 · ROAS 324→515%",None),
   ("올영세일 사전 캠페인 목표 검증","사전 장바구니가 본행사 매출로 이어지는지 미검증","장바구니 vs 구매 목표 동시 운영·연계 분석","본행사 장바구니 ROAS 1,180% vs 구매 699%",None),
   ("IP·모델 맞춤 매체 발굴 (브링그린·컬러그램)","팬덤 소구 필요한데 범용 매체 제안","X 커스텀라이크+GFA / 인스티즈 신규 발굴","검색 9K→140K(15배) · 인스티즈 노출 733만",None),
 ]),
 ("데이터","리뷰·VOC를 직접 수집·분석",[
   ("올리브영 VOC 리뷰 크롤링 리포트","기존 리스닝 툴로 구매자 리뷰 맥락·진정성 파악 한계","파이썬 크롤러 자체 개발 + R 빈도·키워드 분석, ‘상황·목적’ 중심 분류 → 검색·SNS·리뷰 연계 종합 리포트","올영 타 브랜드 추가 리포트 요청 연쇄","올더베러 VOC 딥다이브 대시보드(히어로 3종·리뷰 1,104건)로 이미 구현"),
 ]),
 ("크리에이티브","감도 × 성과를 잇는 소재 시스템",[
   ("‘써먹히는’ 소재 가이드 (브링그린·웨이크메이크)","감도 vs 성과 기준 모호 → 컨펌 지연·QC 혼선","목적 중심 분류(USP형/프로모션형)+비주얼 키워드화(CLEAN·COOL·DYNAMIC)","빅&스몰웨이브 실행 구조로 정착·내부 공유",None),
   ("웨이크메이크 산리오 콜라보 스프레이AI","3~40개 브랜드 동시 콜라보·색조 경쟁","맞춤 크리에이터+스프레이AI·메타 파트너십·틱톡 미러링","파트너십 ROAS 127→356→589% · CPC -38%",None),
 ]),
 ("인플루언서 · 시딩","맞춤 크리에이터·시딩으로 바이럴·전환",[
   ("바이오힐 보(US) 시딩 콘텐츠 기획","B&A 컷 확보 불가(인플 대부분 20대)·짧은 일정 내 바이럴","‘노인 필터’ 역발상+감정 트리거 자막, 나노 인플·실시간 QC·부정댓글 모니터링","틱톡 75만 뷰 · CPV 3원(업계 1/5) · 공유 601건","올더베러 마이크로/나노 시딩 앵글·역발상 훅 설계에 적용"),
 ]),
]
cw=2.98; gap=0.13; x0=0.5; hy=1.86; hh=0.6; cardtop=2.62; bottom=7.04; area=bottom-cardtop
strip=0.5
for ci,(pillar,role,cards) in enumerate(cols):
    x=x0+ci*(cw+gap)
    rrect(x,hy,cw,hh,BLACK,rad=0.12)
    tbox(x+0.16,hy+0.07,cw-0.3,0.5,[[(pillar,13,True,WHITE)],[(role,8.5,True,MG)]])
    n=len(cards); cgap=0.12; chh=(area-(n-1)*cgap)/n
    for j,(title,prob,method,result,extra) in enumerate(cards):
        cy=cardtop+j*(chh+cgap)
        rrect(x,cy,cw,chh,WHITE,LINE,1.5,rad=0.07)
        compact = chh<1.65
        if compact:
            tbox(x+0.16,cy+0.1,cw-0.32,chh-strip-0.1,
                 [[(title,11,True,BLACK)],[(method,8.6,False,GRAY)]])
        else:
            paras=[[(title,11.5,True,BLACK)],
                   [("문제  ",8.5,True,GRAY),(prob,9.3,False,DARK)],
                   [("실행  ",8.5,True,GRAY),(method,9.3,False,DARK)]]
            if extra and chh>=3.0:
                paras.append([("올더베러 적용  ",8.5,True,GRAY),(extra,9.3,True,BLACK)])
            tbox(x+0.16,cy+0.1,cw-0.32,chh-strip-0.12,paras)
        # 성과 strip
        rrect(x,cy+chh-strip,cw,strip,SOFT,rounded=False)
        rrect(x,cy+chh-strip,0.07,strip,BLACK,rounded=False)
        tbox(x+0.16,cy+chh-strip+0.04,cw-0.3,strip-0.06,
             [[("성과  ",8.5,True,GRAY),(result,9.8,True,BLACK)]])

prs.save(DECK)
print("page11 shapes:",len(list(prs.slides)[10].shapes))
