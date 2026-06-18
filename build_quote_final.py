# -*- coding: utf-8 -*-
"""올더베러 견적서 최종본 — BATi(마이크로 앰배서더) 반영 / 어필리에이트 제거 / 총 2.5억.
4개 탭(견적제안·견적상세·월별 액션플랜·마진 계산)을 단일 마스터 수치로 일관 생성."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
ACC = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
PCT = '0.0%'
DARK="404040"; OLIVE="62A81B"; DGREEN="082A08"; GREEN_HL="82DC28"
LGRN="EAF3DA"; SUB="D9D9D9"; WHITE="FFFFFF"; INK="262626"; GREY="808080"
thin=Side(style="thin",color="BFBFBF")
bd=Border(left=thin,right=thin,top=thin,bottom=thin)

# ============================================================== MASTER DATA
# 항목: (채널, 페르소나, 해시태그, 정가, 할인가, 건당실비, 수량, 비고)
ITEMS = [
("나노","IG","웰니스 입문형 2030","#올더베러 #오늘도베러 #웰니스입문",275000,225000,150000,160,""),
("마이크로 (UGC)","IG·TT","직장인 루틴형 2535","#올더베러 #오늘도베러",450000,350000,300000,180,"기대행 수수료 25%에서 5%p 추가 할인"),
("매크로","IG·TT·YT","헬시플레저 광역 4050+","#올더베러 #올영세일",1000000,800000,500000,12,""),
("KOL 무가시딩","IG·YT","30대 웰니스 KOL","#올더베러 #오늘도베러",50000,50000,0,30,"무가 시딩"),
("X (트위터) 시딩","X","리뷰·밈 2030","#올더베러 #올영템",500000,500000,350000,25,"할인 없음"),
("네이버 블로그","Blog","정보탐색 검색자","#올더베러 #웰니스구미추천",300000,200000,200000,15,""),
("EGC 콘텐츠","IG·TT","정보 신뢰 구매고려층","#올더베러 #직원피셜",1000000,450000,0,120,""),
("Half EGC 콘텐츠","IG·TT","건기식 관심 직장인","#올더베러 #웰니스루틴",1250000,650000,250000,24,""),
("콘텐츠 제작 (KV·영상·이미지)","제작","—","—",1085000,1085000,0,6,"KV·디자인·영상·이미지 소재"),
("파워페이지","SNS","구매 고의도 검색자","#올영추천템 #웰니스추천",800000,700000,600000,6,""),
("커뮤니티 체험단","뷰티앱·커뮤니티","카테고리 관여 高 사용자","#올리브영 #웰니스어워드",5000000,4000000,4000000,2,""),
("챌린저스","챌린저스 앱","올영 액티브 유저","#올영1위 #올더베러챌린지",5000000,4500000,3500000,2,""),
("BATi 마이크로 앰배서더","IG·TT","우수 마이크로 앰배서더","#올더베러 #오늘도베러",500000,350000,300000,60,"10명 × 6회(월1회) = 60건"),
("스프레이ai 솔루션","솔루션","—","—",4200000,1015000,0,6,"월 1식 · 6개월"),
]

# 월별 수량 분배 [7,8,9,10,11,12]
DIST = {
"나노":[19,38,22,19,26,36],
"마이크로 (UGC)":[22,43,26,22,30,37],
"매크로":[0,4,0,0,4,4],
"KOL 무가시딩":[5,5,5,5,5,5],
"X (트위터) 시딩":[0,8,0,0,8,9],
"네이버 블로그":[0,5,0,0,5,5],
"EGC 콘텐츠":[20,20,20,20,20,20],
"Half EGC 콘텐츠":[4,4,4,4,4,4],
"콘텐츠 제작 (KV·영상·이미지)":[1,1,1,1,1,1],
"파워페이지":[0,2,0,0,2,2],
"커뮤니티 체험단":[0,1,0,0,1,0],
"챌린저스":[0,1,0,0,1,0],
"BATi 마이크로 앰배서더":[10,10,10,10,10,10],
"스프레이ai 솔루션":[1,1,1,1,1,1],
}

# 월별 앵글(목적) 텍스트 [7,8,9,10,11,12]; "" = 해당 월 미진행
ANGLE = {
"나노":["진성 후기 볼륨 착수 · #오늘도베러 해시태그 자산화","다수 진성 후기로 '많이 보이는 제품' 인식 형성","구매 인증·리뷰 볼륨 집중","9월 위닝 소재 낙수·후기 자산 확산","진성 후기 재확대·대세감","구매 인증·쟁여두기 소구"],
"마이크로 (UGC)":["멜라나잇(자기 전)·올리브오일(아침 공복) 위닝 메시지 LMF 테스트","위닝 메시지 대량 확산 — 세일 1주 전 집중","세일 전환 소구 + '올영 1위' 랭킹 결합 소재","건기식(루테인·바나바잎·올인원) 확장 메시지 테스트","블프 재점화·위닝 소재 재투입","전 라인업 대표성 콘텐츠·세일 클로징"],
"매크로":["","세일 대세감 점화·도달 확대","","","블프 대세감 점화","세일 대세감·대표성"],
"KOL 무가시딩":["웰니스 코어 KOL 관계 형성·자발 콘텐츠","오가닉 콘텐츠 확대","전환 시점 자발 후기","관계형 콘텐츠 유지","오가닉 재가동","연말 자발 후기"],
"X (트위터) 시딩":["","실시간 트렌드·밈 바이럴","","","실시간 트렌드 바이럴","실시간 바이럴"],
"네이버 블로그":["","검색 후기·정보 탐색 상위 노출","","","검색 상위·추천 점유","검색 점유 유지"],
"EGC 콘텐츠":["직원 피셜 — 올리브오일 '이 스펙 이 가격' 신뢰형","제품 RTB 신뢰형 콘텐츠","세일 클로징 신뢰 소재","건기식 라인 신뢰 콘텐츠","RTB+혜택 결합 소재","대표성 신뢰 소재"],
"Half EGC 콘텐츠":["건기식 정보+공감 (루테인 모니터 앞 직장인)","건기식 인정 기능성 정보+공감","건기식 전환 정보 콘텐츠","건기식 정보+공감 확장","건기식 정보 콘텐츠","건기식 정보 콘텐츠"],
"콘텐츠 제작 (KV·영상·이미지)":["브랜드 KV·기본 소재 제작","세일 소재·B&A 컷 제작","세일 본편·마지막날 소재","건기식 라인 소재 제작","블프 소재 제작","세일 클로징 소재"],
"파워페이지":["","'추천템' 맥락·정보성 노출","","","추천 맥락 재노출","추천 맥락 노출"],
"커뮤니티 체험단":["","체험단·어워드로 '검증된 후기' 볼륨","","","검증 후기 보강",""],
"챌린저스":["","구매·리뷰 인증 챌린지 가동(랭킹 빌드업)","","","챌린저스 2차·신뢰 자산 보강",""],
"BATi 마이크로 앰배서더":["앰배서더 풀 온보딩·1차 협업 콘텐츠 발행","세일 前 앰배서더 위닝 메시지 협업","세일 피크 앰배서더 전환 콘텐츠","건기식 라인 앰배서더 협업","블프 앰배서더 재점화 협업","연말 대표성 앰배서더 협업"],
"스프레이ai 솔루션":["AI 스프레이 솔루션 운용(콘텐츠 자동 생성·배포 보조)"]*6,
}

LOOK = {it[0]: it for it in ITEMS}
MONTH_TITLE = [
"  ▣ 7월 · 착수 / 첫 점화 (Phase 1)   —   위닝 메시지 발굴 + 일반식품 시딩 도화선 점화 · 타겟: 웰니스 입문형",
"  ▣ 8월 · 세일 사전 부스팅 (1차 웨이브)   —   올영세일 前 인지·각인 확산 + 선택 증거 확보 · 타겟: 입문형+루틴형",
"  ▣ 9월 · 올영세일 1차 피크   —   확보한 증거로 구매 전환 + 카테고리 랭킹 상단 · 타겟: 실질 전환형",
"  ▣ 10월 · 유지 / 건기식 라인 확장   —   건기식 라인 확장 + 위닝 소재 낙수 · 타겟: 직장인 루틴형",
"  ▣ 11월 · 블프 (2차 웨이브)   —   재점화 — 많이 보이고 많이 찾는 브랜드 · 타겟: 입문형+루틴형",
"  ▣ 12월 · 올영세일 2차 피크 / 대표성 착지   —   웰니스 카테고리 대표 브랜드 착지 · 타겟: 실질 전환형",
]
MONTH_ORDER = ["마이크로 (UGC)","나노","매크로","KOL 무가시딩","X (트위터) 시딩","네이버 블로그",
"EGC 콘텐츠","Half EGC 콘텐츠","파워페이지","커뮤니티 체험단","챌린저스","BATi 마이크로 앰배서더",
"콘텐츠 제작 (KV·영상·이미지)","스프레이ai 솔루션"]

wb = openpyxl.Workbook()

# ============================================================== helpers
def C(ws,r,c,v=None,b=False,col=INK,fill=None,al="center",fmt=None,sz=10,wrap=False,border=True,it=False):
    cell=ws.cell(r,c)
    if v is not None: cell.value=v
    cell.font=Font(name=FONT,size=sz,bold=b,color=col,italic=it)
    if fill: cell.fill=PatternFill("solid",fgColor=fill)
    cell.alignment=Alignment(horizontal=al,vertical="center",wrap_text=wrap)
    if border: cell.border=bd
    if fmt: cell.number_format=fmt
    return cell

def mrg(ws,r1,c1,r2,c2):
    ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)

# ============================================================== 견적상세
ds = wb.active; ds.title="견적상세"
ds.sheet_view.showGridLines=False
widths={1:2.5,2:24,3:13,4:34,5:18,6:22,7:8,8:13,9:13,10:8,11:12,12:13,13:8,14:14,15:16}
for col,w in widths.items(): ds.column_dimensions[get_column_letter(col)].width=w

mrg(ds,2,2,2,15); C(ds,2,2,"올더베러 26년 하반기 견적 상세 — 월별 캠페인   (단위: 원 / VAT 별도)",b=True,col=DGREEN,al="left",sz=13,border=False)
mrg(ds,3,2,3,15); C(ds,3,2,"2026년 3,4Q : 07월 – 12월 / 총 예산 2.5억 원 / 인플루언서 시딩·바이럴 콘텐츠 통합 제안",col=GREY,al="left",sz=10,border=False)

# --- BATi 소개 블록 (최상단) ---
r=5
mrg(ds,r,2,r,15); C(ds,r,2,"■ BATi — 마이크로 앰배서더 프로그램 (신규)",b=True,col=DGREEN,fill=LGRN,al="left")
r=6
bati_txt=("BAT는 다년간의 인플루언서 시딩을 통해 우수한 마이크로 인플루언서 실(實)리드 풀을 보유하고 있으며, "
"'마이크로 MCN 모듈'을 통해 안정적이면서도 고퀄리티의 콘텐츠 발행이 가능합니다.  해당 인원들과 함께 "
"'마이크로 앰배서더' 프로세스를 도입하여 6개월간 각 인원당 월 1회 이상(총 6회)의 협업을 진행합니다.  "
"e.g. 일본 마이크로 인플루언서의 SK-II 다회차 협업 사례")
mrg(ds,r,2,r+2,15); C(ds,r,2,bati_txt,al="left",wrap=True,sz=10,col=INK)
ds.row_dimensions[r].height=58

# --- 견적 요약(이미지) 박스 ---
r=10
mrg(ds,r,2,r,15); C(ds,r,2,"■ 견적 요약 (6개월 취합 · VAT 별도)",b=True,col=DGREEN,fill=LGRN,al="left")
sum_hdr=["할인 원고료 및 할인가 비용","총 마크업(마진) 비용","총 원고료 + 마크업 비용","예상 마크업(%)"]
sum_fill=[OLIVE,OLIVE,DARK,GREEN_HL]; sum_tc=[WHITE,WHITE,WHITE,INK]
# 4개 컬럼을 B:E / F:H / I:K / L:O 로 배치
spans=[(2,4),(5,7),(8,11),(12,15)]
for (c1,c2),h,f,t in zip(spans,sum_hdr,sum_fill,sum_tc):
    mrg(ds,11,c1,11,c2); C(ds,11,c1,h,b=True,col=t,fill=f,wrap=True)
sum_val=["=N28-N29","=N29","=N28","=N29/N28"]  # 실비합 / 마진합 / 총견적 / 마진율  (placeholders -> set below)
# 실제 값은 전체합계 합계행을 참조: 총견적합=N(합계행), 마진합=L합, 실비합=총-마진
# 합계행 위치는 아래에서 확정(전체합계 표 시작 r0). 우선 표 생성 후 다시 채움.

# --- 전체 합계 표 ---
r0=14
mrg(ds,r0,2,r0,15); C(ds,r0,2,"■ 전체 합계 (6개월 취합)",b=True,col=DGREEN,fill=LGRN,al="left")
hdr=["항목","채널","캠페인 및 활용 목적 / 앵글","타겟 페르소나","필수 해시태그","목표수량",
     "원고료 및 정가","할인 원고료 및 할인가","할인율(%)","건당 실비","마크업(건당)","마크업(%)","총 견적(합계)","비고"]
hr=r0+1
fillmap={8:OLIVE,9:OLIVE,11:GREEN_HL,12:GREEN_HL,13:DARK}
tcmap={8:WHITE,9:WHITE,11:INK,12:INK,13:WHITE}
for i,h in enumerate(hdr,2):
    C(ds,hr,i,h,b=True,col=tcmap.get(i,WHITE),fill=fillmap.get(i,DARK),wrap=True)
C(ds,hr-0,16,"")  # keep border clean (col P unused)
rr=hr+1
first=rr
for nm,ch,per,tag,jeong,hal,sil,qty,memo in ITEMS:
    C(ds,rr,2,nm,b=True,al="left",wrap=True)
    C(ds,rr,3,ch); C(ds,rr,4,"6개월 합산",al="left",sz=9,col=GREY)
    C(ds,rr,5,per,al="left",wrap=True,sz=9)
    C(ds,rr,6,tag,al="left",wrap=True,sz=9,col=DGREEN)
    C(ds,rr,7,qty,b=True)
    C(ds,rr,8,jeong,fmt=ACC)
    C(ds,rr,9,hal,fmt=ACC)
    C(ds,rr,10,f"=1-(I{rr}/H{rr})",fmt=PCT)
    C(ds,rr,11,sil,fmt=ACC,col=GREY)
    C(ds,rr,12,f"=I{rr}-K{rr}",fmt=ACC,col=DGREEN)   # 마크업(건당)=할인가-실비
    C(ds,rr,13,f"=IF(I{rr}=0,0,L{rr}/I{rr})",fmt=PCT)
    C(ds,rr,14,f"=G{rr}*I{rr}",fmt=ACC,b=True)        # 총 견적 = 할인가×수량
    C(ds,rr,15,memo,al="left",wrap=True,sz=8,col=GREY)
    rr+=1
last=rr-1
# 합계행
trow=rr
mrg(ds,trow,2,trow,6); C(ds,trow,2,"총 합계 (할인가, VAT 별도)",b=True,col=WHITE,fill=DARK,al="right")
C(ds,trow,7,f"=SUM(G{first}:G{last})",b=True,col=WHITE,fill=DARK)
C(ds,trow,8,f"=SUMPRODUCT(G{first}:G{last},H{first}:H{last})",fmt=ACC,b=True,col=WHITE,fill=DARK)
C(ds,trow,9,f"=N{trow}",fmt=ACC,b=True,col=WHITE,fill=DARK)
C(ds,trow,10,f"=1-(N{trow}/H{trow})",fmt=PCT,b=True,col=WHITE,fill=DARK)
C(ds,trow,11,f"=SUMPRODUCT(G{first}:G{last},K{first}:K{last})",fmt=ACC,b=True,col=WHITE,fill=DARK)
C(ds,trow,12,f"=N{trow}-K{trow}*1",fmt=ACC,b=True,col=WHITE,fill=DARK)  # 마진합 = 총견적-실비합
C(ds,trow,13,f"=L{trow}/N{trow}",fmt=PCT,b=True,col=WHITE,fill=DARK)
C(ds,trow,14,f"=SUM(N{first}:N{last})",fmt=ACC,b=True,col=WHITE,fill=DGREEN)
C(ds,trow,15,"",fill=DARK)
vrow=trow+1
C(ds,vrow,13,"부가세(10%)",b=True,al="right"); C(ds,vrow,14,f"=N{trow}*0.1",fmt=ACC,b=True)
vrow2=trow+2
C(ds,vrow2,13,"합계(VAT 포함)",b=True,al="right",col=WHITE,fill=DARK); C(ds,vrow2,14,f"=N{trow}*1.1",fmt=ACC,b=True,col=WHITE,fill=DARK)

# 요약 박스 값 채우기 (전체합계 합계행 trow 참조)
# 실비합 = K합 = K{trow}; 마진합 = L{trow}; 총견적 = N{trow}; 마진율 = L{trow}/N{trow}
realsum=[f"=K{trow}",f"=L{trow}",f"=N{trow}",f"=L{trow}/N{trow}"]
realfmt=[ACC,ACC,ACC,PCT]
for (c1,c2),v,f in zip(spans,realsum,realfmt):
    mrg(ds,12,c1,12,c2); C(ds,12,c1,v,b=True,fmt=f,sz=12,col=DGREEN)

# --- 월별 상세 ---
r=trow+4
mrg(ds,r,2,r,15); C(ds,r,2,"■ 월별 상세 (7월 ~ 12월)",b=True,col=DGREEN,fill=LGRN,al="left")
r+=1
mhdr=["항목","채널","캠페인 및 활용 목적 / 앵글","타겟 페르소나","필수 해시태그","목표수량","단가(할인가)","금액"]
for mi in range(6):
    # 헤더
    for i,h in enumerate(mhdr,2):
        C(ds,r,i,h,b=True,col=WHITE,fill=DARK,wrap=True)
    # 빈 col 10-15 정리 (병합 안함, 그냥 둠)
    r+=1
    mrg(ds,r,2,r,15); C(ds,r,2,MONTH_TITLE[mi],b=True,col=DGREEN,fill=LGRN,al="left",wrap=True)
    r+=1
    blk_first=r
    for nm in MONTH_ORDER:
        q=DIST[nm][mi]
        if q==0: continue
        ch,per,tag,jeong,hal,sil,qty,memo = LOOK[nm][1],LOOK[nm][2],LOOK[nm][3],LOOK[nm][4],LOOK[nm][5],LOOK[nm][6],LOOK[nm][7],LOOK[nm][8]
        ang=ANGLE[nm][mi]
        C(ds,r,2,nm,b=True,al="left",wrap=True,sz=9)
        C(ds,r,3,ch,sz=9); C(ds,r,4,ang,al="left",wrap=True,sz=9)
        C(ds,r,5,per,al="left",wrap=True,sz=9); C(ds,r,6,tag,al="left",wrap=True,sz=8,col=DGREEN)
        C(ds,r,7,q,b=True); C(ds,r,8,hal,fmt=ACC)
        C(ds,r,9,f"=G{r}*H{r}",fmt=ACC,b=True)
        # col 10-15 비우기(테두리 유지 위해 채움)
        for cc in range(10,16): C(ds,r,cc,"")
        r+=1
    blk_last=r-1
    mrg(ds,r,2,r,6); C(ds,r,2,f"{mi+7}월 소계",b=True,fill=SUB,al="right")
    C(ds,r,7,f"=SUM(G{blk_first}:G{blk_last})",b=True,fill=SUB)
    C(ds,r,8,"",fill=SUB)
    C(ds,r,9,f"=SUM(I{blk_first}:I{blk_last})",fmt=ACC,b=True,fill=SUB)
    for cc in range(10,16): C(ds,r,cc,"",fill=SUB)
    r+=2
# 월별 총합계
mrg(ds,r,2,r,6); C(ds,r,2,"월별 총 합계 (할인가, VAT 별도)",b=True,col=WHITE,fill=DARK,al="right")
C(ds,r,9,f"=N{trow}",fmt=ACC,b=True,col=WHITE,fill=DGREEN)
mon_total_row=r
r+=3

# --- 무상 지원 ---
mrg(ds,r,2,r,15); C(ds,r,2,"■ 무상 지원 (Value-add · 0원)",b=True,col=DGREEN,fill=LGRN,al="left")
r+=1
freebies=[
("고성과 리포트 대시보드","무상","캠페인 성과·랭킹 트래킹 실시간 제공"),
("올리비아 (AI 어시스턴트)","무상","AI 콘텐츠·리포팅 어시스턴트 제공"),
("VOC 딥다이브 리포트","무상","리뷰 크롤링·키워드 분석 리포트"),
("Claude AI 심의 사전검수","무상","심의 반려율↓·리드타임 단축"),
("월간 성과 리포팅","무상","캠페인 완수 후 리포팅 제공"),
("위닝 콘텐츠 2차 활용","무상","시딩 콘텐츠 2차 활용 라이선스"),
]
for nm,ch,desc in freebies:
    C(ds,r,2,nm,b=True,al="left"); C(ds,r,3,ch)
    mrg(ds,r,4,r,13); C(ds,r,4,desc,al="left",sz=9,col=GREY)
    C(ds,r,14,"-"); C(ds,r,15,"")
    r+=1
mrg(ds,r,2,r,15); C(ds,r,2,"※ 스프레이ai 솔루션은 본 견적 내 유상 항목으로 반영(할인가 1,015,000원/식). 그 외 솔루션은 대행 기간 무상 지원.",al="left",sz=8,col=GREY,it=True)

# ============================================================== 견적제안
qp=wb.create_sheet("견적제안"); qp.sheet_view.showGridLines=False
for col,w in {2:2.5,3:16,4:2,5:120}.items(): qp.column_dimensions[get_column_letter(col)].width=w
def QL(r,label,val,head=False):
    if label:
        C(qp,r,3,label,b=True,col=(WHITE if head else DGREEN),fill=(DARK if head else None),al="left",border=False)
    if val is not None:
        C(qp,r,5,val,al="left",wrap=True,sz=10,border=False)
mrg(qp,3,3,3,5); C(qp,3,3,"BAT  |  견적 제안서",b=True,col=WHITE,fill=DGREEN,al="left",sz=15,border=False)
QL(5,"· 프로젝트명","올리브영 PB 올더베러 브랜드 빌딩 캠페인")
QL(7,"· 집행 기간","2026년 7월 1일 ~ 2026년 12월 31일 (6개월) · 인플루언서 시딩·바이럴 통합 운영")
QL(9,"· 집행 내용","① 인플루언서 시딩(나노·마이크로 코어) 중심 진성 UGC 대량 확산")
QL(10,None,"② BATi 마이크로 앰배서더(10명×6회) 다회차 협업으로 고퀄리티 콘텐츠 지속 발행")
QL(11,None,"③ 바이럴(챌린저스·커뮤니티·파워페이지) + EGC/Half EGC 신뢰 콘텐츠")
QL(12,None,"④ 위닝 콘텐츠·시딩 자산을 올영세일 구매 전환으로 직접 연결")
QL(14,"· 견적 요약","· 총 견적 : KRW 250,000,000원 (VAT 별도)")
QL(15,None,"· 할인 원고료(실집행) 138,350,000원 + 총 마크업(마진) 111,650,000원  ·  예상 마크업률 44.7%")
QL(16,None,"· 목표 수량 : 크리에이터 시딩 422건 + BATi 60건 + EGC 120 · Half EGC 24 + 콘텐츠 제작 6 + 바이럴 10 + 솔루션 6")
QL(18,"· A/C 견적 상세","· 크리에이터 시딩 : 125,600,000원  (나노 160 · 마이크로 180 · 매크로 12 · KOL 30 · X 25 · 블로그 15)",head=False)
QL(19,None,"· BATi 마이크로 앰배서더 : 21,000,000원  (10명 × 6회 = 60건 × 350,000원)")
QL(20,None,"· 콘텐츠 EGC : 69,600,000원  (EGC 120 × 450,000 · Half EGC 24 × 650,000)")
QL(21,None,"· 콘텐츠 제작(KV·영상·이미지) : 6,510,000원  (6식 × 1,085,000)")
QL(22,None,"· 바이럴 : 21,200,000원  (파워페이지 6 · 커뮤니티 2 · 챌린저스 2)")
QL(23,None,"· 스프레이ai 솔루션 : 6,090,000원  (6식 × 1,015,000)")
QL(24,None,"· 총 합계 : 250,000,000원 (VAT 별도)")
QL(26,"· 단가 안내","· 나노 22.5만 / 마이크로 35만 / 매크로 80만 / BATi 35만 (정가 50만 → 할인 35만)")
QL(27,None,"· EGC 45만 / Half EGC 65만 / 콘텐츠 제작 108.5만 · 마크업(마진)은 할인가 − 실집행 원고료")
QL(29,"· BATi 안내","· BAT 보유 우수 마이크로 인플루언서 실리드 풀 + 마이크로 MCN 모듈 기반 '마이크로 앰배서더' 6개월 다회차 협업")
QL(30,None,"· 1명당 월 1회 이상(총 6회) 협업 · e.g. 일본 마이크로 인플루언서 SK-II 다회차 협업 사례")
QL(32,"· 무상 지원","· 고성과 리포트 대시보드 / 올리비아 AI 어시스턴트 / VOC 딥다이브 / Claude AI 심의 검수 / 월간 리포팅 / 위닝 2차 활용")
QL(34,"· 특이사항","· 시장 단가 변동 및 모집 기간에 따라 총 원고료는 소폭 변동될 수 있습니다.")
QL(35,None,"· 인플루언서 섭외 & 콘텐츠 업로드 & 원고료 지급까지 BAT 담당 ※ 제품 발송은 올리브영 PB")
QL(36,None,"· 항목별 상세는 '견적상세'(월별), 집행 일정은 '월별 액션플랜' 시트를 확인 바랍니다.")

# ============================================================== 월별 액션플랜
ap=wb.create_sheet("월별 액션플랜"); ap.sheet_view.showGridLines=False
for col,w in {1:30,2:11,3:7,4:9,5:9,6:7,7:9,8:9,9:8,10:14}.items(): ap.column_dimensions[get_column_letter(col)].width=w
C(ap,1,1,"월별 실행 타임라인 · 액션플랜",b=True,col=DGREEN,sz=13,border=False)
aph=["항목","단가","7월","8월\n세일사전","9월\n세일★","10월","11월\n블프★","12월\n세일★","합계","금액"]
for i,h in enumerate(aph,1): C(ap,3,i,h,b=True,col=WHITE,fill=DARK,wrap=True)
ap_rows=[
("■ 크리에이터 시딩",None,None),
("나노",225000,"나노"),
("마이크로 (UGC)",350000,"마이크로 (UGC)"),
("매크로",800000,"매크로"),
("KOL 무가시딩",50000,"KOL 무가시딩"),
("X (트위터) 시딩",500000,"X (트위터) 시딩"),
("네이버 블로그",200000,"네이버 블로그"),
("■ BATi 앰배서더",None,None),
("BATi 마이크로 앰배서더",350000,"BATi 마이크로 앰배서더"),
("■ 콘텐츠",None,None),
("EGC 콘텐츠",450000,"EGC 콘텐츠"),
("Half EGC 콘텐츠",650000,"Half EGC 콘텐츠"),
("콘텐츠 제작 (KV·영상·이미지)",1085000,"콘텐츠 제작 (KV·영상·이미지)"),
("■ 바이럴 · 전환",None,None),
("파워페이지",700000,"파워페이지"),
("커뮤니티 체험단",4000000,"커뮤니티 체험단"),
("챌린저스",4500000,"챌린저스"),
("■ 솔루션",None,None),
("스프레이ai 솔루션",1015000,"스프레이ai 솔루션"),
]
r=4; data_rows=[]
for label,price,key in ap_rows:
    if price is None:
        mrg(ap,r,1,r,10); C(ap,r,1,label,b=True,col=DGREEN,fill=LGRN,al="left")
        r+=1; continue
    C(ap,r,1,label,al="left",sz=9); C(ap,r,2,price,fmt=ACC,sz=9)
    for mi in range(6):
        q=DIST[key][mi]
        C(ap,r,3+mi,(q if q>0 else "-"))
    C(ap,r,9,f"=SUM(C{r}:H{r})",b=True)
    C(ap,r,10,f"=I{r}*B{r}",fmt=ACC,b=True)
    data_rows.append(r); r+=1
# 월 시딩 총건수 (크리에이터 시딩만: 나노~블로그)
seed_rows=[rr for rr,(l,p,k) in zip([4,5,6,7,8,9,10],ap_rows[0:7])]  # not used
r+=0
totrow=r+1
C(ap,totrow,1,"총 견적 (할인가, VAT 별도)",b=True,col=WHITE,fill=DARK,al="right")
mrg(ap,totrow,1,totrow,9)
C(ap,totrow,10,"+".join(f"J{rr}" for rr in data_rows),fmt=ACC,b=True,col=WHITE,fill=DGREEN)
ap.cell(totrow,10).value="="+ "+".join(f"J{rr}" for rr in data_rows)

# ============================================================== 마진 계산
mg=wb.create_sheet("마진 계산(확인 후 삭제)"); mg.sheet_view.showGridLines=False
for col,w in {3:24,4:12,5:11,6:11,7:8,8:14,9:14,10:9}.items(): mg.column_dimensions[get_column_letter(col)].width=w
mh=["항목","할인가(단가)","건당 실비","건당 마진","수량","매출","마진","마진율"]
for i,h in enumerate(mh,3): C(mg,3,i,h,b=True,col=WHITE,fill=DARK,wrap=True)
r=4; mg_rows=[]
for nm,ch,per,tag,jeong,hal,sil,qty,memo in ITEMS:
    C(mg,r,3,nm,al="left",sz=9)
    C(mg,r,4,hal,fmt=ACC); C(mg,r,5,sil,fmt=ACC); C(mg,r,6,f"=D{r}-E{r}",fmt=ACC)
    C(mg,r,7,qty,b=True)
    C(mg,r,8,f"=D{r}*G{r}",fmt=ACC); C(mg,r,9,f"=F{r}*G{r}",fmt=ACC,b=True,col=DGREEN)
    C(mg,r,10,f"=IF(H{r}=0,0,I{r}/H{r})",fmt=PCT)
    mg_rows.append(r); r+=1
tr=r
C(mg,tr,3,"합계",b=True,col=WHITE,fill=DARK,al="right")
C(mg,tr,7,f"=SUM(G4:G{r-1})",b=True,col=WHITE,fill=DARK)
C(mg,tr,8,f"=SUM(H4:H{r-1})",fmt=ACC,b=True,col=WHITE,fill=DARK)
C(mg,tr,9,f"=SUM(I4:I{r-1})",fmt=ACC,b=True,col=WHITE,fill=DGREEN)
C(mg,tr,10,f"=I{tr}/H{tr}",fmt=PCT,b=True,col=WHITE,fill=DARK)
C(mg,tr,4,"",fill=DARK);C(mg,tr,5,"",fill=DARK);C(mg,tr,6,"",fill=DARK)
r=tr+2
for txt in [f"매출(총 견적) : =H{tr}","총 실비 : =H{tr}-I{tr}","총 마진(마크업) : =I{tr}","블렌디드 마진율 : =I{tr}/H{tr}"]:
    lab=txt.split(" : ")[0]
    C(mg,r,3,lab,b=True,al="left",border=False,col=DGREEN)
    if "마진율" in lab:
        C(mg,r,5,f"=I{tr}/H{tr}",fmt=PCT,b=True,al="left",border=False)
    elif "실비" in lab:
        C(mg,r,5,f"=H{tr}-I{tr}",fmt=ACC,b=True,al="left",border=False)
    elif "마진" in lab:
        C(mg,r,5,f"=I{tr}",fmt=ACC,b=True,al="left",border=False)
    else:
        C(mg,r,5,f"=H{tr}",fmt=ACC,b=True,al="left",border=False)
    r+=1
C(mg,r+1,3,"※ 본 시트는 내부 마진 확인용입니다(확인 후 삭제).",col=GREY,sz=8,it=True,al="left",border=False)

wb.save("proposal/올더베러_견적서_BAT.xlsx")
print("saved -> proposal/올더베러_견적서_BAT.xlsx")
