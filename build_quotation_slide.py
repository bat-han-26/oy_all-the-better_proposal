# -*- coding: utf-8 -*-
"""견적서 (Quotation) 1p — 최종 확정 금액(총 250,000,000 / VAT 포함 275,000,000)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

FONT="맑은 고딕"
INK=RGBColor(0x22,0x22,0x22); SUBINK=RGBColor(0x88,0x88,0x86); WHITE=RGBColor(0xFF,0xFF,0xFF)
BLACK=RGBColor(0x1A,0x1A,0x1A); BAND=RGBColor(0xEC,0xEC,0xEC); LINE=RGBColor(0xDD,0xDD,0xD8)

prs=Presentation(); prs.slide_width=Inches(13.33); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])

def setfont(run,sz,col,bold=False):
    run.font.name=FONT; run.font.size=Pt(sz); run.font.bold=bold; run.font.color.rgb=col
    rPr=run._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
    if ea is None: ea=rPr.makeelement(qn('a:ea'),{}); rPr.append(ea)
    ea.set('typeface',FONT)

def tb(x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=False; tf.vertical_anchor=anchor
    tf.margin_left=Pt(2);tf.margin_right=Pt(2);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0)
    p=tf.paragraphs[0]; p.alignment=align
    for t,sz,col,bold in runs:
        r=p.add_run(); r.text=t; setfont(r,sz,col,bold)
    return box

def rect(x,y,w,h,fill=None,line=None,lw=0.75):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False
    return sp

# ---- columns ----
X0,W0=0.55,7.75          # 세부 항목
XQ,WQ=8.30,1.20          # 수량 (center)
XP,WP=9.50,1.55          # 단가 (right)
XA,WA=11.05,1.73         # 금액 (right)
RIGHT=XA+WA              # 12.78

# ---- title ----
tb(0.55,0.26,11.0,0.62,[("견적서 ",27,INK,True),("(Quotation)",20,INK,True)])
tb(0.57,0.86,11.5,0.34,[("올리브영 PB 올더베러 브랜드 빌딩 캠페인  ·  단위 원 (VAT 별도)",12,SUBINK,False)])

# ---- header ----
y=1.20; HH=0.36
rect(X0,y,RIGHT-X0,HH,fill=BLACK)
tb(X0+0.2,y,W0-0.2,HH,[("세부 항목",12,WHITE,True)])
tb(XQ,y,WQ,HH,[("수량",12,WHITE,True)],align=PP_ALIGN.CENTER)
tb(XP,y,WP-0.15,HH,[("단가",12,WHITE,True)],align=PP_ALIGN.RIGHT)
tb(XA,y,WA-0.18,HH,[("금액",12,WHITE,True)],align=PP_ALIGN.RIGHT)
y+=HH

GROUPS=[
("제작비",[
  ("콘텐츠 제작 (KV·디자인·영상·이미지)","6","1,000,000","6,000,000"),
  ("EGC 콘텐츠","120","500,000","60,000,000"),
  ("Half EGC 콘텐츠","24","750,000","18,000,000"),
]),
("크리에이터 시딩",[
  ("나노","90","270,000","24,300,000"),
  ("마이크로 (UGC)","140","420,000","58,800,000"),
  ("매크로 (인스타·유튜브)","12","880,000","10,560,000"),
  ("KOL 무가시딩","30","100,000","3,000,000"),
  ("X (트위터) 시딩","24","550,000","13,200,000"),
  ("네이버 블로그","15","220,000","3,300,000"),
  ("BATi 마이크로 앰배서더 (시딩)","60","500,000","30,000,000"),
]),
("콘텐츠 · 바이럴",[
  ("파워페이지 콘텐츠 바이럴","6","750,000","4,500,000"),
  ("커뮤니티 체험단 (탑 2건)","2","4,300,000","8,600,000"),
  ("챌린저스","2","4,800,000","9,600,000"),
]),
("솔루션",[
  ("스프레이ai 솔루션","6","23,333","140,000"),
]),
]
BH=0.275; IH=0.265
for gname,items in GROUPS:
    rect(X0,y,RIGHT-X0,BH,fill=BAND)
    tb(X0+0.2,y,RIGHT-X0-0.3,BH,[(gname,11.5,INK,True)])
    y+=BH
    for nm,qty,price,amt in items:
        rect(X0,y,RIGHT-X0,IH,fill=WHITE,line=LINE,lw=0.5)
        tb(X0+0.2,y,W0-0.3,IH,[(nm,10.5,INK,False)])
        tb(XQ,y,WQ,IH,[(qty,10.5,SUBINK,False)],align=PP_ALIGN.CENTER)
        tb(XP,y,WP-0.15,IH,[(price if price!="-" else "-",10.5,SUBINK,False)],align=PP_ALIGN.RIGHT)
        tb(XA,y,WA-0.18,IH,[(amt,10.5,INK,True)],align=PP_ALIGN.RIGHT)
        y+=IH

# ---- totals ----
TH=0.32
def total_row(label,amt,dark=False,big=False):
    global y
    if dark:
        rect(XQ,y,RIGHT-XQ,TH+0.02,fill=BLACK)
        tb(XQ,y,(XA-XQ)-0.15,TH+0.02,[(label,12.5,WHITE,True)],align=PP_ALIGN.RIGHT)
        tb(XA,y,WA-0.18,TH+0.02,[(amt,13,WHITE,True)],align=PP_ALIGN.RIGHT)
    else:
        tb(XQ,y,(XA-XQ)-0.15,TH,[(label,11.5,INK,True)],align=PP_ALIGN.RIGHT)
        tb(XA,y,WA-0.18,TH,[(amt,11.5,INK,True)],align=PP_ALIGN.RIGHT)
    y+=TH+(0.02 if dark else 0)

y+=0.06
total_row("공급가액 소계","250,000,000")
total_row("부가가치세 (10%)","25,000,000")
y+=0.02
total_row("합계 금액 (VAT 포함)","275,000,000",dark=True)

# ---- footer ----
tb(0.55,7.12,6,0.3,[("© 2026 BAT",10,SUBINK,False)])

prs.save("proposal/올더베러_견적서_Quotation_1p.pptx")
print("saved. table bottom y=",round(y,2))
