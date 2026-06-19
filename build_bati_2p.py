# -*- coding: utf-8 -*-
"""BATi 2p — ① 마이크로 앰배서더 프로그램(개념) ② 운영 시스템(핸드북 기반 보강)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from template_lib import *
from template_lib import _set_font

prs=Presentation(); prs.slide_width=Inches(13.33); prs.slide_height=Inches(7.5)
BL=prs.slide_layouts[6]

def line2(s,x1,y1,x2,y2,color=GREENA,w=1.4):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w); c.shadow.inherit=False; return c

def chip(s,x,y,w,h,en,kr,fill=GREENA):
    rect(s,x,y,w,h,fill=fill,rounded=True)
    textbox(s,x,y+0.07,w,0.26,[(en,9,LIME,True,1.2)],align=PP_ALIGN.CENTER)
    textbox(s,x,y+0.30,w,h-0.34,[(kr,12.5,WHITE,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def card(s,x,y,w,h,no,en,title,bullets,accent=GREENA):
    rect(s,x,y,w,h,fill=BEIGE,line=LGRAY,line_w=1.0,rounded=True)
    rect(s,x,y,0.12,h,fill=accent)
    textbox(s,x+0.3,y+0.16,w-0.5,0.26,[(f"{no}  {en}",9.5,accent,True,1.0)])
    textbox(s,x+0.3,y+0.42,w-0.5,0.4,[(title,14.5,INK,True)])
    bullets_tb=s.shapes.add_textbox(Inches(x+0.3),Inches(y+0.92),Inches(w-0.55),Inches(h-1.05))
    tf=bullets_tb.text_frame; tf.word_wrap=True
    tf.margin_left=Pt(0);tf.margin_top=Pt(0);tf.margin_right=Pt(0);tf.margin_bottom=Pt(0)
    for i,b in enumerate(bullets):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(5); p.line_spacing=1.02
        r=p.add_run(); r.text="·  "+b; _set_font(r,10.5,RGBColor(0x3A,0x3A,0x36),False)

# ============================================================ SLIDE 1
s=prs.slides.add_slide(BL)
content_header(s,"BATi","AMBASSADOR","마이크로 앰배서더 — 6개월 다회차 협업 프로그램")
textbox(s,1.4,1.84,10.53,0.4,[("BAT가 보유한 우수 마이크로 인플루언서 ‘실(實)리드 풀’을 ‘마이크로 앰배서더’로 전환합니다",12,GRAY,False)],align=PP_ALIGN.CENTER)

# 3 concept cards
cy=2.45; ch=1.62; cw=3.71; xs=[0.9,4.81,8.72]
concepts=[
 ("01","REAL-LEAD POOL","실(實)리드 풀",
  "다년간의 인플루언서 시딩으로 축적한 우수 마이크로 인플루언서 실리드 보유 — 검증된 풀에서 즉시 가동"),
 ("02","MICRO MCN","마이크로 MCN 모듈",
  "MCN 운영 모듈로 안정적·고퀄리티 콘텐츠를 지속 발행 — 단발 협업의 품질 편차를 제거"),
 ("03","AMBASSADOR","마이크로 앰배서더",
  "6개월간 1인당 월 1회 이상(총 6회) 다회차 협업 — 브랜드 맥락이 누적된 관계형 콘텐츠"),
]
for (no,en,t,d),x in zip(concepts,xs):
    rect(s,x,cy,cw,ch,fill=BEIGE,line=LGRAY,line_w=1.0,rounded=True)
    oval(s,x+0.3,cy+0.26,0.5,fill=GREENA)
    textbox(s,x+0.3,cy+0.26,0.5,0.5,[(no,13,WHITE,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    textbox(s,x+0.95,cy+0.24,cw-1.1,0.24,[(en,8.5,GREENA,True,1.0)])
    textbox(s,x+0.95,cy+0.46,cw-1.1,0.32,[(t,13.5,INK,True)])
    textbox(s,x+0.3,cy+0.95,cw-0.55,ch-1.05,[(d,10,RGBColor(0x3A,0x3A,0x36),False)])

# 타임라인 + stat band
by=4.35; bh=1.72
rect(s,0.9,by,7.0,bh,fill=WHITE,line=LGRAY,line_w=1.0,rounded=True)
textbox(s,1.2,by+0.18,6.5,0.3,[("마이크로 앰배서더 1인 · 6개월 협업 사이클",12.5,INK,True)])
tlx=1.35; tgap=1.0; tly=by+0.95
months=["7월","8월","9월","10월","11월","12월"]
for i in range(6):
    cx=tlx+i*tgap
    if i<5: line2(s,cx+0.17,tly+0.17,cx+tgap-0.17,tly+0.17,GREENA,1.6)
for i,m in enumerate(months):
    cx=tlx+i*tgap
    oval(s,cx,tly,0.34,fill=GREENA)
    textbox(s,cx-0.16,tly,0.66,0.34,[("●",8,WHITE,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    textbox(s,cx-0.25,tly+0.4,0.84,0.24,[(m,9.5,INK,True)],align=PP_ALIGN.CENTER)
    textbox(s,cx-0.25,tly+0.62,0.84,0.22,[("협업 1회",8,GRAY,False)],align=PP_ALIGN.CENTER)

# stat box
rect(s,8.12,by,4.31,bh,fill=GREENA,rounded=True)
textbox(s,8.12,by+0.22,4.31,0.34,[("본 캠페인 적용 규모",10.5,LIME,True)],align=PP_ALIGN.CENTER)
textbox(s,8.12,by+0.52,4.31,0.6,[("10명 × 6회 = 60건",26,WHITE,True)],align=PP_ALIGN.CENTER)
textbox(s,8.32,by+1.16,3.91,0.44,[("앰배서더 10인이 6개월간 매월 1회씩 협업",10,RGBColor(0xCF,0xE6,0xD8),False)],align=PP_ALIGN.CENTER)

# SK-II 사례 + bottom strip
rect(s,0.9,by+bh+0.16,11.53,0.46,fill=BEIGE,line=LGRAY,line_w=0.8,rounded=True)
textbox(s,1.2,by+bh+0.16,11.0,0.46,[("e.g.  ",10.5,GREENA,True),("일본 마이크로 인플루언서의 SK-II 다회차 앰배서더 협업 사례 — 반복 노출로 신뢰·전환을 누적",10.5,INK,False)],anchor=MSO_ANCHOR.MIDDLE)
textbox(s,0.9,by+bh+0.74,11.53,0.34,[("단발성 바이럴 ×   →   지속 관계 기반 네트워크 ○      ·      브랜드 캠페인의 성과와 효율을 동시에 향상",11,MAIN,True)],align=PP_ALIGN.CENTER)

# ============================================================ SLIDE 2
s=prs.slides.add_slide(BL)
content_header(s,"BATi","OPERATIONS","검증된 운영 시스템 — 크리에이터 그로스마케팅")
textbox(s,1.4,1.84,10.53,0.4,[("‘마이크로 MCN 운영 핸드북’ 기반 — 감각이 아닌 표준화·데이터·자동화로 운영되는 시스템",12,GRAY,False)],align=PP_ALIGN.CENTER)

gx=[0.9,6.95]; gy=[2.4,4.35]; gw=5.48; gh=1.78
pillars=[
 ("01","NETWORK","느슨한 네트워크",[
   "고비용 계약 모델 × → 비계약·저비용 구조",
   "인플루언서는 가볍게 온보딩, 풀은 무한 확장",
   "단발 바이럴이 아닌 ‘생태계·관계’ 자산화"]),
 ("02","TIERING","데이터 기반 티어링",[
   "친분 × · 데이터 ○ 로 티어 승급/강등 결정",
   "종합 퍼포먼스·인게이지먼트·콘텐츠 만족도 기준",
   "개인별 기준점 · 가이드 미준수 시 강등"]),
 ("03","INCENTIVE","인센티브 시스템",[
   "업로드 트래킹·조회수 임계치 기준 인센티브 지급",
   "캠페인별 단가 별도 · 전담 매니저가 내역 관리",
   "고성과 생산자에게 보상 집중 → 품질 상향"]),
 ("04","GOVERNANCE","양해각서(MOU) 운영",[
   "법적 부담 없는 느슨한 합의 · 단일 커뮤니케이션 창구",
   "브랜드 가이드·일정은 반드시 준수",
   "경쟁 카테고리 충돌 통제(원칙 MOU 명시)"]),
]
pos=[(gx[0],gy[0]),(gx[1],gy[0]),(gx[0],gy[1]),(gx[1],gy[1])]
acc=[GREENA,MAIN,GREENA,MAIN]
for (no,en,t,bl),(x,y),a in zip(pillars,pos,acc):
    card(s,x,y,gw,gh,no,en,t,bl,accent=a)

# vision bar
rect(s,0.9,6.42,11.53,0.66,fill=MAIN,rounded=True)
textbox(s,1.2,6.42,11.0,0.66,[("우리는 크리에이터를 ‘채널’로 소비하지 않습니다  —  ",11.5,WHITE,True),("브랜드와 함께 성장하는 ‘관계와 시스템’을 구축합니다",11.5,LIME,True)],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)

prs.save("proposal/올더베러_BATi_2p.pptx")
print("saved", len(prs.slides._sldIdLst))
