#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""올더베러 3·4분기 월별 액션 장표 (편집 가능 네이티브 PPT)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BLACK=RGBColor(0x11,0x11,0x11); DARK=RGBColor(0x33,0x33,0x33); GRAY=RGBColor(0x70,0x70,0x70)
MG=RGBColor(0xB0,0xB0,0xB0); LG=RGBColor(0xED,0xED,0xED); WHITE=RGBColor(0xFF,0xFF,0xFF)
FONT="맑은 고딕"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
SW,SH=prs.slide_width,prs.slide_height

def rect(s,x,y,w,h,color):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tf=s.shapes.add_textbox(x,y,w,h).text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,(t,sz,b,c) in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(3)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
    return tf
def cell(c,t,sz=9.5,b=False,col=BLACK,fill=WHITE,al=PP_ALIGN.CENTER):
    c.fill.solid(); c.fill.fore_color.rgb=fill
    c.margin_left=Pt(5); c.margin_right=Pt(5); c.margin_top=Pt(1); c.margin_bottom=Pt(1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=c.text_frame.paragraphs[0]; p.alignment=al; r=p.add_run(); r.text=t
    r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col

MONTHS=[7,8,9,10,11,12]
# per-month quantities (index 0=7월)
QTY={ # 항목: 월별
 "매크로":[0,5,0,0,5,5],"마이크로":[15,35,20,15,30,35],"나노":[25,55,35,25,40,45],
 "KOL 무가시딩":[10,10,10,10,10,10],"X(트위터)":[0,5,0,0,5,5],
 "파워페이지":[0,3,0,0,3,3],"커뮤니티 체험단(탑)":[0,1,0,0,1,1],"네이버 블로그":[0,5,0,0,5,5],
 "EGC 콘텐츠":[2,2,1,2,1,2],"어필리에이트(누적)":[10,20,30,40,50,60],
}
QORDER=["매크로","마이크로","나노","KOL 무가시딩","X(트위터)","파워페이지","커뮤니티 체험단(탑)","네이버 블로그","EGC 콘텐츠","어필리에이트(누적)"]
SEEDTOTAL=[50,110,65,50,90,100]
BUDpct=[10,20,20,10,18,22]; BUDamt=["25,000,000","50,000,000","50,000,000","25,000,000","45,000,000","55,000,000"]

DATA={
7:("First Shot","착수 · 콘텐츠/메시지 테스트","1차 웨이브 준비",
   "위닝 메시지 발굴 · 시딩 도화선 점화",
   ["콘텐츠·이미지 소재 LMF 테스트 → 위닝 메시지 선별",
    "일반식품 + 심의 불필요 건기식 시딩 즉시 라이브",
    "건기식 심의 파이프라인 착수 (가이드→섭외→콘티→접수)",
    "챌린저스 1차 가동 · 블로그/검색 키워드 선점"],
   "위닝 메시지 확정 · 후기·콘텐츠 1차 적재"),
8:("Proof Take-off","올영세일 사전 부스팅","★ 1차 웨이브 발생",
   "세일 전 인지·각인 확산 + 선택 증거 확보",
   ["시딩 대량 릴리즈(월 110건) — 세일 1주 전 집중",
    "매크로 합류 · 파워페이지·커뮤니티 체험단(탑)·네이버 블로그 바이럴",
    "건기식 심의 완료분 1차 업로드",
    "‘올영 1위’ 키워드 플레이 · 챌린저스 랭킹 견인"],
   "1차 웨이브로 사전 부스팅 · 검색·후기 증거 확산"),
9:("Acceleration","올영세일 구매 전환","★ 1차 웨이브 피크",
   "확보한 증거로 올영세일 구매 전환 폭발",
   ["올영세일 집중 — 일반+건기식 콘텐츠 동시 투입",
    "퍼포먼스 매체(별첨)·어필리에이트 가동 전환 유도",
    "건기식 2차 업로드 세일 직전 라이브",
    "랭킹/후기/혜택 결합 소재 운영"],
   "1차 웨이브 매출 피크 · 카테고리 랭킹 상위/1위"),
10:("Orbit Expansion","유지 · 포트폴리오 확장","유지 (낙수효과)",
   "구매 증거 기반 건기식 라인업 확장",
   ["건기식 6종 본격 콘텐츠 · 확장 메시지 테스트",
    "9월 위닝 소재 낙수 운영(유지)",
    "크로스셀 CRM · 후기 자산 확산"],
   "매출 유지 · 건기식 라인 인지 형성"),
11:("Boost-up","블프 · 재점화","★ 2차 웨이브 발생",
   "많이 보이고 많이 찾는 브랜드로 재점화",
   ["시딩 재확대(월 90건) · 매크로·체험단(탑)·파워페이지 재가동",
    "블프 프로모션 연계 퍼포먼스 · 건기식 RTB+혜택 결합",
    "챌린저스 2차 · 검색 자산 보강"],
   "2차 웨이브 시작 · 탐색량·관심층 확대"),
12:("Category Landing","올영세일 · 대표성 착지","★ 2차 웨이브 피크",
   "라인업으로 웰니스 카테고리 대표 브랜드 착지",
   ["12월 올영세일 집중 — 전 라인업 대표성 콘텐츠",
    "세일 1주 전 시딩 대량 릴리즈(월 100건)·바이럴",
    "쟁임/재구매 CRM · 퍼포먼스 2차 피크"],
   "2차 웨이브 매출 피크 · 대표성 착지"),
}

# ---------- intro slide: wave timeline ----------
s=prs.slides.add_slide(BLANK)
rect(s,0,0,SW,Inches(1.4),BLACK)
txt(s,Inches(0.6),Inches(0.32),Inches(12),Inches(1.0),
    [("3·4분기 월별 액션플랜",30,True,WHITE),
     ("7월 착수 → 8·9월 1차 웨이브(올영세일) → 10월 유지 → 11·12월 2차 웨이브(블프·올영세일)",14,False,MG)])
wave=["착수","★1차 웨이브","★1차 웨이브","유지","★2차 웨이브","★2차 웨이브"]
shade=[DARK,BLACK,BLACK,LG,BLACK,BLACK]
tcol=[WHITE,WHITE,WHITE,BLACK,WHITE,WHITE]
bw=Inches(1.95); gap=Inches(0.12); x0=Inches(0.6); y0=Inches(2.4)
for i,m in enumerate(MONTHS):
    x=Emu(int(x0)+i*(int(bw)+int(gap)))
    rect(s,x,y0,bw,Inches(1.5),shade[i])
    txt(s,x,Emu(int(y0)+int(Inches(0.2))),bw,Inches(0.5),[(f"{m}월",22,True,tcol[i])],PP_ALIGN.CENTER)
    txt(s,x,Emu(int(y0)+int(Inches(0.78))),bw,Inches(0.5),[(wave[i],12,True,tcol[i])],PP_ALIGN.CENTER)
    txt(s,x,Emu(int(y0)+int(Inches(1.6))),bw,Inches(0.4),[(f"{BUDpct[i]}% · {BUDamt[i].replace(',000,000','백만').replace('000,000','')}",10,False,GRAY)],PP_ALIGN.CENTER)
txt(s,Inches(0.6),Inches(5.2),Inches(12),Inches(1.5),
    [("· Phase 1 (7~9월) — 빠른 반응 확보 : 일반식품·심의불필요 건기식 선행 + 8·9월 1차 웨이브로 올영세일 전환",13,False,DARK),
     ("· Phase 2 (10~12월) — 대표성 확장 : 건기식 라인 확장 + 11·12월 2차 웨이브(블프·올영세일)로 대표 브랜드 착지",13,False,DARK),
     ("· 예산 총 2.5억(공급가) · 퍼포먼스 매체는 별도(별첨)",13,False,GRAY)])

# ---------- monthly slides ----------
for idx,m in enumerate(MONTHS):
    stage,desc,wavetag,mission,actions,goal=DATA[m]
    s=prs.slides.add_slide(BLANK)
    # header
    rect(s,0,0,SW,Inches(1.15),BLACK)
    txt(s,Inches(0.55),Inches(0.16),Inches(7),Inches(0.95),
        [(f"{m}월  ·  {stage}",26,True,WHITE),(desc,13,False,MG)])
    # wave tag (right)
    tagw=Inches(3.1)
    rect(s,Emu(int(SW)-int(tagw)-int(Inches(0.5))),Inches(0.34),tagw,Inches(0.5),
         (LG if "유지" in wavetag else WHITE))
    txt(s,Emu(int(SW)-int(tagw)-int(Inches(0.5))),Inches(0.4),tagw,Inches(0.4),
        [(wavetag,14,True,BLACK)],PP_ALIGN.CENTER)
    # mission quote
    rect(s,Inches(0.55),Inches(1.45),Inches(0.08),Inches(0.55),BLACK)
    txt(s,Inches(0.75),Inches(1.45),Inches(11.8),Inches(0.6),[("“"+mission+"”",16,True,DARK)],anchor=MSO_ANCHOR.MIDDLE)
    # left: actions
    txt(s,Inches(0.55),Inches(2.35),Inches(7.1),Inches(0.4),[("이 달의 핵심 액션",14,True,BLACK)])
    rect(s,Inches(0.57),Inches(2.78),Inches(1.0),Inches(0.04),BLACK)
    tf=s.shapes.add_textbox(Inches(0.55),Inches(2.95),Inches(7.1),Inches(3.4)).text_frame; tf.word_wrap=True
    for i,a in enumerate(actions):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(9)
        r=p.add_run(); r.text="■  "+a; r.font.name=FONT; r.font.size=Pt(13.5); r.font.color.rgb=DARK
    # goal box
    rect(s,Inches(0.55),Inches(6.35),Inches(7.1),Inches(0.75),LG)
    txt(s,Inches(0.7),Inches(6.45),Inches(6.8),Inches(0.6),
        [("성과 목표",10,True,GRAY),(goal,13,True,BLACK)])
    # right: quantities table
    txt(s,Inches(7.95),Inches(2.35),Inches(4.8),Inches(0.4),[("이 달의 진행 물량 (건)",14,True,BLACK)])
    rect(s,Inches(7.97),Inches(2.78),Inches(1.0),Inches(0.04),BLACK)
    rows=QORDER
    t=s.shapes.add_table(len(rows)+2,2,Inches(7.95),Inches(2.95),Inches(4.85),Inches(3.4)).table
    t.columns[0].width=Inches(3.55); t.columns[1].width=Inches(1.30)
    cell(t.cell(0,0),"항목",9.5,True,WHITE,DARK,PP_ALIGN.LEFT); cell(t.cell(0,1),f"{m}월",9.5,True,WHITE,DARK)
    for ri,name in enumerate(rows,1):
        v=QTY[name][idx]
        isaff = name=="어필리에이트(누적)"
        cell(t.cell(ri,0),name,9,False,BLACK,(WHITE if ri%2 else RGBColor(0xF6,0xF6,0xF6)),PP_ALIGN.LEFT)
        disp = (str(v) if v>0 else "·")
        cell(t.cell(ri,1),disp,10,v>0,(BLACK if v>0 else MG),(WHITE if ri%2 else RGBColor(0xF6,0xF6,0xF6)))
    # seeding total row
    cell(t.cell(len(rows)+1,0),"월 시딩 총건수",10,True,WHITE,BLACK,PP_ALIGN.LEFT)
    cell(t.cell(len(rows)+1,1),f"{SEEDTOTAL[idx]}",11,True,WHITE,BLACK)
    for r in t.rows: r.height=Inches(0.27)
    # budget callout
    rect(s,Inches(7.95),Inches(6.35),Inches(4.85),Inches(0.75),BLACK)
    txt(s,Inches(8.1),Inches(6.43),Inches(4.6),Inches(0.6),
        [("이 달 예산",10,True,MG),(f"{BUDamt[idx]} 원   ({BUDpct[idx]}%)",16,True,WHITE)])

prs.save("proposal/올더베러_월별장표_Q3Q4.pptx")
print("slides:",len(prs.slides._sldIdLst))
